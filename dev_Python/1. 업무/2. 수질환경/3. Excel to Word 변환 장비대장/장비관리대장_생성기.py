#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
장비관리대장 생성기 (Python 판)
================================
엑셀(장비대장 / 장비구성 시트) + 사진 폴더(파일명 = 관리번호)  →  ledger.docx 양식의 시험장비관리대장(Word)

[설치]   pip install openpyxl lxml pillow pillow-heif
[실행]   python 장비관리대장_생성기.py                (창이 열립니다)
         ※ 결과: Word 문서 + 장비카탈로그_날짜.html (관리번호·물품명 검색/필터 카탈로그)
[명령행] python 장비관리대장_생성기.py --xlsx 장비.xlsx --images 사진폴더 --out 결과.docx
         옵션: --zip(장비별 파일→ZIP)  --date dot|dash|kr|han  --limit 1024  --target 300
               --only 04-장비-05,04-장비-30   --template 내양식.docx

- 양식(ledger.docx)은 이 파일 맨 아래에 내장되어 있어 따로 필요 없습니다. 양식을 바꾸려면 --template 사용.
- 사진: JPG/PNG/HEIC(아이폰)/WEBP/GIF/BMP/TIFF/AVIF 지원. 1MB 이상이면 300KB 이하로 축소, 비 JPG/PNG는 JPG로 변환.
"""
import argparse, base64, copy, io, zlib, os, re, sys, threading, unicodedata, zipfile
from datetime import date, datetime, timedelta

from lxml import etree
import openpyxl
from PIL import Image, ImageOps

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
    HAVE_HEIF = True
except Exception:          # pillow-heif 미설치
    HAVE_HEIF = False
try:
    import pillow_avif  # noqa: F401
except Exception:
    pass

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
XML_NS = 'http://www.w3.org/XML/1998/namespace'
IMG_EXT = {'jpg', 'jpeg', 'jpe', 'jfif', 'png', 'gif', 'bmp', 'webp', 'heic', 'heif', 'tif', 'tiff', 'avif'}
MAX_COMP_ROWS = 10


def q(n):
    return '{%s}%s' % (W_NS, n)


# ------------------------------------------------------------------ 공통 유틸
def nfc(s):
    return unicodedata.normalize('NFC', '' if s is None else str(s))


def norm_key(s):
    s = nfc(s)
    s = re.sub('[‐-―−ー－]', '-', s)
    return re.sub(r'\s+', '', s).lower()


def norm_hdr(s):
    return re.sub(r'\s+', '', nfc(s)).lower()


def clean_xml(s):
    s = '' if s is None else str(s)
    s = re.sub('[\u0000-\u0008\u000b\u000c\u000e-\u001f￾￿]', '', s)
    return s.replace('\r\n', '\n').replace('\r', '\n')


def sval(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def fmt_bytes(n):
    return '%.2fMB' % (n / 1048576) if n >= 1048576 else '%dKB' % round(n / 1024)


def safe_name(s):
    return re.sub(r'[\\/:*?"<>|\r\n]+', '_', nfc(s)).strip()


# ------------------------------------------------------------------ 엑셀 읽기
def col_idx(hdr, names, last=False):
    keys = [norm_hdr(n) for n in names]
    found = -1
    for i, h in enumerate(hdr):
        if h in keys:
            found = i
            if not last:
                return i
    return found


def fmt_date(v, style='dot'):
    y = m = d = None
    if isinstance(v, (datetime, date)):
        y, m, d = v.year, v.month, v.day
    elif isinstance(v, (int, float)) and v > 0:
        t = datetime(1899, 12, 30) + timedelta(days=float(v))
        y, m, d = t.year, t.month, t.day
    elif isinstance(v, str):
        mm = re.search(r'(\d{4})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})', v)
        if mm:
            y, m, d = int(mm[1]), int(mm[2]), int(mm[3])
    if y is None:
        return '' if v in (None, '') else str(v)
    if style == 'dash':
        return '%d-%02d-%02d' % (y, m, d)
    if style == 'kr':
        return '%d. %d. %d.' % (y, m, d)
    if style == 'han':
        return '%d년 %d월 %d일' % (y, m, d)
    return '%d.%02d.%02d' % (y, m, d)


def fmt_price(v):
    if v in (None, ''):
        return ''
    try:
        n = v if isinstance(v, (int, float)) else float(re.sub(r'[,\s원]', '', str(v)))
        return '{:,}'.format(int(round(n)))
    except Exception:
        return str(v)


def pick_sheets(names):
    led = next((n for n in names if '대장' in n and '폐기' not in n), names[0])
    comp = next((n for n in names if re.search('component|구성', n, re.I)), '')
    return led, comp


def parse_ledger(rows):
    hi = next((i for i, r in enumerate(rows) if any(norm_hdr(c) == '관리번호' for c in r)), -1)
    if hi < 0:
        raise ValueError("장비대장 시트에서 '관리번호' 열을 찾지 못했습니다.")
    hdr = [norm_hdr(c) for c in rows[hi]]
    C = dict(
        no=col_idx(hdr, ['순번']), id=col_idx(hdr, ['관리번호']), item=col_idx(hdr, ['물품번호']),
        name=col_idx(hdr, ['물품명(내부)', '물품명', '품명']), spec=col_idx(hdr, ['규격', '형식및규격']),
        maker=col_idx(hdr, ['제조사', '제작회사']), country=col_idx(hdr, ['제조국']),
        date=col_idx(hdr, ['취득일(확정)']), date2=col_idx(hdr, ['구입일(내부)']),
        price=col_idx(hdr, ['내부금액(원)']), qty=col_idx(hdr, ['수량']), vendor=col_idx(hdr, ['도입회사명', '회사명']),
        addr=col_idx(hdr, ['도입처주소', '주소']), tel=col_idx(hdr, ['연락처', '전화번호']),
        usage=col_idx(hdr, ['용도']), cond=col_idx(hdr, ['상태'], last=True),   # 마지막 '상태' = 입고시 상태
        note=col_idx(hdr, ['내용', '주요성능및특징']))
    if C['id'] < 0 or C['name'] < 0:
        raise ValueError('필수 열(관리번호, 물품명)을 찾지 못했습니다.')

    def get(r, k):
        return r[C[k]] if C[k] >= 0 and C[k] < len(r) else None

    out = []
    for i in range(hi + 1, len(rows)):
        r = rows[i]
        rid = sval(get(r, 'id')).strip()
        if not rid:
            continue
        d = get(r, 'date')
        out.append(dict(
            row=i + 1, no=sval(get(r, 'no')), id=rid, key=norm_key(rid), item=sval(get(r, 'item')).strip(),
            name=re.sub(r'\s*\n\s*', ' ', sval(get(r, 'name'))).strip(), spec=sval(get(r, 'spec')).strip(),
            maker=sval(get(r, 'maker')).strip(), country=sval(get(r, 'country')).strip(),
            date=d if d not in (None, '') else get(r, 'date2'), price=get(r, 'price'), qty=get(r, 'qty'),
            vendor=sval(get(r, 'vendor')).strip(), addr=sval(get(r, 'addr')).strip(),
            tel=sval(get(r, 'tel')).strip(), usage=sval(get(r, 'usage')).strip(),
            cond=sval(get(r, 'cond')).strip(), note=sval(get(r, 'note')).rstrip(), img=None))
    return out


def parse_components(rows):
    comp = {}
    if not rows:
        return comp
    hi = next((i for i, r in enumerate(rows) if any(norm_hdr(c) == '관리번호' for c in r)), -1)
    if hi < 0:
        raise ValueError("장비구성 시트에서 '관리번호' 열을 찾지 못했습니다.")
    hdr = [norm_hdr(c) for c in rows[hi]]
    ci = col_idx(hdr, ['관리번호'])

    def fz(names, word):
        i = col_idx(hdr, names)
        if i < 0:
            i = next((k for k, h in enumerate(hdr) if k != ci and word in h), -1)
        return i
    cn = fz(['내용', '장비명', '구성품', '구성품명', '품명', '명칭', '물품명', 'name'], '명')
    cq = fz(['수량', 'qty', 'quantity'], '수량')
    cu = fz(['단위', 'unit'], '단위')
    cs = col_idx(hdr, ['연번', '순번', 'no'])
    if cn < 0:
        raise ValueError("장비구성 시트에서 장비명('내용') 열을 찾지 못했습니다.")
    for i in range(hi + 1, len(rows)):
        r = rows[i]
        rid = sval(r[ci]).strip() if ci < len(r) else ''
        name = re.sub(r'\s*\n\s*', ' ', sval(r[cn])).strip() if cn < len(r) else ''
        if not rid or not name:
            continue
        try:
            sn = float(r[cs]) if cs >= 0 and r[cs] not in (None, '') else i
        except Exception:
            sn = i
        comp.setdefault(norm_key(rid), []).append(dict(
            name=name, qty=sval(r[cq]).strip() if cq >= 0 and cq < len(r) else '',
            unit=sval(r[cu]).strip() if cu >= 0 and cu < len(r) else '', sn=sn))
    for v in comp.values():
        v.sort(key=lambda x: x['sn'])
    return comp


def load_excel(path, ledger_sheet=None, comp_sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    led, comp = pick_sheets(wb.sheetnames)
    led = ledger_sheet or led
    comp = wb.sheetnames and (comp if comp_sheet is None else comp_sheet)
    rows = lambda n: [list(r) for r in wb[n].iter_rows(values_only=True)]
    records = parse_ledger(rows(led))
    comps = parse_components(rows(comp)) if comp else {}
    return records, comps, wb.sheetnames, led, comp


# ------------------------------------------------------------------ 사진 찾기·변환
def scan_images(folder):
    out = []
    for root, _, files in os.walk(folder):
        for f in files:
            ext = f.rsplit('.', 1)[-1].lower() if '.' in f else ''
            if ext in IMG_EXT:
                out.append(dict(path=os.path.join(root, f), name=f, key=norm_key(f.rsplit('.', 1)[0])))
    out.sort(key=lambda e: [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', e['name'])])
    return out


def find_image(rec, index):
    exact = [e for e in index if e['key'] == rec['key']]
    if exact:
        return exact[0], len(exact)
    rx = re.compile('^' + re.escape(rec['key']) + r'(?![0-9a-z가-힣])')
    part = [e for e in index if rx.match(e['key'])]
    return (part[0], len(part)) if part else (None, 0)


def sniff(b):
    if b[:2] == b'\xff\xd8': return 'jpg'
    if b[:4] == b'\x89PNG': return 'png'
    if b[:3] == b'GIF': return 'gif'
    if b[:2] == b'BM': return 'bmp'
    if b[:4] == b'RIFF' and b[8:12] == b'WEBP': return 'webp'
    if b[:4] in (b'II*\x00', b'MM\x00*'): return 'tiff'
    if b[4:8] == b'ftyp':
        return 'avif' if b[8:12] in (b'avif', b'avis') else 'heic'
    return 'unknown'


def _jpeg_bytes(im, quality):
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=quality, optimize=True)
    return buf.getvalue()


def _to_rgb(im):
    if im.mode in ('RGBA', 'LA') or (im.mode == 'P' and 'transparency' in im.info):
        im = im.convert('RGBA')
        bg = Image.new('RGB', im.size, (255, 255, 255))
        bg.paste(im, mask=im.split()[-1])
        return bg
    return im.convert('RGB')


def _resized(im, cap):
    sc = min(1.0, cap / max(im.size))
    if sc >= 1.0:
        return im
    return im.resize((max(1, round(im.width * sc)), max(1, round(im.height * sc))), Image.LANCZOS)


def _compress(im, target):
    best = None
    for cap in (2000, 1600, 1280, 1024, 800, 640):
        r = _resized(im, cap)
        lo_b = _jpeg_bytes(r, 40)
        if len(lo_b) > target:
            best = (lo_b, r.size, 40)
            continue
        hi_b = _jpeg_bytes(r, 92)
        if len(hi_b) <= target:
            return hi_b, r.size, 92
        lo, hi, ok = 40, 92, (lo_b, 40)
        while hi - lo > 1:
            mid = (lo + hi) // 2
            b = _jpeg_bytes(r, mid)
            if len(b) <= target:
                lo, ok = mid, (b, mid)
            else:
                hi = mid
        return ok[0], r.size, ok[1]
    return best


def process_image(path, limit_kb=1024, target_kb=300):
    limit, target = limit_kb * 1024, target_kb * 1000
    raw = open(path, 'rb').read()
    orig = len(raw)
    kind = sniff(raw)
    if kind == 'unknown':
        raise ValueError('지원하지 않는 이미지 형식')
    if kind == 'heic' and not HAVE_HEIF:
        raise ValueError('HEIC 변환에는 pillow-heif 설치가 필요합니다 (pip install pillow-heif)')
    im = Image.open(io.BytesIO(raw))
    orient = im.getexif().get(0x0112, 1) if kind == 'jpg' else 1
    label = kind.upper()
    if kind in ('jpg', 'png') and orig < limit and orient == 1:
        return dict(bytes=raw, ext=kind, w=im.width, h=im.height, note='%s %s 그대로' % (label, fmt_bytes(orig)))
    im = _to_rgb(ImageOps.exif_transpose(im))
    im = _resized(im, 2400)
    data, size, tag = _jpeg_bytes(im, 92), im.size, ''
    if orig >= limit or len(data) >= limit:
        data, size, qq = _compress(im, target)
        tag = ' (%dx%d, 화질 %d)' % (size[0], size[1], qq) + (' ※목표 초과' if len(data) > target else '')
    act = ('회전 보정' if orient > 1 else '축소') if kind in ('jpg', 'png') else '변환'
    return dict(bytes=data, ext='jpg', w=size[0], h=size[1],
                note='%s %s → JPG %s%s [%s]' % (label, fmt_bytes(orig), fmt_bytes(len(data)), tag, act))


# ------------------------------------------------------------------ 글자 크기 자동 맞춤
def char_w(c):
    o = ord(c)
    if c == ' ': return 0.32
    if o < 128:
        if c.isdigit(): return 0.52
        if 'A' <= c <= 'Z': return 0.66
        if 'a' <= c <= 'z': return 0.52
        if c in ".,:;'`|!": return 0.34
        if c in '()[]-/': return 0.4
        return 0.6
    return 1.0


def count_lines(text, usable, size):
    lines = 0
    for para in clean_xml(text).split('\n'):
        w, l = 0.0, 1
        for ch in para:
            cw = char_w(ch) * size * 1.04
            if w + cw > usable:
                l += 1
                w = cw
            else:
                w += cw
        lines += l
    return lines


def fit_size(text, width_tw, height_tw, max_pt, min_pt):
    usable, avail = (width_tw - 140) / 20, (height_tw - 40) / 20
    s = max_pt
    while s >= min_pt:
        if count_lines(text, usable, s) * s * 1.22 <= avail:
            return s
        s -= 0.5
    return min_pt


half_pt = lambda s: int(round(s * 2))


# ------------------------------------------------------------------ DOCX 편집
def kids(n, name):
    return [c for c in n if isinstance(c.tag, str) and etree.QName(c).localname == name]


rows_of = lambda t: kids(t, 'tr')
cells_of = lambda tr: kids(tr, 'tc')
pars_of = lambda tc: kids(tc, 'p')

SZ_AFTER = ['szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang',
            'eastAsianLayout', 'specVanish', 'oMath']


def set_sz(rpr, hp):
    for c in list(rpr):
        if etree.QName(c).localname in ('sz', 'szCs'):
            rpr.remove(c)
    sz, szcs = etree.Element(q('sz')), etree.Element(q('szCs'))
    sz.set(q('val'), str(hp))
    szcs.set(q('val'), str(hp))
    ref = next((c for c in rpr if etree.QName(c).localname in SZ_AFTER), None)
    if ref is not None:
        ref.addprevious(sz)
        ref.addprevious(szcs)
    else:
        rpr.append(sz)
        rpr.append(szcs)


def run_rpr(p):
    runs = kids(p, 'r')
    if runs:
        rp = kids(runs[0], 'rPr')
        return copy.deepcopy(rp[0]) if rp else None
    ppr = kids(p, 'pPr')
    pr = kids(ppr[0], 'rPr') if ppr else []
    if pr:
        rp = etree.Element(q('rPr'))
        for c in pr[0]:
            if etree.QName(c).localname not in ('ins', 'del', 'moveFrom', 'moveTo', 'rPrChange'):
                rp.append(copy.deepcopy(c))
        return rp
    return None


def set_par(p, text, sz=None):
    rpr = run_rpr(p)
    for r in kids(p, 'r'):
        p.remove(r)
    if sz:
        if rpr is None:
            rpr = etree.Element(q('rPr'))
        set_sz(rpr, sz)
    t = clean_xml(text)
    if t == '':
        return
    r = etree.SubElement(p, q('r'))
    if rpr is not None:
        r.append(rpr)
    for i, ln in enumerate(t.split('\n')):
        if i:
            etree.SubElement(r, q('br'))
        e = etree.SubElement(r, q('t'))
        e.set('{%s}space' % XML_NS, 'preserve')
        e.text = ln


def drawing_xml(rid, did, name, cx, cy):
    name = name.replace('&', '&amp;').replace('<', '&lt;').replace('"', '&quot;')
    return (
        '<w:r xmlns:w="%s" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:drawing>'
        '<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="%d" cy="%d"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:docPr id="%d" name="%s"/>'
        '<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
        '<pic:nvPicPr><pic:cNvPr id="%d" name="%s"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic>'
        '</wp:inline></w:drawing></w:r>') % (W_NS, cx, cy, did, name, did, name, rid, cx, cy)


class DocBuilder:
    def __init__(self, template_bytes):
        self.zin = zipfile.ZipFile(io.BytesIO(template_bytes))
        self.files = {i.filename: self.zin.read(i.filename) for i in self.zin.infolist() if not i.is_dir()}
        self.root = etree.fromstring(self.files['word/document.xml'])
        self.body = self.root.find(q('body'))
        ch = list(self.body)
        self.sect = next(c for c in ch if etree.QName(c).localname == 'sectPr')
        self.pristine = [c for c in ch if c is not self.sect]
        sig = ','.join(etree.QName(c).localname for c in self.pristine)
        if sig != 'tbl,tbl,p,p,p,tbl,p,p':
            raise ValueError('양식 구조가 예상과 다릅니다: ' + sig)
        for c in self.pristine:
            self.body.remove(c)
        self.rels, self.exts, self.seq, self.docpr, self.relof = [], set(), 0, 5000, {}

    def add_photo(self, photo, key):
        if key in self.relof:
            return self.relof[key]
        self.seq += 1
        rid, path = 'rIdPhoto%d' % self.seq, 'media/photo%d.%s' % (self.seq, photo['ext'])
        self.files['word/' + path] = photo['bytes']
        self.rels.append('<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                         'relationships/image" Target="%s"/>' % (rid, path))
        self.exts.add(photo['ext'])
        self.relof[key] = rid
        return rid

    def append_record(self, rec, comps, photo, photo_key, first, date_style='dot'):
        f = [copy.deepcopy(n) for n in self.pristine]
        for n in f:
            for e in list(n.iter(q('lastRenderedPageBreak'))):
                e.getparent().remove(e)
            for e in n.iter():
                if isinstance(e.tag, str) and etree.QName(e).localname in ('docPr', 'cNvPr') and e.get('id'):
                    self.docpr += 1
                    e.set('id', str(self.docpr))
        f = f[:7]          # 마지막 빈 문단은 쓰지 않음
        t0, t1, p_front, p_break, p_head, t2, p_back = f
        name_line, id_line = '물품명 : ' + rec['name'], '관리번호 (%s)' % rec['id']
        R = rows_of(t1)
        c0, c1 = cells_of(R[0]), cells_of(R[1])
        hs = half_pt(fit_size(name_line, 6773, 620, 12, 8.5))
        set_par(pars_of(c0[0])[0], name_line, hs if hs < 24 else None)
        set_par(pars_of(c0[1])[0], id_line)
        set_par(pars_of(c1[1])[0], '물품번호 ' + rec['item'] if rec['item'] else '')

        r4, H3 = cells_of(R[4]), 1039

        def fill(tc, text, w, h, mx=10, mn=6.5):
            s = fit_size(text, w, h, mx, mn)
            set_par(pars_of(tc)[0], text, half_pt(s) if s < mx else None)
        fill(r4[0], rec['spec'], 1563, H3)
        ps = pars_of(r4[1])
        mk, co = rec['maker'], ('(%s)' % rec['country'] if rec['country'] else '')
        s = min(fit_size(mk, 1563, H3 / 2, 10, 6.5), fit_size(co, 1563, H3 / 2, 10, 6.5))
        set_par(ps[0], mk, half_pt(s) if s < 10 else None)
        set_par(ps[1], co, half_pt(s) if s < 10 else None)
        fill(r4[2], fmt_date(rec['date'], date_style), 1563, H3)
        set_par(pars_of(r4[3])[0], fmt_price(rec['price']))
        fill(r4[4], rec['vendor'], 1563, H3)
        fill(r4[5], rec['addr'], 1563, H3)
        fill(r4[6], rec['tel'], 1791, H3)
        fill(r4[7], rec['usage'], 2814, H3)

        r6 = cells_of(R[6])
        used, is_new = '중고' in rec['cond'], bool(re.search('신품|신규|새', rec['cond']))
        set_par(pars_of(r6[3])[0], ('■' if is_new else '□') + ' 신품')
        set_par(pars_of(r6[4])[0], ('■' if used else '□') + '중고품')

        r7 = cells_of(R[7])
        ps = pars_of(r7[0])
        s = fit_size(rec['note'], 6773, 2840 + 400 - 12 * 20, 10, 6)
        set_par(ps[1] if len(ps) > 1 else ps[0], rec['note'], half_pt(s) if s < 10 else None)

        comps = comps[:MAX_COMP_ROWS]
        comp_rows = R[10:15]
        total = max(5, len(comps))
        while len(comp_rows) < total:
            nr = copy.deepcopy(comp_rows[-1])
            comp_rows[-1].addnext(nr)
            comp_rows.append(nr)
        row_h = 2615 // total if total > 5 else 523
        base = 9 if total > 5 else 10
        for i, tr in enumerate(comp_rows):
            for h in tr.iter(q('trHeight')):
                h.set(q('val'), str(row_h))
            cs, cp = cells_of(tr), (comps[i] if i < len(comps) else None)
            set_par(pars_of(cs[0])[0], str(i + 1), half_pt(base) if total > 5 else None)
            nm = cp['name'] if cp else ''
            s = min(base, fit_size(nm, 3685, row_h, base, 6.5))
            set_par(pars_of(cs[1])[0], nm, half_pt(s) if s < 10 else None)
            set_par(pars_of(cs[2])[0], cp['qty'] if cp else '', half_pt(base) if total > 5 else None)
            set_par(pars_of(cs[3])[0], cp['unit'] if cp else '', half_pt(base) if total > 5 else None)

        p_photo = pars_of(r7[1])[0]
        for r in kids(p_photo, 'r'):
            p_photo.remove(r)
        if photo:
            rid = self.add_photo(photo, photo_key)
            box_w, box_h = (7733 - 140) * 635, (2840 + row_h * total + 523 * 2 - 330) * 635
            sc = min(box_w / photo['w'], box_h / photo['h'])
            cx, cy = round(photo['w'] * sc), round(photo['h'] * sc)
            self.docpr += 1
            p_photo.append(etree.fromstring(drawing_xml(rid, self.docpr, '사진 ' + rec['id'], cx, cy).encode('utf8')))

        rs = kids(p_head, 'r')
        kids(rs[0], 't')[0].text = name_line
        kids(rs[1], 't')[0].text = id_line

        if not first:
            br = copy.deepcopy(self.pristine[3])
            for e in list(br.iter(q('lastRenderedPageBreak'))):
                e.getparent().remove(e)
            self.body.append(br)
        for n in f:
            self.body.append(n)

    def finish(self):
        self.body.append(self.sect)
        self.files['word/document.xml'] = etree.tostring(self.root, xml_declaration=True, encoding='UTF-8', standalone=True)
        rels = self.files['word/_rels/document.xml.rels'].decode('utf8')
        self.files['word/_rels/document.xml.rels'] = rels.replace('</Relationships>', ''.join(self.rels) + '</Relationships>').encode('utf8')
        ct = self.files['[Content_Types].xml'].decode('utf8')
        mime = {'jpg': 'image/jpeg', 'png': 'image/png'}
        for e in self.exts:
            if 'Extension="%s"' % e not in ct:
                ct = ct.replace('<Default ', '<Default Extension="%s" ContentType="%s"/><Default ' % (e, mime[e]), 1)
        self.files['[Content_Types].xml'] = ct.encode('utf8')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            order = ['[Content_Types].xml'] + [n for n in self.files if n != '[Content_Types].xml']
            for n in order:
                z.writestr(n, self.files[n])
        return buf.getvalue()



# ------------------------------------------------------------------ 카탈로그(HTML)
def _thumb_data_url(photo, max_edge=800, quality=74):
    im = Image.open(io.BytesIO(photo['bytes']))
    im = _resized(_to_rgb(ImageOps.exif_transpose(im)), max_edge)
    return 'data:image/jpeg;base64,' + base64.b64encode(_jpeg_bytes(im, quality)).decode()


def _date_parts(v):
    if isinstance(v, (datetime, date)):
        return v.year, v.month, v.day
    s = fmt_date(v, 'dash')
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})$', s)
    return (int(m[1]), int(m[2]), int(m[3])) if m else None


def build_catalog_html(records, comps, photos, date_style='dot', thumb_edge=800, source=''):
    import json
    dup = {}
    for r in records:
        dup[r['key']] = dup.get(r['key'], 0) + 1
    images, key_of, out = {}, {}, []
    for rec, ph in zip(records, photos):
        pk = None
        if ph and rec.get('img'):
            ck = rec['img']['path']
            if ck not in key_of:
                try:
                    key_of[ck] = 'i%d' % len(key_of)
                    images[key_of[ck]] = _thumb_data_url(ph, thumb_edge)
                except Exception:
                    key_of[ck] = None
            pk = key_of[ck]
        dp = _date_parts(rec['date'])
        try:
            qn = float(re.sub(r'[^0-9.]', '', sval(rec.get('qty')))); qn = int(qn) if qn == int(qn) else qn
        except Exception:
            qn = 1
        try:
            pn = None if rec['price'] in (None, '') else float(re.sub(r'[,\s원]', '', str(rec['price'])))
            pn = int(pn) if pn is not None and pn == int(pn) else pn
        except Exception:
            pn = None
        out.append(dict(
            no=rec['no'], id=rec['id'], item=rec['item'], name=rec['name'], spec=rec['spec'], maker=rec['maker'],
            country=rec['country'], date=fmt_date(rec['date'], date_style),
            dateSort=('%04d%02d%02d' % dp) if dp else '', year=dp[0] if dp else None,
            price=pn, qty=qn if qn and qn > 0 else 1, vendor=rec['vendor'], addr=rec['addr'], tel=rec['tel'],
            usage=rec['usage'], cond=rec['cond'], note=rec['note'],
            comps=[dict(n=c['name'], q=c['qty'], u=c['unit']) for c in comps.get(rec['key'], [])],
            ph=pk, dupLabel=('순번 %s' % (rec['no'] or rec['row'])) if dup[rec['key']] > 1 else ''))
    data = dict(records=out, images=images, source=source, generated=datetime.now().strftime('%Y.%m.%d'))
    payload = json.dumps(data, ensure_ascii=False).replace('<', '\\u003c')
    tpl = zlib.decompress(base64.b64decode(CATALOG_TPL_B64)).decode('utf8')
    return tpl.replace('/*__DATA__*/null', payload)


# ------------------------------------------------------------------ 실행 본체
def generate(records, comps, index, out_path, template_bytes, zip_mode=False, date_style='dot',
             limit_kb=1024, target_kb=300, log=print, progress=lambda a, b: None,
             catalog=False, thumb_edge=800, source=''):
    """records: 변환할 장비 목록. 반환: 저장된 파일 경로"""
    n = len(records)
    log('선택 %d건 변환 시작 (%s)' % (n, '장비별 파일→ZIP' if zip_mode else '하나의 파일'))
    cache, photos, stat = {}, [], dict(photo=0, none=0, conv=0, err=0)
    for k, rec in enumerate(records):
        ph = None
        if index:
            e, multi = find_image(rec, index)
            rec['img'] = e
            if e:
                if e['path'] not in cache:
                    try:
                        cache[e['path']] = process_image(e['path'], limit_kb, target_kb)
                        log('[%s] %s: %s' % (rec['id'], e['name'], cache[e['path']]['note']))
                        if multi > 1:
                            log('[%s] 같은 관리번호 사진 %d장 — 이름순 첫 장 사용' % (rec['id'], multi))
                    except Exception as ex:
                        cache[e['path']] = None
                        log('[%s] 사진 오류 %s — %s' % (rec['id'], e['name'], ex))
                        stat['err'] += 1
                ph = cache[e['path']]
        if ph:
            stat['photo'] += 1
            stat['conv'] += '→' in ph['note']
        else:
            stat['none'] += 1
        photos.append(ph)
        progress(k + 1, n * 2)
    for rec in records:
        if len(comps.get(rec['key'], [])) > MAX_COMP_ROWS:
            log('[%s] 구성내역 %d행 중 %d행만 반영' % (rec['id'], len(comps[rec['key']]), MAX_COMP_ROWS))
    dup = {}
    for r in records:
        dup[r['key']] = dup.get(r['key'], 0) + 1
    cat_html = build_catalog_html(records, comps, photos, date_style, thumb_edge, source) if catalog else None
    if zip_mode:
        out_path = os.path.splitext(out_path)[0] + '.zip'
        used = set()
        with zipfile.ZipFile(out_path, 'w', zipfile.ZIP_STORED) as z:
            for k, rec in enumerate(records):
                b = DocBuilder(template_bytes)
                b.append_record(rec, comps.get(rec['key'], []), photos[k], rec['img']['path'] if rec['img'] else '', True, date_style)
                nm = safe_name(rec['id']) + ('_' + (rec['no'] or str(rec['row'])) if dup[rec['key']] > 1 else '') + '_' + safe_name(rec['name'])[:30].strip().rstrip('. ')
                while nm in used:
                    nm += '_x'
                used.add(nm)
                z.writestr(nm + '.docx', b.finish())
                progress(n + k + 1, n * 2)
            if cat_html:
                z.writestr('장비카탈로그.html', cat_html)
    else:
        b = DocBuilder(template_bytes)
        for k, rec in enumerate(records):
            b.append_record(rec, comps.get(rec['key'], []), photos[k], rec['img']['path'] if rec['img'] else '', k == 0, date_style)
            progress(n + k + 1, n * 2)
        open(out_path, 'wb').write(b.finish())
    log('완료: %s (%s)' % (out_path, fmt_bytes(os.path.getsize(out_path))))
    if cat_html and not zip_mode:
        cp = os.path.join(os.path.dirname(os.path.abspath(out_path)), '장비카탈로그_%s.html' % datetime.now().strftime('%Y%m%d'))
        open(cp, 'w', encoding='utf8').write(cat_html)
        log('완료: %s (%s)' % (cp, fmt_bytes(os.path.getsize(cp))))
    log('사진 반영 %d건 · 사진 없음 %d건 · 변환/축소 %d건 · 오류 %d건' % (stat['photo'], stat['none'], stat['conv'], stat['err']))
    return out_path


# ------------------------------------------------------------------ 내장 카탈로그 양식 (zlib+base64)
CATALOG_TPL_B64 = (
    "eNq1O/2TE0d2v+9f0Yw5ZgZGs9J+AdKuCF6TxGXsc4EvVa5lA6OZljTWaEbMtPbDslLExq61oep8znEQ3+LDlTg+X3BqzcEZV5xK"
    "lfOf8ONKW/kX8l53z5c+lj2SFIU00/36vdev33drl4+98vPVt95+8wJpsrZXnVnGL+JZfmNFaQUKDlDLga82ZRaxm1YYUbaidFm9"
    "cEaJh32rTVeUDZdudoKQKcQOfEZ9ANt0HdZcceiGa9MCfzGI67vMtbxCZFseXSkhEuYyj1aHt3YP7t4bfvHV4IebZPjD44MPdgZf"
    "7u5//3R5VgDMLEdsG7/LYRCwXqFQa5Rfqi/Ul+pnKoWCbYUOvNbr8Oz6rfJLJXvu9Pw8vLW7jMLUorNUO30W3j3Xp+WXnLO0SJfg"
    "tRZaPkwX60u1hXr8LlAIbFFQZ+WX6Hx9jtrw2vT4RMlCqqHluN2oXCp2tvozf9GmjmsRrRPSOg2jgh14QQgbbdI2LRPHClt6L8t7"
    "aa50umQnvJdqc/NztZh9ukTteill/6xlna0tJezPOfNnFxZS9ufrtaUztTz7xTOlxWK6g9Lp+eLcktzBorVgFYv9/szJXi3YKkTu"
    "u67fKNeC0KFhAUb6qAhGLXC2e20rbLh+udif4a81y241wqALVDesUMOt6BW+VfkOtPVKHXSgXFrobM2WzEWivG55ja5P/ipgTddW"
    "DOV8p+NRcvkVOULeoAGMvhGwgFy2/Ii8dkkxou2I0Xah6xoRDBUiGrr1/gzqIw3HuUAZ6pVkA4wF7XKps0WiwHMdImBQdHqlE0Sg"
    "goFfjphrt7YrLOiUi5V3gW+HbpXnYJ9m04FdbwmVLZfmzsDxVmIxEKvLgkrHchyUWGkOaJSW8PhNrqY9x406nrVdrnt0q2J5bsMv"
    "uLCPqFyzIoocVBpWh6+rIEhhM4RX/IC9lXooNzwNWi6dzRDtm2hpvayYuV4IQcsFc+Yi56Puegz0L+GkEbpOBT8KwAeMMIqq2W37"
    "UTmkHWoxDbdUqLvMaLs+bFwrnYYdG6V6qOuc2zO4ybmEnwKKDHU+tz/qO0g8LwC+RccNqc1FLuhynPPAq1mvRh3Lz+4aqYxvsz/j"
    "+p0uW2PbHbrC6BZbNzIDfrddo+G6AfIFOhxb2fWboDAsOajTsAWUqNCQKaoh1Uda9WkEP5K2S0UpFn9WAQlKvSlKpsv1wO5Gkjnx"
    "0gu6jNvx3Agb3Hr1ipwuBPU6+FvEZIbBZl6yKMOlkSOwwe3ScEyvzE4IHvgo+gBHzvWbxA+czMJEMuAPumBmfl7eh8t3stWOyTN7"
    "aFwjxg/G7oYRLOoEbpaXcjPYQO8gwLN4hWT7AswErsdPVsh+bJFgaTrKGbNmhXPjhzNmIBNPZ9Sk+hwbMaNODwHLpf6MZ9WoZ9qt"
    "cRKLkymMyqZtuf6RHdqCdGhkiQc10/ZZHAJQW4uE72uCiSIkqfVGDxNQoK79+d7I82J3NDeXd0eoEsgZKM/hgeBIZi5mxAuYHuhP"
    "3Qs2y03Xcag/IsvKEZwbA7WQEcYszUWS0edpplhWD8J2mT+hYN7WCrBTlGCn2bOiDtABPgFzeWF2frJrykkZniDxyqvG4REETnZE"
    "AJw4cduNXsbJNanbaDLxHNTeQb4geoAANjIiqnmB3cLd13qxcp1F3SrJD1AYUAtOfVPgO10sViZIJsMgDxq2384tW4JlUkNLXEMF"
    "FFjQEeLl2IFjdCkkg9Tz3E7kRpXNJsixANHKpmU/EH4ViLbGFbDpxU4k7xNjfUMhxwIBaxKpw2YvIcnN8f9YrfszzKpBcpKqn2d1"
    "IlqOH7IRbETckLQzp5fzyMXUIx+aaXFRcv9U9midVWCHkHJB8i/GwOUBX83e5IRsigQmnAPDvJSwsDfi9+JxYXqETXAWmBnrfeaY"
    "fi/DbIhqNfHETdaEXEOawiKGX2kKC9x0Rk0hfyIYRnPGMdGEgQj4RbadiHwRxZ3hbqolw8p24FheKtC6u0WdiuvzJCJLLWzULK1o"
    "4D9zcTF1G34A+WmcCi8W06iwlLUU1NCYmBk1M5lJbPOO15jmmdM4dHZCGMpLbIpDcpr5UPhOFzSnvl2QhWeZn1mhRtkmBYOeECCT"
    "bRVlsDuSOstMmDNAmnPZrPV0Lld3akeJdnFwW4qDGymZ8/VQRLiljJMQ5YWoLrVUfKdxod5DalOTuT5mJ27Od+clfOa5KfGkKINJ"
    "bhwDkIsXCDt95Gs0rozzNqJYjnekLFacq5TlYuyykgNy2ITI0HecpNTN2T6U8zxdq2wCdxCTqNUq888CDkDVNp/VhIWUEE+ligRP"
    "z4y6tUwhLTsFkJZDciUqWoPnfYa04Kw9HnPb2FexfAanuTwrmyDLs7Izg05O9mloWF123A1ie1YUrShNR6nOEJId4iWqUl1ulg7r"
    "t8DsMtZl8SqsPRXiOvKpCkzALHwB5lEKsvjklGGG567JnCLwVvcf3xj8y8PBo52De0+JNvz4x4Nf7QJxMrz7cPC7h2R/b1ePafAS"
    "ivAyT0EnKPiov+oohOtZM/Bg3yvK8N5OmRQXCmJDheKiQeaLCncqdgD6QRkggHoKuedMHcrg4OGPB5/tDP7w4Ytx94bVppP4G/zp"
    "5v7jXxlk+ODG8POnL8zdwUdPDu58Q/Yf3Rh+8DHR9r/f3X/0zU/fDx/sDr/cG77/8KfvB7+8Ofziw+GjezD6+TfwBkPvP4bH53F+"
    "3t+exPhlK2RBCEYJvO/cG3698+KSFZzd3YOHmBdRGgsG3qZWyDWMjx0Jpdj3/pNvJ+FbBdfFwu2pKLOqqxBuXCsKdyvCm5S5Icxl"
    "6/pEi5/uDO98SbLkOi9bkCwmeMSCXH2FNRT4BcARdDBGg7P1ugDa9V2mVAe3fr+/d2N5VsyNwrCAWZ5SHT6+D3Q1Aftfd+FABr/7"
    "WE8XpRvlbPJ95nfK+wFKXglEE0WoQed11x9Vgye7w49uKwTksKIkIvi7SeqUx2RtjWMa3L6RYor5y4rxFyiNcSmOSa0E0vjtp9Pk"
    "BTGlCD5AIKZOdfD17edB8xXV4W+e5ABTiaLHS/3eJAeIbnyi97NbIxK3m9Ru1YItaXpvNgPcIQEDHn4NLvmLncEnvybCnwHnh1pC"
    "LKqRijQMNielP7J1lNrPncGDhyPakjmNiHf1n38YLsSbrGefJmi8L1BSFzsNzIGQrsT+4v6PGmrgrQfDnV192gqu1q8osWEOPvps"
    "eP8GLDgU/nwK//6/jcJP9hnZ2AgiBkGIthKX1sYqpLlKPB34oE0/PB78AyAVQHngt7AsU6oQYdL5zEHkSEUdJT+ZQVRj/iUKaT6g"
    "unPz4OYeGT7e2X+6d/CPv55IF8BXow3Y+aO9/T/+SFYv/w2PSV98lUCn6h0rvcwwZpaxnTSq9bYvIwg+jGQGOBx00+HlWYEhu56n"
    "PTLH4I+5TAbKCTGHD3krXI7s0O2w6gwk/xEjr5x/6zxZIbMnr17Fx6tXT876Xc+ryOnjMAcJ7kqVOIHdbYM5mA3KLngUH1/eftXR"
    "XChRJLAfhG2AjxD8MoNcraFF5Nw5oqq6iXNgWe9STX3jL1dhgAUXg00arloR1XQzpNzlabNrz258Wnh2486znc/WZxuGWlAzk1ei"
    "UzimJhRpZE8jmGI8sVxVEBexEU7rqSfUsnrCancqqqEu47PH8LGKjw3+qODj9W4AL/01e11PCF4CcignQG9DMhslW7fYa3R7lJcs"
    "6w6yTto43TYhuF1mkCJoS4ZaVBE9ZAUEbwQBhYp1n2oQzwXEK2RtHTjvhvBUKFVmLpn1ILxg2U1NCw1XR3Q9EppXXTyoCn+C4+JH"
    "oYUmHg6OoQdJR/ENSBKcsfzteGItNLFlZoRm22pBgh2atsgF4GmD+k6AQxCV8YtRDz67kdWg8O0HDL/QZ/JV7U4EODqavVK1TV83"
    "3wlcX1OJqq+njzEDPEYDC8AXdy9kBdgBDSTnxFeZxBMnEeY62ybvvUdKsJ70AcdxTcUcGxUKkrJVUdACumvHe/ycGtSnIThGp0+G"
    "H9wf3vyO/PS9jBHkeO+S6VG/wZr9/e8eXyOniMbXREE3BHrnyDUElojEYP8aMMQVcKbe9XnQINh5vUw9DURCuNcz0F9Geo9Qz3R9"
    "oP/Xb71+EXkacadK9XiPL+gnDhSZwMVcfBt4uvEigAVt1zb0DHAsWRWFMRPzASLhWaGqG0QdPrg5fPQYtGnNNE0f9OsyZdoljj5E"
    "9KG5DaC6vAPTXg4Cj1q+DieFUUzTLKOmr1RrBQsPLEtB5olHJCI1aSodxC5NiQUtMCLCN8/1ckMHvfRcpoFh6hwpN7LIBCtrw8pR"
    "nJX0aJqehmphEAbHAfqGVobIUZI4AXJz69oxFkk90ElIWTf0SYTaCZamSZ5IUOc4BLjOr+xdv0srPLAIoJCbGOz+Em1c2OpoKuj5"
    "KcJi5iXv3AnZGfdknjx17m+P9/qa/t7alfUrV7jbu3Ll+AmwEXm+V65EJ1UdkKk6CFltuKouCONeogRXSA2itQ0CfgbKC+4bZteW"
    "T6zPgm1ETGvroNJt0N9ry9iIBYVq9zG0wOM1jq+PFpnsv5+K0ep0vG2NS1Bs1XWQMp4U14VXHdgcV2rQBr+dm8PSLjMr/E0yC+VT"
    "MllJ0G+jt0vVWMyDFwzkcKJ78UynKWd4TgjjPFWkToqyCwCnAIJnypmFkFSLpZjCx+Pgg8ANq6kLOjUKcJJ0cbG1FS+2tp6zOAMA"
    "i1PGalgASSy8GIrBEESGgEuxinNr6vGjB0WEU5B6S06cIMfwNQraVGPC5iASgPuxva5DI43peqLcdXAwUnMBC5xXFgu+5rBgtDgK"
    "HjjZLB58zeHBgQwa9OITsMvho7EOegKkZLSVnowcA/Fvh1NWgArBisQfcWA7mALcETuB+NOcBoHqc0wGLOCca4R813uyNBDn3IFz"
    "lGeN+sHDHipJEgKTQMd9UieOg5Moj9MGPjtkmavzyAIBmrIlQKuc1Um4+/xTjrNQ6GE/Y5xSWdFtZ1VVTNpt3GfPdcoycojESLMw"
    "FTG9AH/EtAr5gRVSTU7VcIqfO0D5QaEGH+hGQDFiJBZPWUaW1/igobbA3g3ONBZA8RKYxtfLwOV776HvzS+28rMGEaVQZrU4lXPn"
    "CiW9AOCZVwl8PgFOZ1/16/h7rW1YUhsf7K9F67FVi7AH4hJ5EKRXYN5gwdx/6uh9hUA3A/RPPg+GmdQI8mJQGJ9n0bgvaQMgjMJr"
    "l0SogIJczYTCURKonXCMUH6AcPJZSg3zEuRRJkdQ41QhQSKzRAT60cQJd4BxMbMGIiUghyJmBLmarVL4tQ2Ue1/uAZbh3U/J4Osv"
    "0hJ+f+8GGd79aPjJk8GtncGtfzZFAaNWpHJWuKoCXZE2r8SJs7S755PHppVSxRjNGcfgDDl1SxdZV7ZW4wUqKIxVaK0ox3utfr7a"
    "6jQxk0M3gTnjsttuEC+wsHu1onjWu9sKiUIbF/JE0m1D0hytIfh6XyGWxzAT5Hll3MmAXd+/rfaTqnC0HWXX8gzYuJHjPch2MOU3"
    "MDzrfeTI6XYu8sYH5rLjLeqqJpLKFFDvxw1PkehKHvKVazshxi2Qq9MkQKy/YwJYWWQLuZNXfPjAVAdrgX62RBWfmeQWtTk+fHRG"
    "kFdHlBz1mBn2G/htLnzxO4BlFuJjVYh7eRYe8TXfipGDma5LDMa7x8lr0kTW9p98q6fr4jZMuo43TWQXMgUU7WbxOouMzcZM8osK"
    "UE+pAJOUlIWjasmcUVWMxYAXscoLKibqAXMk9v8fNcujzykWDMdGwCGOqFEJRgHPi1rBbxz9kd2Et7hE6evjPEnuIWSkiOMt+qjk"
    "4KbjAnYCaV4nx+N4yGPKLQ4bvoWeZpQ9m4cHHeq/jt0erSWScd4WaFXiAgReUEvWWuuirkmyoDRA17HUxriFClTjSuIwYNRCE2Rg"
    "vzHXNW6UjjygiogWjtcYixbZtlMTr8rmkn1DYO+LElq888YDoAWQw3uC7TdDuqFUn929MbEb134DL2Cqz37zp8nTq14QYSv01r/u"
    "P90jz3bvjPYKJzaiHelVc06s5jbGnPv/xpvnKAMlD5DXNTXre6DEE6LjM9IBZWYYbefn/vAhnxDSFei4j+KjaCZyNHFVfCKxh2QG"
    "HBifSSxBEIk9GZ8TNpCZAGcWkxrt36iYo2QtAzssooVTJSVhfeIyhqCEYTy2PT3mC/D/8cHw1i4ZfvD3Bx/sSvZ8J8/bo3t8QjSp"
    "RqfI8J9+HH50m0Ng9ypGfXdvcP8/4qUMfZIY5/6YDyZm6/BGNvKInS6uB835KuL9HLKVm98NPvmKDPZ+SQ4++WH4NSgbTMLJ5i7H"
    "Jv3YQMmYm0AtbY5TlF5IEhZNNVlfxQyIPOnJQ2ThfcigvpWkjxT8YP+PdtIwJK5NMnFOnEwa0G79frg7PU6lPPIQZcveJA9RUg7y"
    "10YOrVtdj8Wxyj1VGnWY2DMcH7s+Yaw7xan2J3vUfE4jhqRn4010bByg3C5iqAVdgSqnGWyKZgvCcL8CQAGUpq7dAvdn4wh3yRUO"
    "gX4rB6BxMaSOW0OfXQDtl1ln3Hf6We5VIEMv93xkp0hpbHk2aKQsaiIrH99qSNvBBk12C0EnSatAChfArhjCYSsVKgZkBuyDit5z"
    "3CBbIdRkVtigzOQUI6apayI5WVdFnw2q+5T3Uwx9iRUBfCvTy5W8PZcsoIvp8dw/XavndiwwJ5cX43hbdNsJNv0Es6xlJsoJu36W"
    "60expLLhlTPUwjsArEQuRLbVAVUZ4YUnrnnI8yFEv4u0DsJ+ATWZhvAS/tZoHONkXYH/a7yJZ8h2nSEac4boeRmie7We3DyIC6Hj"
    "ePUzQaD8xhbEyZuGusDN+3hG2rgz4k6dIRtyRtz8MkRj4ajE7KblN2iWGhwcv1GcYDi9/A0LN7IYdsTuA18VRiguHCdai4CRvVHQ"
    "M0E7hp9OnHskdRr2EfJj7B1KPb7VnEj/zztjKfWkqclvPcaasny8Mqkpm5kZacrCjGxNjXO/Gk1yn2njGUMO3ompw51dCF/Adz5z"
    "GkmXshlSkhRl86Bs5pPNdFQR/ZJBkeVow99+quPkWFoiwcjB57cBr6A3noFkk44k0VCn5RGc0GiEV9fTXP56fE2iKtjFkD2gjbEb"
    "UAWLIkURtYaipgjAUUSyZkgva9YwFTF4hZdc6fFy7JDrQcwKDZ7BGTLXM0SedtSrw5HUYaV6TWQGJiQ/kA70+df1Pn52+9eSmx95"
    "PQKmsJ7pUNrRBmrff//7f6JY1lBnDGKa+CdB0TonYuFGLf54PcZiZO5bQkSZIgTw9P7bDinsVl6Ba6rFzdBshhRLq19cuigBfs5/"
    "2wzvGl4JvewFNW0NGFs3SA9/zVJW8f5pFkYq8V+J8j8SVfviZtQyMSphtY5bibOzbwZf3r8qfrAgfolgAgKwsoQ5THpMMCuQ+WrT"
    "9RzN4txxc0JDAzpvuW0KsV2L3QKyDFlL0MqwLPbDl0p/g0aKvxAsFjFiVGZi08UfWMofFECRJX5aOSv+NvZ/AOAf5ho="
)

# ------------------------------------------------------------------ 내장 양식 (ledger.docx, base64)
TEMPLATE_B64 = (
    "UEsDBBQABgAIAAAAIQBqtrmFigEAAFwGAAATAAgCW0NvbnRlbnRfVHlwZXNdLnhtbCCiBAIooAACAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAC0lUlPwzAQhe9I/IfIV9S45YAQatoDyxEqUSSurjNJ"
    "XbzJnm7/nkmXCKG0QbS9REpm3ntfJvKkP1wZnSwgROVsxnpplyVgpcuVLTP2MX7p3LMkorC50M5CxtYQ2XBwfdUfrz3EhNQ2ZmyK"
    "6B84j3IKRsTUebBUKVwwAuk2lNwL+SVK4Lfd7h2XziJY7GDlwQb9JyjEXGPyvKLHW5KZL1nyuO2rojKmTKWfeSgZb5QE0PGXRniv"
    "lRRIdb6w+S+wzg4qJeWmJ06VjzfUcCChqhwO2OneaJpB5ZCMRMBXYaiLL13Iee7k3JAyPW7TwOmKQkmo9ZWbD05CjPSZjE7rihHK"
    "7vmbOOQ8ojOfRnOFYEbB+dg7Gac2rfwgoIJ6hgdnYedmAoHozz+M2roVIuJaQzw/wda3PR4QSXAJgJ1zK8ISJu8Xo/hh3gpSUO5Y"
    "TDScH6O2boVA2l6wvZ5+KjY2xyKpc3MAaRuGf7z2fndV6o7/08mrE8n65PeDai3mkDdk882/YfANAAD//wMAUEsDBBQABgAIAAAA"
    "IQAekRq37wAAAE4CAAALAAgCX3JlbHMvLnJlbHMgogQCKKAAAgAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAArJLBasMwDEDvg/2D0b1R2sEYo04vY9DbGNkHCFtJTBPb2GrX/v082NgCXelhR8vS05PQ"
    "enOcRnXglF3wGpZVDYq9Cdb5XsNb+7x4AJWFvKUxeNZw4gyb5vZm/cojSSnKg4tZFYrPGgaR+IiYzcAT5SpE9uWnC2kiKc/UYySz"
    "o55xVdf3mH4zoJkx1dZqSFt7B6o9Rb6GHbrOGX4KZj+xlzMtkI/C3rJdxFTqk7gyjWop9SwabDAvJZyRYqwKGvC80ep6o7+nxYmF"
    "LAmhCYkv+3xmXBJa/ueK5hk/Nu8hWbRf4W8bnF1B8wEAAP//AwBQSwMEFAAGAAgAAAAhAMyXg5XMFQAAJugAABEAAAB3b3JkL2Rv"
    "Y3VtZW50LnhtbOxdW4/cRnZ+D5D/QMyT8iBP3Ys1WGnBq9eInQi2k33mdHM0Hfc0G+weXRIEUGw5UCwDaztWrPVKjoxkL15oAUWW"
    "1zai/Jk8Tvf8h60qkj096hvZQ3azZ1aAhuwi61SdU+ec71SxLj/56a2DtnEjjHutqHNlC74Gtoyw04iarc71K1t/965/2dwyev2g"
    "0wzaUSe8snU77G399Opf/sVPbu40o8bhQdjpG5JEp7dzs9u4srXf73d3trd7jf3wIOi9dtBqxFEv2uu/1ogOtqO9vVYj3L4Zxc1t"
    "BCDQd904aoS9nizPCTo3gt5WSq5xKx+1ZhzclJkVQbLd2A/ifnjrhAYsTIRui21zkhBagpDkEMFJUrgwKbatajVBiCxFSNZqghJd"
    "jtIU5thylNAkJb4cJTxJyVyO0oQ6HUwqeNQNO/LhXhQfBH35M76+fRDE7x12L0vC3aDf2m21W/3bkiZgGZmg1XlviRrJXCMKB7hZ"
    "mALfPoiaYRs3MyrRla3DuLOT5r88yq+qvpPkTy+jHGE7X7GyOLEd3uq3e/0sb5xHdkl2N3UsWmrbcdiWcow6vf1Wd+QdDpalJh/u"
    "Z0RuzBPAjYN29t7NLsxparNcm5s0wwnBPNVP2+6gndR8PkUIcrSmIjHKkacKp8vManIgNfik4KVEMyZcmNP5ZATQBAHWCHOCRUbD"
    "TGlsN06sW9Fp5TSrjE7SKopO60SwMKcPfLUyYwSah4VIIJzVQ11U9jFavWa/uV+MXNZG2ypv0A/2g97IaBKKezkdQUaRjFFMFKwd"
    "NUb+TNEMiwmNjgjePhhrw+71sxnq63F02D2h1jobtTdOXPZNFTwVoJUa/LgT6p2tMu/sB13pyQ8aO29c70RxsNuWNZLma0gLNHQL"
    "qL9SkdVF34a3dLrSn/Rmr61umoeGcolbV2UQuBs1b6trf7edXq7F6c3PjZvKy8iIUv683ZWlBYf9aGs7ffxmcDs67I+e7bVuhc3R"
    "Qydst98KEkpRNyEE6Qml5q0gebcd7ika8jEbK2j0eDfq96ODOfnj1vX92QS2Jyoj6x1F78n3bgTSBACxdKa9Vtzrvx2pQtTPdpD+"
    "OnnoRO3DAxVhZ8+zBP1KJ/qZLWPs0a+/T37BkzqMpPp63Gqq2+vyKmmkrBEAEn5OJ2NA2QmNLGs/lo9lxN98WxYITIQ8yGW5kOx0"
    "gzh4QxaMoW+7wHaTVBkC9VUqT/+pOqrs78YqP7WoTy2tDP04rWb8szCVqxYThGLESvZKI/mb/UqVRTMy2Uo3rHbreicj15BoHsYZ"
    "wZRE9zQDFrYJ95W05zBwmv8kyQ33gsN2f/LJtVeY7Sbl9rpBQ1qXfCnYk7VKm7DdUvaOyOjH24fK3DID2B5lT/4k953oWhxFe8nz"
    "NC0NAuRtd6fVUZSMZqvXf1eXo+7s0d2bozvF15a20Z2g09iPYsU8Y57rYttPH4TNlpYJcTyIudA8dXekoFSPTnW+TMoRlWQat2Ve"
    "CilMNEy+s7cXNvpe8qYygi1DiWvLSJjfVX+TN2UP8VpsKFDF0mMTYRK8ZXSCAymJo+9/GPzmpaHVQr7Z+Jsbr8dBd7/V8GP5XPEe"
    "SF0+SXlTwkUvC3+XiJ6SmKUTOftB53po9bqShZGBzS//rKWOkXIlmBqH8SSsLCbVbTX6h3Eoqcm7ne6oWvLuzNQ6N661tBmpH1IU"
    "CxrtJH37JIsioIQ5QW+33er6rXbbiKP+z1v9fY1DSvRKLuqhEe+EB7uhLFDqqXZXwU4vbrwtW0ipFyKUmlrFIGUCaTVDDCOmVU1w"
    "xtMs/TjsN/bV7Z4sTmVX9Rl7oOuWVSepXE+aobF78y3ZM0qtU+W/tRcfqKtEVOOWVurbqVIHykJS8xAIocw8MABC69JJ5q70+q+H"
    "EnzUjeRN1kcTD2682Utrlr2iktsd9bcTqbolT5MUWZ58PytZ684/UcwIEIhdtiyXXybENS/btrxzHE8QDBmhnvPPqjBp5739oBnd"
    "/NvdXkO2dzPT5bxR+1j/ESTanNRN1iW76tpp0Sphprfyv34+pvfjvxOjS/yZ9nYjN6c8X+IfU+8+Dyw0vJWAFsAn2LIcq6ZocXPn"
    "HxpTuckwJKvNNQ3IpocQ9nSpKYjsaiK9f8yIJEGBSnF6WRoyU6IZQF8d3n9kHH/x0Bh+9Wtj8ONd4+jFHWPwm6fG4OM7Kk2Lc2qD"
    "aZjPgg5NazI4fKd/ux1mZQd72nj0gz9HjauLGtWglObnVDKl2qO++jJl2t9PJtMpyRIdpqSa016FAM4gXCSZi2nJyIS6HkXjXwg4"
    "xi508sS/NmTY8bW5zYp/8Ukt5oe/jHM8Rd0UT+90g5FLS9vnxlthfH1kRHHY6wdxPzWGhi07f2HcO7EF/VKn1T5lAacTR3p/OjnT"
    "9vFUbepjZRT2uYiY2Kausq6Kfe4MJyjBO5FEoP1xKK3L6rVklPdedPmv3x53htOj9Ql6qcYvpidFcnXw9OXxZ/cGv//Q2DGGXzw9"
    "+vbl0Q/PLg2+fXF8/87xB58eP7h7fPfZ4OEng48+Tx4P//e7v5rtcmdqlFSoPBqV1r3+asNsD3NuedWrzQnmakZSnuqoSxKYJS4P"
    "nt87fviDcQmQyxKdJWBfBnSOxqQgnccbUhcjTFyyCd7wDGpcjZvzfQeYYKX6elKtEhW2Pu6mknbiyHFdTO3z5VdiP+r0VYi/3+rI"
    "Gmb5U7d9JqeTAph2OgRCQAQklxEwGQUIleJ2AACOiRaZTvluR4eYi3U4bZo5QVhxcLMQtDBE63EWOaFJ8z2uCccPHwzvPzYGz35x"
    "qtnPoom51V6Wp4Y0dL9aNlZXyj+Mb4RbV405dVnM0tH3j46efzNbi4vqTplKgjgXHhVqHHeDlGT45JEx/OpT4/jLj43h+0/P0Dq7"
    "sX6hf/WSpvn1M+Pouz8sExqvoK2AxSyXuap/vkFtNfjFXdlWH5bRRoMPJalffW4MH79cooUknOiRjwWeOO35l9lw2Ld9IeAKvh+V"
    "3HCy3Y6e3ZHOy7g0/NUny5iFBHJRQOiFJUsYABCpIjbPJIzh84fL6LEahqrW0zAEBfHNFQynlIoKX36jJCtDB+P4ox+PfngmoeH4"
    "wR9mi1hdkrzxVJbGfa/jQeDjNLXe4eO8luWWR33Z/a6oZbOii3O0uOoM2dil2NzAqiMTcAc7VSF3jqovA37zOCKO5zJgVdW5OENj"
    "FPV1RFgQMrbSvnoJ3aQk8jUGvz8dWeWDkNJkh6U3Qbwqk6xIdsP/emkM//XjZeTGxTQrKo6vXGBbINUGG9Xrumsc//JzY/D8nnH8"
    "8Icl5Dc3dJnrQAVGDLoLPqGV4G40Ius6x1OJjtfJEdghatZGOSGBOZpcV9cRJQwhpP6iiXyVDz9XMn6Z8+NIjYdyABJIMLrgw8p5"
    "bp3hk0fDr1XUP6WBJgMY4dqenjpxQaWlRrykuOo73kUti7kuvMDuZvj97wb//nj4+OWl418+GD55sExDFYz8C4eAHne57a9pBKYO"
    "jTR4/8Xgj3eOfrg3fPD1ssNkK7Am7APhIbNWMWcZ0tejlOpLgPT75faGShQ+FwQLKagLbCW6nYbPHxqy+1Vu76vUCNdksqNRq/Hk"
    "MnDki2eDx/9X2+FmgDhgGKT9q/MjdT0ePVvk6pK8Gk9lcVxCmNrI9RYMyNejrznXvmzXhHQF/fiiHOUYfrYdm1lWVeOEVVadU8eU"
    "NlbVioQcVa84CIUUYh+tYoFeDgciDf+rD4++faJWPAw/+JfjDx4tg3XLtjUW1LbtykYBqlRT4COPOHZVEWqeqs+LMOZamEcg5l5V"
    "EzPzWNjSQ6sc2MI2q/8iWADusMU5grisyXooWdW3erijtu8Qa9HKj3pihmlj3zE3Ee6Q7TkMWuv8xg1YnlkvmR4VnrTH1MIFUI/p"
    "Gf2r//8fT4zh/SfHn91bBuiAnqF65k9qMgCQpPR6/3rI5Ovhf38qo4AlpbKs7kPfAUDYm+hxECUI07VOKVga/k0IKDHXOb9jafin"
    "JrPsFShMAfiXEYnnM71WthT4N4lebr8Y/4usEZopUcAJlt2dFSzSSZnZ1ZU8y2qx7I2byeyI4ZefG8O7/zP46NfZZLrhbx9M8WIT"
    "fQ8obGCBFUylKTyemc/tFltwVOo8fOy6QAYup0WHCHUZwEqgeURHGLFcvYfBSHRp0pjo5g4Qn3HUa7qGlTGU9v7T4W/vGuOrEWej"
    "agFPQ7mJLAbKWoxIkR7BKNfRLKFNzPQxx6R283dPNaluyKPvnipP8/6L4RdzJuqWarDzBEc9YUHgVD/HtICOAmQ5DndybZ9Vpo7q"
    "DR0mZVv8C7GFZO3t2k2wPKWMXzxT09qW0EDM9DT5qkY2OQEOZbx2i/gmLXm5b69qa48SdAya1KVCrGmqyCtikhK599AY/Oe/LSEP"
    "CISuQUUjBxQQLHuHVUX7ReU0uP+74aM5X8hmiqlsv08c23EhqNcgqI3ViPgCLKqv38eCCwexFeyDlUvXYP2cu2xeghhaMzKu1FFz"
    "0zV901/zMF0tfDFmLgZuZR+USxVF6XG2cADClc0aXMrfEoyJ59hljTqt3N9S35OdPq8eE/z6V+ds2zFTFBX7W4w96EB3zYi0Wn/L"
    "XGbBdXe+auFvEQbcpFzZ2oXztxwJ7qhZX3WKby0HmP6iPfvq628ZYMJyK/vok1OfRv4W18/fUoihx92LFN8CJEzLBWteEVsLf4sZ"
    "Rja01jzWsB5/S6HLXLuyCYlL+Vtm+55r0zR18/wtMKll2Yvi85X5W1I/fwsFtwAWa9oeci3+lpue4HTdCzZq4W+Rjz2CVrEHcv38"
    "rYRdZrte9RNwCvhbSqFJkLWgt1lff4tsDAV16zKpfs4H/5miqNjfIpO4yK7sq0GpFleSv8WYC8LhmgeZauFviWeZDPsL4qnz6W+R"
    "C3wPoKrWs074W3VJjv14JeLxHEQtUbrfT5PG2qA3OmtlN9yLYinF5ICVM5/TtZs8GX1NvzR88NXgmxfTVpFPtAKEAkCv9MkB87if"
    "wu8r7Ia3goaeBjed37EdhtKXMkHEIwXtBtfD7OEs5ilnwvdg6UNak8z3g91ees2qnmy2LX91o57UQkJYxkz28qS+6JGiTH4nQpuj"
    "MCPRrOMkAH2ATNhphnHYvCbbw47D4D2du4zzJsqrpxR4WqtiJxeMPMqfDxJK2K3NQUKqOzeCp1EyBnT6qZRo6gFD2Jx6sg9JpoJP"
    "0J56whDEdOoZQ2lywkzGQ44Bb4f7DuK5NgwbCwJmdgiSMxQ1QmZwn/x9NQ4iVAdsi9Bfn9m1RCAkuzmUV7cyWhc4OxAq30mmAWHO"
    "mcn3HupjzLSzMwZ/vCMdo3H03SfD5y+M4eMXgyeLjzbLozuMe4iKRas0c+sOQSMNXqA7ydjKRL+hzN0fHIqpt3GbSD5+aQy/+mR2"
    "484UqXZk1YqUU2o7bNHUsbqJdPC+NJkvl9ldD3E4TaSlzZp0PEu41pp7ekXFefT8mdTQaXvhLRKnRslqNRT5lucTfX7wBTH66eMv"
    "JYqUUV92BvmGibS2Rk+443misgXwFYlzodGrS/JuPJWF8XgVutRbuN3rimOOubAHhVQKtMaF/3OxfV7VoQUcx6Pr3PtEdaSmVL34"
    "FzlKgAvrYTdqicGLwf0fl3PauhNZgkiIbXKHmvXYVa5/9fjjO3NPXJ4tj3lxwTwBqIkgPgVrXNo/F3/nqjOAEJp4nfspqeGGErSQ"
    "WjKErW4P9dUaZjkiAYJ47sJTXWpjmOqSvBpPrdp4Y7u+LexFR07lxm6Oc++gNB27C6MIp5bH6JrXPebTxqmgP5szSpDjVXbi3gqR"
    "Hnq+a7lkM74FlwTlQC0RIc5GLCGYBdeFeeYeYK675o5lXlucgvMzOSPcw7aJ17lTWFlIxqDvMGcz1nmUxbMje02isk0Xl+G5AEJj"
    "07PVbM35ta8vQkOTmy7ga54eVgVCM9eD2KnMx68QoTnxGCD+RnjushAay8jR9DZjWlZZCI1MZHN7M3guhtCIY5/ZeK22WA5aEQ4g"
    "MzclWi6HZ+gAYRFUp2i5CELLmjti0V4F9UVogpjHgX0OERpBiWuyF71Or1BSHxoARl1nIxYll4XQ1CcmcsSF6kMzz8OeiTdicU4x"
    "hMYmBdhk6zyWuSS04qYALgQbEi2XwzNjNkRg3SfRLovQAPiWx61c5wTUEaGZKaCLyTlEaOwyRKrbjn2FCM1dRjjV+/lfGISG0BfE"
    "qmxctFyey+pDYwvayfbJ5wyhgWcKDt11bqxfVh9aeIDgdc8tWS1CY5vbmK/7cLxlERrZjmuzRfM864vQmAhm483w/sUQGhLmCtta"
    "68SWchCa2L5p+mAjNpopDaGFBGiyGVvZlYXQnMmuJtcrJ88ZQjOT+dhzq98Oc45elvRN1nMs30MbMp5V1ii3SRyA66SXBRCaQWp6"
    "RJ8etJEIDVwV3nopV+dqlNuHJnbEeUBo3zVdwjZiY6rSJn07snbQ3IzRxJIQGvs+t9x1HXNfJUJD20EeX+vB1WWhlU+o5ZHNGNsp"
    "a9zAx47p8zr14gogNLek+8R4Y/vQyMLIccg5RGguXGyCtZ40WxZCQ9/hpt4v68IgtDp8AQhzM6KSsvrQVFjUrOxE6pJtsdB3aBdj"
    "5K7gSIk5elnSN1lfVjg5QfrCIDQlCDCnVmsMCiA0MF1PmKSsnX1W/x3asoBv2hsxmljwO7TDkccqmwO3QoSGvtQylmnexUBoDIHw"
    "GalqBKRcnsuaKcawoASfQ4QmQKovrGy17QoRmlMCkWldqLnc1HeAz2mdxg0KIDT0Kaf2otrXF6GpZ2JXBhmb4RWKIDTAHkVuZYfE"
    "rBChgQWdTTnupsSZYpwBsBEzGMtCaGohy2abccxGwT40thyLLDoiZhMQGkAkmMkuFEIjRzDq6J2Ba47Q6XEXp63KdR2XLjoMt8aj"
    "3AgjQHMdTzRif0L6p5/UZZSbQOhwkWeFxzKcrRChmUBYglWemRprb6OyEJp7DiPWol38asJzSQgNTeShfJt21cAWi41yOwKIXCu9"
    "K7PFkka5iQDAF3nGOWpgiyVFJdTGjsvyjHKviucCCE1saHMsFiBcjfvQBApokjyj3BuG0Aw61Mo3l7vmCM05ZZab6xC7c4PQFFNu"
    "O+6FQmjGfUo4rJMnnGOLxfYUsxlxcs1SrzlCYy4wpvpQmAuD0IhyF5qU1YjnFKHVZdoRaJxS08ZipYeAVX0E2uCz73MegUZ8x/Hh"
    "qzP7qIxSfKBnTp9w73qEUdWJTZKu6ThltkhOvz4uEq1ZvbDRvxYXpP+OzKRSASYMskS2199R27krFWYm1tn31bCqSA7xjOKW1I4r"
    "W+2g0+w1gq4+Fk3leitQhfejrpQBTRhVwZXKqvs96WlRam8oDVLqbKmTh/th0By12V4UnTTg9cN++kMX1IjaPSNpetmMavGNTm5G"
    "DXXuUdrG11r9hqw0ZtmBRIlw9O1u1Lytb2SWwwPJy9U/AQAA//8DAFBLAwQUAAYACAAAACEAVjVqJy8BAADIBAAAHAAIAXdvcmQv"
    "X3JlbHMvZG9jdW1lbnQueG1sLnJlbHMgogQBKKAAAQAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACslEFPhDAQhe8m/gfS"
    "uxRWXY1Z2Isx2ati4rULA1RpS9pB5d87kSywcZd46HFeM+99mU672X6rJvgE66TRCYvDiAWgc1NIXSXsNXu6umeBQ6EL0RgNCevB"
    "sW16ebF5hkYgNblati4gF+0SViO2D5y7vAYlXGha0HRSGqsEUmkr3or8Q1TAV1G05nbuwdIjz2BXJMzuCsrP+hb+423KUubwaPJO"
    "gcYTERypF8hQ2AowYb/lIMYhmTF+muHaJ4PDvqEhjhBDvRR/5zO+NBozsW9mYxilJYiVTwjdqT1Y2rEJYpSWIGKfEHnn0Kg3Shsh"
    "wnBSuURQi2ux9kkjFT2LiURBIcUgxuF7W51juPXJ8AX7F0CkS5jt50xcGsaN1zfyh+KgHBD40f+T/gAAAP//AwBQSwMECgAAAAAA"
    "AAAhAIWpIi+BQAAAgUAAABUAAAB3b3JkL21lZGlhL2ltYWdlMS5qcGf/2P/hABhFeGlmAABJSSoACAAAAAAAAAAAAAAA/+wAEUR1"
    "Y2t5AAEABAAAAB4AAP/hA/hodHRwOi8vbnMuYWRvYmUuY29tL3hhcC8xLjAvADw/eHBhY2tldCBiZWdpbj0i77u/IiBpZD0iVzVN"
    "ME1wQ2VoaUh6cmVTek5UY3prYzlkIj8+IDx4OnhtcG1ldGEgeG1sbnM6eD0iYWRvYmU6bnM6bWV0YS8iIHg6eG1wdGs9IkFkb2Jl"
    "IFhNUCBDb3JlIDYuMC1jMDA2IDc5LjE2NDY0OCwgMjAyMS8wMS8xMi0xNTo1MjoyOSAgICAgICAgIj4gPHJkZjpSREYgeG1sbnM6"
    "cmRmPSJodHRwOi8vd3d3LnczLm9yZy8xOTk5LzAyLzIyLXJkZi1zeW50YXgtbnMjIj4gPHJkZjpEZXNjcmlwdGlvbiByZGY6YWJv"
    "dXQ9IiIgeG1sbnM6eG1wTU09Imh0dHA6Ly9ucy5hZG9iZS5jb20veGFwLzEuMC9tbS8iIHhtbG5zOnN0UmVmPSJodHRwOi8vbnMu"
    "YWRvYmUuY29tL3hhcC8xLjAvc1R5cGUvUmVzb3VyY2VSZWYjIiB4bWxuczp4bXA9Imh0dHA6Ly9ucy5hZG9iZS5jb20veGFwLzEu"
    "MC8iIHhtbG5zOmRjPSJodHRwOi8vcHVybC5vcmcvZGMvZWxlbWVudHMvMS4xLyIgeG1wTU06T3JpZ2luYWxEb2N1bWVudElEPSJ1"
    "dWlkOjVEMjA4OTI0OTNCRkRCMTE5MTRBODU5MEQzMTUwOEM4IiB4bXBNTTpEb2N1bWVudElEPSJ4bXAuZGlkOkI3QjQ5N0VBNzMx"
    "RjExRUJBOEIzRkRENTlBODdDOEJEIiB4bXBNTTpJbnN0YW5jZUlEPSJ4bXAuaWlkOkI3QjQ5N0U5NzMxRjExRUJBOEIzRkRENTlB"
    "ODdDOEJEIiB4bXA6Q3JlYXRvclRvb2w9IkFkb2JlIElsbHVzdHJhdG9yIENTNiAoTWFjaW50b3NoKSI+IDx4bXBNTTpEZXJpdmVk"
    "RnJvbSBzdFJlZjppbnN0YW5jZUlEPSJ4bXAuaWlkOjBBODAxMTc0MDcyMDY4MTE4M0QxRDMzRUNDMTczRjA1IiBzdFJlZjpkb2N1"
    "bWVudElEPSJ4bXAuZGlkOjBBODAxMTc0MDcyMDY4MTE4M0QxRDMzRUNDMTczRjA1Ii8+IDxkYzp0aXRsZT4gPHJkZjpBbHQ+IDxy"
    "ZGY6bGkgeG1sOmxhbmc9IngtZGVmYXVsdCI+TG9nbyBEYXRhPC9yZGY6bGk+IDwvcmRmOkFsdD4gPC9kYzp0aXRsZT4gPC9yZGY6"
    "RGVzY3JpcHRpb24+IDwvcmRmOlJERj4gPC94OnhtcG1ldGE+IDw/eHBhY2tldCBlbmQ9InIiPz7/7QBIUGhvdG9zaG9wIDMuMAA4"
    "QklNBAQAAAAAAA8cAVoAAxslRxwCAAACAAIAOEJJTQQlAAAAAAAQ/OEfici3yXgvNGI0B1h36//uAA5BZG9iZQBkwAAAAAH/2wCE"
    "ABALCwsMCxAMDBAXDw0PFxsUEBAUGx8XFxcXFx8eFxoaGhoXHh4jJSclIx4vLzMzLy9AQEBAQEBAQEBAQEBAQEABEQ8PERMRFRIS"
    "FRQRFBEUGhQWFhQaJhoaHBoaJjAjHh4eHiMwKy4nJycuKzU1MDA1NUBAP0BAQEBAQEBAQEBAQP/AABEIAk8EnQMBIgACEQEDEQH/"
    "xAClAAEAAwEBAQEAAAAAAAAAAAAABQYHBAMBAgEBAAMBAQAAAAAAAAAAAAAAAAIDBAEFEAEAAQICBAYOBwcEAgMAAAAAAgEDBAUR"
    "0ZIGITESFFQWQWGBkbEicrLSE1NzkzVRcaFSYnQVMkKCI0OjNMHCMyTwouFjBxEBAAIAAggFBAMBAQEAAAAAAAECEQMhMVGRElIT"
    "BEFhcTIUgaEiM7FictFCgv/aAAwDAQACEQMRAD8A0AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAELvTdu2svtytTlblW9GlaxrWNdHJn9Dlpw"
    "iZ2I3tw1m2vBNDPOfY3pF3blrOfY3pF3blrVdeNjP8uOWd7QxnnPsb0i7ty1nPsb0i7ty1nXjYfLjlne0MZ5z7G9Iu7ctZz7G9Iu"
    "7ctZ142Hy45Z3tDGec+xvSLu3LWc+xvSLu3LWdeNh8uOWd7QxnnPsb0i7ty1nPsb0i7ty1nXjYfLjlne0MZ5z7G9Iu7ctZz7G9Iu"
    "7ctZ142Hy45Z3tDGec+xvSLu3LWc+xvSLu3LWdeNh8uOWd7QxnnPsb0i7ty1nPsb0i7ty1nXjYfLjlne0MZ5z7G9Iu7ctZz7G9Iu"
    "7ctZ142Hy45Z3tDGec+xvSLu3LWvmCrWWDsSlXTWtuFa1rx1ryaJ0zOLHRhgsys6MyZiIwwe4OTNJShl2JlGtYypblWlaV0VpXQn"
    "M4RitmcImdjrGec+xvSLu3LWc+xvSLu3LWp68bGb5ccs72hjPOfY3pF3blrOfY3pF3blrOvGw+XHLO9oYzzn2N6Rd25azn2N6Rd2"
    "5azrxsPlxyzvaGM859jekXduWs59jekXduWs68bD5ccs72hjPOfY3pF3blrOfY3pF3blrOvGw+XHLO9oYzzn2N6Rd25azn2N6Rd2"
    "5azrxsPlxyzvaGM859jekXduWs59jekXduWs68bD5ccs72hjPOfY3pF3blrOfY3pF3blrOvGw+XHLO9oYzzn2N6Rd25azn2N6Rd2"
    "5azrxsPlxyzvaGM859jekXduWta92bt27l1ZXZynL1kqaZVrWujRH6UqZsWnDBPLz4vbhwwTACxeChY3G4yOMvxjfuUpS5OlKUnL"
    "RSnKr23jz7G9Iu7ctanrRsZZ7uInDhne0MZ5z7G9Iu7ctZz7G9Iu7ctZ142Hy45Z3tDGec+xvSLu3LWc+xvSLu3LWdeNh8uOWd7Q"
    "xnnPsb0i7ty1nPsb0i7ty1nXjYfLjlne0MZ5z7G9Iu7ctZz7G9Iu7ctZ142Hy45Z3tDGec+xvSLu3LWc+xvSLu3LWdeNh8uOWd7Q"
    "xnnPsb0i7ty1nPsb0i7ty1nXjYfLjlne0MZ5z7G9Iu7ctZz7G9Iu7ctZ142Hy45Z3tDGec+xvSLu3LWc+xvSLu3LWdeNh8uOWd7Q"
    "xnnPsb0i7ty1nPsb0i7ty1nXjYfLjlne0MBc1AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAACD3t+XW/fR8yacQe9vy6376PmTQzPZPorzv129FRAZHmgAAAAAAAAAAAAAAADQ8D/hYf3UPNozxoeB/wsP7"
    "qHm0X5GuWrtNdvSHu483+WYr3UvA7HHm/wAsxXupeBdbVPo1X9tvSVBAYnlgAAAAAAAAAAAAAAAC4bq/LK+9l4IqeuG6vyyvvZeC"
    "K3J9/wBF/a/s+kpoBpb2eY7/ADcR72fnVeD3x3+biPez86rwYp1y8q2ufUAccAAAAAAAAAAAAAAAAaUA3PWAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAEHvb8ut++j5k04g97fl1v30fMmhmeyfRXnfrt"
    "6KiAyPNAAAAAAAAAAAAAAAAGh4H/AAsP7qHm0Z40PA/4WH91DzaL8jXLV2mu3pD3ceb/ACzFe6l4HY483+WYr3UvAutqn0ar+23p"
    "KggMTywAAAAAAAAAAAAAAABcN1fllfey8EVPXDdX5ZX3svBFbk+/6L+1/Z9JTQDS3s8x3+biPez86rwe+O/zcR72fnVeDFOuXlW1"
    "z6gDjgAAAAAAAAAAAAAAADSgG56wAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAg97fl1v30fMmnEVvBgcTjsFCzho0lON2k60rWlOCkZU7P1o3jGs4K82JnLtEadClCU6tZv7KO3HWdWs39lHbjrZeC3LO"
    "5g6WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3IsSnVrN/ZR246zq1m/so7"
    "cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWd"
    "Ws39lHbjrOC3LO46WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3ItoeB/ws"
    "P7qHm0UnGZRjMDa9di+Rat9ispx017VKUrpr3F1wFaVwOGrTirahWmzRdk0tGMzExE+MtHa1tW1sYmNEa3Q483+WYr3UvA7HNmFm"
    "5fwN+zbppnchKMacXDWi2dU+jTb2z6Sz4SnVrN/ZR246zq1m/so7cdbJwW5Z3PO6WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HS"
    "zOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrO"
    "C3LO46WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3IsSnVrN/ZR246zq1m/"
    "so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3ItcN1fllfey8EUDd3fzOzCty7GFu3HhlOVyFKUp261qnd1J"
    "Rllk+TKkqRvTpyqcWmlI8S3KpaJxms4bcF3b0tXM/KJjROtNgL21nmO/zcR72fnVeCZxW72a3MTeuQtUrGc5SjXlx4q1rWnZeXVr"
    "N/ZR2462SaWxnRO55s5WZjP4217EWJTq1m/so7cdZ1azf2UduOtzgtyzuc6WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25"
    "FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO4"
    "6WZyW3IsSnVrN/ZR246zq1m/so7cdZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3IsSnVrN/ZR246zq1m/so7cd"
    "ZwW5Z3HSzOS25FiU6tZv7KO3HWdWs39lHbjrOC3LO46WZyW3LsA2PTAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAB5YjE4fC2q3sTcjat045Trooqmab8xpptZXb5VeL192nB/DDX3"
    "k6Zd7+2Pr4I2vWuuVpxWMwuDtVvYq7Gzbp+9KujT2qfSqma781rptZVDRTi5xcpw/wAMNfeVXF4zFY27W9irsrtyvZlXi7VKcVO4"
    "8WzL7WtdNvyn7KL50zoro/l64nFYnF3a3sTdlduV/enXT3KfQ1PLPl2E9zb8yjJ2sZZ8uwnubfmUR7uMK0w2y7ka7OoBiaAAAAAA"
    "AAAAAAAAAfmUowjWU60jGNNNZVropSiu5rvpgcLptYGnO71ODlU4LUf4v3u530qUtecKxi5a0V0zOCxTnC3CVy5KkIRpplKVdFKU"
    "+mtaq1mm+2Dw+m3l8edXacHrK6aWqV8Mv/OFUcxznMczlpxd6soadMbUfFtx+qNP9XE15faxGm84+Uame2fM6K6PN15hm2PzK5y8"
    "XerOmnTG3TghHyY04F13H+Sy99PwRZ+0Dcf5LL30/BFLuYiMrCIwjGHMmZm+M7FiAYGoAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAHndvWrFut29ONu3Hhl"
    "OdaRjTu1VjNd+MParK1lkPXzpweunppbp9UeCtfsTpl3vOFYxRtetdcrNfxFjDWq3sRcjatx45zrSlPtVfNd+LUOVayuHrJcXr7l"
    "K0j/AAx4693QqeOzLG5hd9bjL0rsqfs0rwRj5MacFHM15fa1jTf8p2eCi+fM6K6P5e+Mx+Mx131uLuyvT7HKrwU8mnFTuPAGmIiN"
    "EaFIA6DWMs+XYT3NvzKMnaxlny7Ce5t+ZRk7zVX1lf2+uzqAYmgAAAAAAAAAAHytaUpprwUpx1QOab4ZbgeVbw9ed36cGiFfEpXt"
    "z1aUq0tacKxi5NorGMzgnq1pGla1ropThrWqvZrvll+D5VrCf9u/Tg0xrot0r5fZ7nfVDM8/zPNK1piLvJtdixDxbfe7PdRzXl9p"
    "Eabzj5Qovnzqro83fmeeZjmkv+1d/l8cbMPFt0/h7PdcANURERhEYQomZnTOkAdBoG4/yWXvp+CLP2gbj/JZe+n4Is/dfr/+oW5H"
    "v+ixAPPagAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAfmc4W41nclSEI00ylKuilKdutVbzXfXBYbTawFOdXuLl100tR7vHLud9KlLXnCsYo2tFdMzgsdy7bs25X"
    "Ls6W7caaZTlWlI0p261VrNN98JY02suhzm5xetlpjbp9XZl9n1qjmGb5hmU+Vi71ZxpXTG3TghH6o04HG2ZfaxGm88XlGpRfPmdF"
    "dHm6sfmmPzG56zGXpXNHDGPFCPkxpwUcoNMRERhEYKZmZ0yAOgAAAA1rAQrDAYaEv2o2rdK/XSNKMwynBSx+Y4fCUpppcnTl+RTx"
    "p12aNWY+8n219ZaO3jXP0fQGNeAAAAAAAhM03ryvLtNuMuc4in9O1WlaUr+KfFT7a9pKtbWnCsYuTaIjGZwTSDzXe3Lcv027VedY"
    "in7luvi0r+KfF3tKn5rvLmeZ1lCdz1OHr/Qt8EdH4q8cu6imrL7TxvP0hRfP8K75SeabxZnmmmF656uxXisW/Fh3ezLuowGqtYrG"
    "FYwhTMzM4zOIAk4AAAAND3Kt1hkcZe0uTlTv8n/azylKyrSlKaa14KUo1bKMHzHLMNha00Stwpy/Lr40v/arL3dsKRG238LsiPym"
    "dkOwBhaQUzFb8Yqxir1imFt1pauShStZS015NaxeXX7GdEt7Ul3xs3Z91fWptXgUfr9jOiW9qR1+xnRLe1J342bs+51qbfsvAo/X"
    "7GdEt7Ujr9jOiW9qR8bN2fc61Nv2XgUfr9jOiW9qR1+xnRLe1I+Nm7Pudam37LwKP1+xnRLe1I6/YzolvakfGzdn3OtTb9l4FH6/"
    "YzolvakdfsZ0S3tSPjZuz7nWpt+y8Cj9fsZ0S3tSOv2M6Jb2pHxs3Z9zrU2/ZeBR+v2M6Jb2pHX7GdEt7Uj42bs+51qbfsvAo/X7"
    "GdEt7Ujr9jOiW9qR8bN2fc61Nv2XgUfr9jOiW9qR1+xnRLe1I+Nm7Pudam37LwAoWAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA+VrSNK1rXRSnDWtVfzXfHLsHyreF/7d+nB4lf5dK9ufZ7iVaWt"
    "OFYxctaKxjM4J+UoxjWUq0pGnDWteClKK7mm+mAwnKtYOnO71ODlUrotUr5X73c76o5nnuZZpKvObtfVcdLMPFt07nZ7qPa8vtIj"
    "TecfKFF8+dVdHm7syzrMc0nysVdrWH7tqPi24/VHW4QaoiIjCIwhRMzOmdIA6AAAAAAA98HgMZjrvqsJalen2eTTgp5VeKndXTIt"
    "z7WDlHFZhWl7ER0Vhapw24V+mv3q/YrzM2lI0zp2eKVKWtq3m5+RywdmuPxUeTiL9NFuFeOFvj4e3JZwebe83tNp8WytYrGEACLo"
    "AACJzXeTLMs0wuXPW4in9C3wypX8VeKPddrWbThWMZcmYiMZnBLIfNd58sy2koVn6/EU/o266a0r+KXFHw9pT813rzPMdNuEua4e"
    "v9O3WumtPxT46/YhWvL7TxvP0j/qm+f4V3ymM03ozTMdMOXzfD1/pWtNNNPxS46+DtIcGqta1jCsYKJmZnGZxAEnAAAAAAAfq1au"
    "3rkbVmErlyVdEYRpWUq17VKLVk25V6co381r6u3Thph418eXlSpxU+rh+pC+ZWkY2n6eKVaWtOiHhufkcsViY5jfjow2Hrptaf37"
    "lOLuR8K+vxatW7NuNq1GkLcKaIwjTRSlKfRSj9vOzcycy2M/SGulIrGAArSZNmfzHF++uefVzOnM/mOL99c8+rmetXVHowzrkASc"
    "AAAAAAAAAAAAAAAAbEA8dvAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAR2d5vbyfB0xVy3W7WU6W4QjWlPGrSsuGteKniuxEzMRGmZcmYiMZSHFw1QOa735bgK1tWK87xFP3YV8SlfxT1aVQzTePM8zrW"
    "N256uxXisW/Fjo/F2Zd1FteX2ka7z9IUXz/Cu+Ulmm8GZZpWsb9zk2dPBYt+LDu9mvdRoNVaxWMIjCFMzMzjM4gCTgAAAAAAPTD4"
    "a/irtLOHtyu3ZcUI001WrKtxpy5N3NJ8iPH6i3XTL+KfFTud9C+ZSkflP08Uq0tbVCr4XB4rGXaWcLaleuV/djTTo7dfoWvKtxqU"
    "0Xc1nprx83t14P456u+tOEwWEwVqlnCWo2YfRGnH26146910MeZ3VraK/jH3X0yYjTbT/Dxw2Fw+EtUs4a3G1ajxRjTRR7Aza9a4"
    "AAB4YvEwwmFu4q5StYWYVnKkeOtI008GkiMdA90Zme8OWZZSsb93l3qcVi340+72Kd1T813wzHG0law3/UsV4PErpuVp259juICt"
    "a1rprw1rx1a8vtJ13nDyhRbP8K75Tuab3Znj+Vbsy5ph6/uW6+PWn4p8fe0ILjBrrStYwrGCibTM4zOIAk4AAAAAABGMpSpGNK1l"
    "XgpSnDWtViyrczH4vk3cZXmlmv7taabtaeT2O73kL3rSMbTg7Ws2nCIxV6EJ3JUhbjWc5V0RjGmmta9qlFlyrcnGYjk3cwlzW1Xh"
    "9XThu1p4I/8AnAtuW5Nl2WQ5OFtUpPRoldl41yX1y1O9kzO6mdFIw851r6ZERptp8nHl+VYDLYcjCWY2614JT45y+uVeF2AzTMzO"
    "MzjK+IiNEADgAAybM/mOL99c8+rmSGf4eWGznGW600abspx8m549Psqj3rVnGsT5Qw21z6gCTgAAAAAAAAAAAAAAADYgHjt4AAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAArW/nyiz+Zj5lxZVa38+UWfzM"
    "fMuLMj9tfVDM9lvRQgHqMYAAAAAAP1bt3Ls6W7caznKuiMY0rWta9qlFlyrcnF4ikbuYT5tbrw+qjordrTt9iP2/UhfMrSMbTg7W"
    "traoVq1Zu37kbVmErlyXBGEKVlKv1UotGV7j4i7ybuZz9TDj9TDRWdfrlxU+1bMBleBy236vCWo2+xKfHOXlSrw1djJmd1adFPxj"
    "b4tFMiI020ubBZfgsvteqwlmNqPZ0U4ZeVKvDXuukGaZmZxnSuww1ADgAAAAI/PvkuO9zPwJBH598lx3uZ+BKnur/qHLe2fRloD1"
    "mEAAAAAABN5VunmeYcm5cjzXD1/qXKeNWn4Ycff0I2vWsY2nB2KzM4RGKEpSta6KcNa8VE/lW5+Y47RcxNOaWK8OmdP5kvqhrW/K"
    "93csyvROzb9Zfpx37njT/h7Ee4lGTM7udVIw85X0yPG26EdlmQ5blcf+ta03ezen41yvd7HcSIMs2mZxmcZXRERGEaABx0AAAAAB"
    "U99snletxzOxHTKzTk4ilOPkdif8PZ/+FJbDWlJUrStNNK8FaVU7PNy51nLE5TSmivDLDVro0e7rXg7lWvt8+IjgvOGGqVGblTM8"
    "VfrCnj1xGFxOFn6vE2p2Z/dnGsa/a8mxnAHQAAAAAAAAAAAAABsQDx28AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAVrfz5RZ/Mx8y4sqtb+fKLP5mPmXFmR+2vqhmey3ooQD1GMAAClK1rSlKaa14KUos"
    "GVbnZjjOTcxVOaWK8Pj0/mVp2odjuo2vWsY2nB2tZtOERigIxlOVIQpWUpV0UjSmmtarFlW5eOxWi7jq80s14eTXhuy/h/d7veW7"
    "LMjy3K4/9a1/N0aK3p+Ncr3ex3EiyZndzOikYec619MiNdtPk4ctyfL8shycJapGVeCVyvjTl9cncDLMzM4zOMroiI0RoAHHQAAA"
    "AAAABH598lx3uZ+BII/PvkuO9zPwJU91f9Q5b2z6MtAeswgAAJXKt3MzzTRO1D1WHr/XucEf4ezLuI2tFYxtOEOxEzOERiikxlW6"
    "+Z5lyblIeow9eH11zg00/DHjr4O2t+VbqZZl2i5OPOsRTh9ZcpwUr+GHFTwptlzO78MuPrK6mR423Qh8q3YyzLeTOMPX4in9a5w1"
    "pX8MeKPh7aYBkta1pxtOMr4iIjCIwAHHQAAAAAAAAAAAH5nCFyPJnGko17Eqaafa8K5bl1a6a4WzWteOtbcdTpHYmY1SYQ5f0zLu"
    "iWfhw1H6Zl3RLPw4anUHFO2XMI2OX9My7oln4cNR+mZd0Sz8OGp1BxTtkwjY5f0zLuiWfhw1H6Zl3RLPw4anUHFO2TCNjl/TMu6J"
    "Z+HDUfpmXdEs/DhqdQcU7ZMI2OX9My7oln4cNR+mZd0Sz8OGp1BxTtkwjY5f0zLuiWfhw1H6Zl3RLPw4anUHFO2TCNjl/TMu6JZ+"
    "HDUfpmXdEs/DhqdQcU7ZMI2OX9My7oln4cNR+mZd0Sz8OGp1BxTtkwjYAOOgAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACtb+fKLP5mPmXFlVrfz5RZ/Mx8y4syP219UMz2W9FCBM5VutmeZaLlY82w9f6"
    "tylaVrT8MeOvg7b0rWrWMbThDJETM4RGKGTmVbpZnj+Tcu05rh68PLuU8atPww4+/oW/Kt2ssyylJwh67EU/r3NFZUr+GnFFLsmZ"
    "3fhlx9ZX0yPG26EXle72WZZSkrNvl36cd+540+593uJQGW1ptOMzjK6IiIwiMABx0AAAAAAAAAAAAR+ffJcd7mfgSCPz75Ljvcz8"
    "CVPdX/UOW9s+jLQd2W5JmWZy0YW1WtvTorel4tun8XZ7j1ZmIjGZwjzYYiZ0RpcKRyvIMyzSVK4e3ybPZvz8WFPqr2e4t+Vbm5fg"
    "+Tdxf/bv04dEqaLdK+R2e73lhpSkaUpSmilOClKMuZ3cRopGPnK+mROu25A5VujluA0XL9Od36fvTpTkUr+GHD9qe4uCj6Mlr2tO"
    "NpxXxWIjCIwAEXQAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAABEbyZTfzfB2cLYlGFY3o3Jzlp0UjSM6V4uPjS47W01mLRrhyYiYwnxQuVbrZZlui5WPOcR"
    "T+rcpStKV/DHip4e2mgdta1pxtOMkRERhEYACLoAAAAAAAAAAAAAAAA5M0w9zFZdicNa0esvW5Qhp4Kaa00cLrHYnCYnYTGMYKzl"
    "W5WCw3Ju4+vOr1OHkcVqlfq45d3vLJCEIRpCEaRhGmiMaU0UpSnYpSj9Dt72vONpxcrWK6owAEXQAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAHneuxsWbl6enkWo1nLRx6I001QPXjJfu3tinpJjM/"
    "l2L9zc8yrJ2nt8muZFptjo2Kc3MtWYw8WgdeMl+7e2KekdeMl+7e2Kekz8X/ABcv+29X17+TQOvGS/dvbFPSOvGS/dvbFPSZ+Hxc"
    "v+2869/JoHXjJfu3tinpJbLMzw2aYauJw1JUt0lWHj00V000V7Fa/SylcN1c+ynLssrYxl/1V2t2UuTyJy4K0jo4YxrTsK87t61r"
    "jSLTOPqnl5szbC0xELkjd4cRew2TYq/YnW3dhGNYzpx08aNHh1u3e6X/AG7voI/Pt5MlxmUYnDYbE8u9cjSkI8i5TTolGvHKFKdh"
    "RTKvx1xpbDijHQstevDOFo1T4qx1jzzptzv01JndTOMzxmbxs4nEzu2qwnXky4tNKcCqJ7cr55D3c/A25tKxl2mK19s+DNS1uKum"
    "dbRAHmtiDxm92VYLFXMLepd9ZaryZcmNK009rxnj14yX7t7Yp6So7x/PMb7yvgojW6vbZc1iZx0xE62a2deJmNGiWgdeMl+7e2Ke"
    "kdeMl+7e2Kekz8S+Ll/23ude/k0Drxkv3b2xT0jrxkv3b2xT0mfh8XL/ALbzr38mk5dvRluZYqOEw9LtLs6VrTlxpSni001/eqmW"
    "dbmQrLPbVafuQuVrs8n/AFaKy5+XWl+GuzFdlWm1cZ2ojejFYjCZNev4a5W1djKFKTjx00ypSqjdY886bc79NS6b4/Ib/lW/Pozl"
    "o7WtZy5xiJ/LxhVnTMW0TMaFw3PzbMsdmd21i8RO9bjYlKkZcXKpO3TT9q5KDuH83vflpefbX5R3MRGZhEYaI1LcmZmmnaOLNc1w"
    "uVYeOIxVJVtznS3TkU015VaVl2a0+67Va38+UWfzMfMuK8usWvWs6plK8zFZmPB++vGS/dvbFPSOvGS/dvbFPSZ+Nvxcv+29n69/"
    "JoHXjJfu3tinpHXjJfu3tinpM/D4uX/bede/k0Drxkv3b2xT0jrxkv3b2xT0mfh8XL/tvOvfyaB14yX7t7Yp6R14yX7t7Yp6TPw+"
    "Ll/23nXv5NBjvtk0pUjSN7TKuiniU7P8Swsgs/8ANb8qnha+z9xlVy+Hhx0461uVebY4+AAzrQAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAHytKVporw0q/HqbPs49"
    "6j0AefqbPs496h6mz7OPeo9AxHn6mzTjhHvUZnvDj7ePzW9ds0pSzD+Xa0cVYw7Pdrpqsm9u8cbVueWYKdJXp05OIuRr+xTswpo/"
    "er2foUlt7XKmPzt46vRmzrxP4x4awTO7GTVzPHxlcjpwmHrSd6teKX3Yd3s9pfP0bKOg4b4MPRWZvcVpbhwxnyRplTaMccGVDVf0"
    "bKOg4b4MPRRm8eWZbZyTFXbOEsW7kYx5M4W4RlTxo8VaUQr3VZmI4Z0zglORMRM4xoZ6ntyvnkPdz8CBT25XzyHu5+Bdnfrv/mVe"
    "X76+rRAHltr8VtWq101hGta8da0o+eps+zj3qPQB5+ps+zj3qHqbPs496j0DEVHfy5bt4bC4eEYxlOcrldFNFdEKcn/cpad3yxtM"
    "TnM7ca6YYaNLVPK/al9tdHcQT08iuGXXz072PNnG87lp3CscrH4nEaOC1apDu3Jaf9i8q5uRhPU5TLESp42JuVlTyIeJT7dKxsPc"
    "WxzbeWjc0ZUYUjz0oPfH5Df8q359GctG3x+Q3/Kt+fRnLV2n65/1KnP9/wBFk3D+b3vy0vPtr8oO4fze9+Wl59tfmfuv2/SFuT7P"
    "qK1v58os/mY+ZcWVWt/PlFn8zHzLiGR+2vqlmey3ooS2bhQhK9jeVGktEbejTTT2ZKmtu4H/ADY3ybfhk3dx+q30/lmyvfC4+ps+"
    "zj3qHqbPs496j0Hm4tjz9TZ9nHvUPU2fZx71HoGI8/U2fZx71D1Nn2ce9R6BiPP1Nn2ce9R6AAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAOPMc0wWWWqXcZOsI"
    "yrohojWVZV49HBRWMfv5KtKwy6xyf/tvcNe5COtZTKvf2xo2+CNr1rrlbsRiLGGtSvYi5G1ajxzlXRRTc83znepLDZVpt268EsTX"
    "gnLyKfu/Xx/UrmNzDG4+563GXpXpdjTxU8mNOCnceEITuSpC3Gs5yrojGNNNa17VKNeX21a6b/lP2UXzpnRXR/L5Wta10143flGT"
    "YvNsRS1Ypybca/zb1aeLCn+tfoomMn3LxOIrG9mVa4ezx+pp/wAkvr+74V1wuEw+DsRw+Gt0tWo8UY+Gv01M3ua10U/K23wgpkzO"
    "m2iFZznG3d1rOEweWQh6udJyuSuUrKUpU5PjVrStOPSiuvGdfds7FfSTe9uS5hmlzDSwduk6WozpPTKkdHKrHR+1XtK91Oz72Efi"
    "Q1o5XRmkTeacU444zpdv1ItMV4sPDB7deM6+7Z2K+k8MdvZmmPwtzCX6WqWrtKUlyY1pXgrSvB41fofep2fewj8SGt44vdnOMFhp"
    "4rEWoxs26UrOtJxrx10cVK9tZEdvjGHBjjoQmc3DTxYIpPblfPIe7n4ECntyvnkPdz8Ced+u/wDmUcv319WiAPLbQABy5ljreX4G"
    "9i7nFajpjT70q8EY92rqULfLOqYvE0y/Dy02MPXTcrTild4v/Xi+vSsycvjvEeGufRDMtw1x3K5duTu3J3bleVO5KspV+msq6a1f"
    "rDYe5isRbw9qmm5dlSEaduVdDzW3cfKazuzzS7HxLem3h9PZlX9qVPqpwd96GZeKUm2zV6stK8Vohb8JhreEw1rDW/2LMIwj/DTR"
    "pewPL16W1B74/Ib/AJVvz6M5aNvj8hv+Vb8+jOW/tP1z/qWXP9/0WTcP5ve/LS8+2vyg7h/N735aXn21+Z+6/b9IW5Ps+orW/nyi"
    "z+Zj5lxZVa38+UWfzMfMuIZH7a+qWZ7LeihLbuB/zY3ybfhkqSw7pZvgcruYmWMnWFLtIUhojWWnk1lp/Z+tvz4mcq0RGM6P5Zsu"
    "Yi8TOhoIg+uOQ+3l8Oeo645D7eXw56nn9LM5LbmrjpzRvTgg+uOQ+3l8Oeo645D7eXw56jpZnJbccdOaN6cEH1xyH28vhz1HXHIf"
    "by+HPUdLM5LbjjpzRvTghI74ZFKVI0vy01rop/Ln2e4m3LVtX3RNcdrsWidUxIAi6AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAjs8y2mZ5bdw1P8Ak0cuzX6Lkf2e"
    "/wASg4XdvO8VXxMJOFPvXaerp/76PsaeLsvPtl1mIiJx2q75cWmJlTcDuFXglmGJ+u3Zp/vnT/asuAyjLsujowliNuvZnxzr9cq8"
    "Ltc+Mx2EwNmt/F3Y2rdOzLs1+ilKcNe4jbNzL6JmZ8odilK6Yj6ugU7H7+aK1hl2H007F296EdaExG9Oe368OKlbp2I26Uho7saa"
    "ftTr2uZOvCvqjOdSPP0aYMnlmuaTrpljL8q/TW7PWUzTM410xxl+ladml2etZ8O3NG5H5EbGsInej5DjPJj58VHsbzZ5YrTk4ycq"
    "U7FzRc8+lXZit78XjstvYHFWYVldpSlLsK1jo0VpLhjXTp4u05HbXras6JiLQ7OdWazGmNCvp7cr55D3c/AgU9uV88h7ufgas79d"
    "/wDMqMv319WiAPLbQfFX3g3vtYaksJlsqXMRxTv04YW/J+9X7EqUtecKwja0VjGXrvTvHHAWpYHBz042dNE5U/pRr/ur2O+oPG+z"
    "nO5OU5yrKcq1rKVa6a1rXjrWr3wGAxWYYmOGwsOXclw/RSNOzKVexR6OXl1yq/eZZb2m8/xD1yjK7+a42GFtcEeO7c7EIU466mn4"
    "XDWcJh7eGsR5Nq1GkY07VHJk2T4fKMLSxa8a5LRW9drxzlqp2KJFiz87qWwj2xq/60ZWXwxp1yAKViD3x+Q3/Kt+fRnLRt8fkN/y"
    "rfn0Zy39p+uf9Sy5/v8Aosm4fze9+Wl59tflB3D+b3vy0vPtr8z91+36QtyfZ9RE7y5XczTK52LPDetypdtR4uVKNK05PdpWqWFN"
    "bTWYtGuNKyYiYmJ8WP3Ldy1OVu7GsLka6JQlStK0r26Vflrt3C4a/wD89mF3RxcuNJeGjy/TMu6JZ+HDU1x3keNfuo+PPMycax+m"
    "Zd0Sz8OGpRd87Nmxm8YWbcbUPUwryYUpGmnTLsUWZXcRmW4YrhoxQvlTWMccUCC/7p4LBXsjs3L2HtXJ1lPTKcIyrwSr2a0WZuZG"
    "XXimMdOCNKcU4Y4KANY/TMu6JZ+HDUfpmXdEs/DhqUfMjlnet+PPMyqz/wA1vyqeFr7lpluXUrpphbOn3cdTqUZ+dGZw4Rhw4rMv"
    "L4MdOOIApWAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAOPNMxs5ZgrmLvcNIcEYdmc6/sxozPMsyxeZ4mWIxU+VKv7Mafswj92NFl3+vz5WDw1K6IaJ3K0+mvBGn"
    "e4VQb+1y4inH42/hlzrzNuHwgpSta6KcNa8VE3gN0M4xkaXJwjhrdeGlb1dEq08immvf0JXcfKrE4XMzvRpOcZ+rsUrw8nRSlZS+"
    "vh0LkjndzNbTWkatcyll5MTHFbx8FOt7gU0fzcdw/RG3wd+s32f/AOfx0eJjq0r+K1pp9k1wFHyM3m+0LOlTYz3G7l5vho1nZ5GK"
    "jTsW66J7MtH2ICcJ251hcjWE410SjKmitK9ulWwqzvplVi9gJZjCNI4jD1jypU45wrXk6K/Vp0rsruZm0Vv46MYV5mTERM18PBQ0"
    "9uV88h7ufgQKe3K+eQ93PwNGd+u/+ZVZfvr6tEcuY463l+DuYy7GU4WqUrWMKaa8NdHZ0Op53rVu/anZu05Vu5GsJxr2Yypoq8yM"
    "MYx1eLbOOGhnmcb1ZhmVJWbdebYWvB6uFfGlT8c+z9XEhEvc3YzWuYXsHYsSnG3LRS9XxYVjXhjXlV4OJY8q3IwuHrG7mM+c3KcP"
    "qo6aW6V7fZl9n1PQ6mTl1jDDbhXWycGZedO+VYyfIMdm06VtR9Xh6V8e/Kni07UfvVaDlWUYPKrHqcNHxq/8l2X7c69uv+jshCFu"
    "FIW40hCNNEYxpopSnapR+mTNz7ZmjVXYvplxXznaAKVgACD3x+Q3/Kt+fRnLRt8fkN/yrfn0Zy39p+uf9Sy5/v8Aosm4fze9+Wl5"
    "9tflB3D+b3vy0vPtr8z91+36QtyfZ9QeOIxWGwsaTxN6FiMq6KSuSpCla/RplWjw/Wco6dhvjQ9JTFZnVErMY2u0cX6zlHTsN8aH"
    "pH6zlHTsN8aHpHDbZO4xjbDtZ9vx86j7mHhkun6zlHTsN8aHpKPvjiMPic2pcw92F636mFOXblScdNKy4NMdLR2tZjM0xMaJVZ0x"
    "wa/FBNG3O+Q2PKuefVnK+7qZll2HySzav4qzauUlPTCdyMZU0yr2JVX91EzlxhGP5K8ify+iyji/Wco6dhvjQ9I/Wco6dhvjQ9Jg"
    "4bbJ3NOMbYdo4v1nKOnYb40PSP1nKOnYb40PSOG2ydxjG2HaPlK0rTTThpXiq+uOgAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAKhv7g5ytYbGxppjbrK1c7XK8aN"
    "fsqpjXMVhrOLw9zDX48u1dpyZRZznm72Lym7WWit3CVr/Lv0p9k/or4W3tc2OHgnRMavNmzqTjxRqnW790c/sZfKeCxcuRh70uXC"
    "7XihPRoryu1XRRe4yjONJQrSUZU00lSumlaMfdeDzbMsBwYTETtR4+RSumGnyJaY/YlndtxzxVnCZ146nMvO4YwmMYauM8hvrnca"
    "aJVtTr9MocP/AK1o+y32zuVNFPVRr9NIcP2yqo+Lmf13revTzaDxcNVO3v3hw92xXLMHOl3lVpXEXI8MaUjXTSFK9mumnCruNzzN"
    "sdSscTiZyhXjt08SFfrjDRSrhXZXbcMxa04zGqIV3zsYwiMMRZ9xcHO5j72MrT+XZt8ilfxzrqpVD5Tk2Nza9S3h46LdK/zL1aeJ"
    "Cn+te00jLcuw+WYSGEw9PFjwylXjlKvHKrvc5sRWaR7rfaHMmkzPF4Q6wGBqAAAAAAAAQe+PyG/5Vvz6M5aTvXYv4jJL1qxblduV"
    "lDRCEaylXRKnYioX6Nm/QcT8Gfot3azEZc4zEfkzZ8TxavBMbh/N735aXn21+UncvAY/C5pduYnDXbMK2JRpK5blCmnl266NMqU+"
    "hdlHczE5mjZC3J9n1Vff35dh/ff7JKK1fH5bgsytxtY2362EJcqNOVKOiujRp8StHD1R3e6J/cu+mnk9xSlIrMWx8kMzKta2MTDN"
    "hpPVHd7on9y76Z1R3e6J/cu+mt+Xl7L/AG/6j0Lbas2Gk9Ud3uif3LvpnVHd7on9y76Z8vL2X+3/AE6FttWbDSeqO73RP7l30zqj"
    "u90T+5d9M+Xl7L/b/p0Lbas2Gk9Ud3uif3LvpnVHd7on9y76Z8vL2X+3/ToW21ZsNJ6o7vdE/uXfTOqO73RP7l30z5eXsv8Ab/p0"
    "LbapWz/w2/Jp4Ho+RpSNKRpwUpTRR9YGkAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAfmUYzjWM6UlGVNFY1pppWj9AK/j9zcpxVazs0lhLlfZ/safIr/poQt/cP"
    "MIVrzfEWrsfx8q3XvaJU+1ehbXuM2v8A6x9dKE5VJ8NzOZbm59SuilqEu3S5H/WtCm5ufVrorZhTt1uR/wBKtGE/l5myu5HoU81C"
    "sbiZpOtPXXrNqPZ0VlOXe0Up9qZwO5GWYetJ4qc8VKnYr4kNmPD9qyCNu4zbf+sPTQ7GVSPDH1edqzasW6WrMI27ceCMIUpGNO5R"
    "6ApWAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAP//ZUEsDBBQABgAIAAAAIQAhWqKEYgYA"
    "ANsdAAAVAAAAd29yZC90aGVtZS90aGVtZTEueG1s7FlLbxNHHL9X6ncY7R38iB2SCAfFjg1tCESJoeI43h3vDpndWc2Mk/hWwbFS"
    "paq06qFIvfVQtUUCqRf6adJStVTiK/Q/M+v1rj0mBoKKWnyw5/H7vx8zu7585SRm6IgISXnS8moXqx4iic8DmoQt71a/d2HNQ1Lh"
    "JMCMJ6TljYn0rmx++MFlvKEiEhME9IncwC0vUirdqFSkD8tYXuQpSWBvyEWMFUxFWAkEPga+MavUq9XVSoxp4qEEx8D25nBIfYL6"
    "mqW3OWHeZfCVKKkXfCYONGtSojDY4LCmf+RYdphAR5i1PJAT8OM+OVEeYlgq2Gh5VfPxKpuXKzkRUwtoC3Q988noMoLgsG7oRDjI"
    "CWu9xvql7Zy/ATA1j+t2u51uLednANj3wVKrSxHb6K3V2hOeBZAdzvPuVJvVRhlf4L8yh19vt9vN9RLegOywMYdfq642tuolvAHZ"
    "YXNe//ZWp7NawhuQHa7O4XuX1lcbZbwBRYwmh3NoHc88MjlkyNk1J3wN4GuTBJiiKoXssvSJWpRrMb7LRQ8AJrhY0QSpcUqG2Adc"
    "BzM6EFQLwBsEF3bski/nlrQsJH1BU9XyPk4xVMQU8uLpjy+ePkan956c3vvl9P7903s/O6iu4SQsUj3//ou/H36K/nr83fMHX7nx"
    "soj//afPfvv1SzdQFYHPvn70x5NHz775/M8fHjjgWwIPivA+jYlEN8gx2ucxGOYQQAbi1Sj6EaZFiq0klDjBmsaB7qqohL4xxgw7"
    "cG1S9uBtAS3ABbw6ultS+CASI5XFuwTcieIScJdz1ubCadOOllX0wigJ3cLFqIjbx/jIJbszE9/uKIVcpi6WnYiU1NxjEHIckoQo"
    "pPf4ISEOsjuUlvy6S33BJR8qdIeiNqZOl/TpoJRNU6JrNIa4jF0KQrxLvtm9jdqcudhvk6MyEqoCMxdLwkpuvIpHCsdOjXHMisjr"
    "WEUuJQ/Gwi85XCqIdEgYR92ASOmiuSnGJXV3oHW4w77LxnEZKRQ9dCGvY86LyG1+2IlwnDp1pklUxH4kDyFFMdrjyqkEL1eInkMc"
    "cLIw3LcpKYX77Nq+RcOSStME0Tsj4SoJwsv1OGZDTAzzykyvjmnyssYdQ9/ODD+/xg2t8tm3D92d9Z1s2VvgBFfNzDbqRbjZ9tzh"
    "IqDvfnfexqNkj0BBOKDvm/P75vyfb86L6vn8W/K0C5sr+OSibdjEC2/dQ8rYgRozcl2a/i3BvKAHi2ZiiPJLfhrBMBNXwoUCmzES"
    "XH1CVXQQ4RTE1IyEUGasQ4lSLuHRwiw7eesNOD+UXWtOHioBjdUuD+zySvFhM2djZqF5oJ0IWtEMlhW2cunNhNUscElpNaPavLTc"
    "ZKc085N5E+oGYf0qobZat6IhUTAjgfa7ZTAJy1sMUWa1NSTCAXEsF+yrGXeeuzeLiXK2Eufj5AmDqZN12c1UE0vKM3Tc8tab9aaH"
    "fJy2vCHclmAYp8BP6k6DWZi0PF9ZA8+uxRmL191ZVatO1ucMLolIhVTbWEaWymxlRCyZ6l9vNrQfzscARzNZTouVtdq/qIX5KYaW"
    "DIfEVwtWptNsj48UEQdRcIwGbCT2MejdsNkVUAmd3uSangjIbbMDs3LhZrUx+8omqxnM0ghn2a5fzUwstHAzznUws4J6+WxG99c0"
    "RVf8eZlSTOP/mSk6c+F+uhLooQ+nuMBI52jL40JFHLpQGlG/J+DcN7JALwRloVVCTL+A1rqSo2nfsjxMQcGFQ+3TEAkKnU5FgpA9"
    "ldl5BrNa1hWzysgYZX0mV1em9ndAjgjr6+pd1fZ7KJp0k8wRBjcbtPI8c8Yg1IX6rl5cbNq86sEzFWTplxVWaPqFo2D9zVRY5gAu"
    "iLMda05cvbnw5Jk9alN4ykD6Cxo3FT6bXk/7fB+ij/JzHkEiXtBdTWdhvjgAne2ilaZZWQlv/xaUy51xdrE4ztHZ+SVqxtkvF/f6"
    "zs5GJV8X88jh6sp8iVYKzyFmNvdHFB/cBdnb8HgzYnZFpjCzgz1hDB7wYJwNmbQtwTpi0tJZsk+GiAYnk7DOeDT7pyc/zPetAG17"
    "TrhyNmGG1zjbnXLi+tnEOYWRDC07JzZPcS4GbCrZ4m2U8xaZe4olb+KyJZR3u8yZvcu6bIlAvYbL1MnLXZZ5quJKPHKiBO5M/rqC"
    "/LWMTMpu/gMAAP//AwBQSwMEFAAGAAgAAAAhACVxG8K0BAAA+w0AABEAAAB3b3JkL3NldHRpbmdzLnhtbLRXbW/bNhD+PmD/wdDn"
    "Odab5RfUKWTLWhMkazCn2GdKom0uoiiQlB232H/fkRItpVGLuFu/xNQ9d88dj3dH5t37Z5oPDpgLwoqF5VzZ1gAXKctIsVtYnx7j"
    "4dQaCImKDOWswAvrhIX1/vrXX94d5wJLCWpiABSFmNN0Ye2lLOejkUj3mCJxxUpcALhlnCIJn3w3oog/VeUwZbREkiQkJ/I0cm07"
    "sBoatrAqXswbiiElKWeCbaUymbPtlqS4+TEW/C1+a5OIpRXFhdQeRxznEAMrxJ6UwrDRH2UDcG9IDt/bxIHmRu/o2G/Y7pHx7Gzx"
    "lvCUQclZioWAA6K5CZAUrWP/FdHZ9xX4braoqcDcsfWqG/n4MgL3FUGQ4ufLOKYNxwgsuzwku4wnOPOQNrFO8GPBdAiy6iIK1zNx"
    "qB9l3uESmcz2l9GZMxopWyTRHolzRdaM2/wyRr/DWBdYztKnLie+LGnjM+GJtmcoXofVU9U1dEcSjng9M5qSpun8ZlcwjpIcwoHS"
    "HkB1DnR06i8csvrRS/ys5Sq3zWKbqwWk/hpG2mfG6OA4LzFPoa9hHtq2NVJAAlHDkIzYH0xuKs5ZVWQfMALZN+GYMdnAGd6iKpeP"
    "KNlIVoKDA4IdT9yGPN0jjlJQ3pQohYZdsUJylhu9TLGuYGBy6OfGQo9PtaoEjtd36MQq2UE29WgGhgJRyMmLcXvPMpidYMrJ2w9P"
    "GehoHL8bwteOGFwlnGT4UZ3FRp5yHMNmNuQzDovsthKSAKMeuv8hgu8FgAvl+SNUz+OpxDFGsoK0/SRn+mTinJT3BM6c3xQZVM3P"
    "cgYn/RcoQxN5j1AuT0smJaMfTuUe9vw/ZHTULSt4D2TCLP6ESjaqtr30JxOvqVyFtojt+YET9CKBHXirPsQZ2xN/2Ye4gW2HvWzu"
    "LPBmUR/iuUE8643ND/ww6mUL1uPlupdtMnOnYW9sU9ddO5NeZO263roPCUMnmvb6WTqQnLgX+WauV0s7CJpWfIlEaz8YT/uQ9dJz"
    "nFkfEq+CmacjGJ3Pns7Vi+aBm5Vq5AGtLVaIJpygwb1684yURsKflqQweIJhcOMusqkSAw6HNSAoyvMYStkAeqN0nhFRRnir1/k9"
    "4ruWt9HgvVKYsrdnLjW/Mf8dJnFZo0eOyrpBjYrj+40lKeQdoUYuqmRjrAq4ajoQjPWPB67z1KbnOJfQaHrQ3SHdsFoXF8NPG9Vi"
    "GAkZCoIW1t9oePvQ9HjON6o/8T0qy7rNk52zsHKy20tHmUn4yuC1rD+SndtgrsbcGtMfKFWbBe1m0cpcI+voeUbmtTLfyPxWNjay"
    "cSsLjCxQsj0MWJ6T4gkmjlkq+ZblOTvi7EOLvxLVSdDzM6wkM7fbA0n1vNao2KMSR/XVCfXIakFzl4rBYY6f4YrGGZHwL0pJMoqe"
    "1Y3t6i5vtHN9M77QVZhSLl8yqAeTmYQvjHVPfBWLutJTAvW7OdGkvamv6m3lRMD0LuFSl4wb7DeNOT5sOr1RDxK/lgdxNLUjrx79"
    "Cq6d3FC0w1FJWkWveYs44zPFuEa+2M4qtmPfHbrL0B76q3A1nEZuPBxH64kdzdzQDlf/NN1t/mu7/hcAAP//AwBQSwMEFAAGAAgA"
    "AAAhAP7qoQaqAAAABAEAABMAKABjdXN0b21YbWwvaXRlbTEueG1sIKIkACigIAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAKyOQQrCMBREr1L+3qa6EClNpSCuRIQouE3S3zaQ5JckFXt7I17B5bxhHtMc384WLwzRkOewLSso0GvqjR85PO7nzQGKmKTv"
    "pSWPHFaMcGwbVQtagsZY5L2PteIwpTTXjEU9oZOxpBl97gYKTqYcw8hoGIzGE+nFoU9sV1V7poyyhsYg52mFn+w/KoEWdcJepNXm"
    "36y7deVTXHLxBVfpMswMWPsBAAD//wMAUEsDBBQABgAIAAAAIQDAWgds4QAAAFUBAAAYACgAY3VzdG9tWG1sL2l0ZW1Qcm9wczEu"
    "eG1sIKIkACigIAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAJyQwWqEMBCG74W+g8w9G5VdtYtxsavCXksLe83GqAGT"
    "kSSWltJ3b6Sn7bGn4Zth5vuZ8vSh5+hdWqfQMEh2MUTSCOyVGRm8vXakgMh5bno+o5EMDMKpenwoe3fsuefOo5UXL3UUGirUS8Pg"
    "q+3SvCiSjLTnrCN1tj+QpzhvSN6lh7o9PyfNvv6GKKhNOOMYTN4vR0qdmKTmboeLNGE4oNXcB7QjxWFQQjYoVi2Np2kcZ1SsQa+v"
    "eoZqy/O7/SIHd49btNWq/1pu6jYrHC1fpk+gVUn/qDa+e0X1AwAA//8DAFBLAwQUAAYACAAAACEAXHGVdc4EAAD7HAAAEgAAAHdv"
    "cmQvbnVtYmVyaW5nLnhtbMyY3267NhTH7yftHSKkXbZgIPzTr/0pbZKp0zRNavcADjgNqg3ImKS53UvtsfYKsw2GJDSMQLq1F3Vi"
    "+3x8OF+fY5Nv398JnmwRzeM0udPAraFNUBKmUZy83ml/vCxvPG2SM5hEEKcJutP2KNe+3//4w7ddkBRkhSifOOGMJA92WXinbRjL"
    "Al3Pww0iML8lcUjTPF2z2zAlerpexyHSdymNdNMAhvyU0TREec45jzDZwlyrcOF7P1pE4Y4bC6CthxtIGXpvGOBiyFT3da8NMgeA"
    "+BOaoI2yLkY5uvCqBbIHgbhXLdJ0GOmDh3OGkcw2yR1GstokbxiptZ1Ie4OnGUr44DqlBDL+lb7qBNK3Irvh4AyyeBXjmO0503AU"
    "BsbJ2wCPuFVNIFZ0McHVSRohbEWKkt5pBU2Cyv6mtheuB6V91dQWCPdbli/n6+id4ZwpW9ondqX5PA0LghImo6ZThHkc0yTfxFld"
    "HchQGh/cKMi2KwBbgtW8XQZ6ptq50jYvZWiAfdyvtCO49LybCIweagpEbdHHheM1lSeE7+Bm4UGhOQgu6Fl8FMBsAZwQ9TwsFMOr"
    "GHrYZLfgxD3TSnFKVQQnbgILetbAU2cOAFFxEcK0lB+iEeYHrDxi0eYynNJIF7aQwQ3M66QpieuehUAR7QNiucFwGtb1TDDRZUGb"
    "1sA9OdAwex2XqD/TtMgaWjyO9tSU7J24PV3AqhL+sAjl45x53sCMV3ISBk+vSUrhCnOPePpOeAZOpALiP9/IopEf0bvsF/un+rDG"
    "4kNUTERJ1O75LRCuckZhyH4ryOTo2xNPJX6b5PCAIn6FpKKzvDDO1gzRB4rgm5giKEkulg22kG+rpfxzHzVdjJACs/hXtEX4ZZ8h"
    "NUc8D0ayu5zGSIbV4KMFjMWjbZYjeCsGYt6oxaQzajIoZ/HL7JLUnREKYwIrNLd84UeZGvsJ3Nb9v4SqF6M1K7uz36l0iIeiatUc"
    "vgaPR5ClXEjgGYaYrzcz40SEQICqYf5tA5NXeRW3HDVd8nW5vGwPQv6veoAReswH67GwDc/nJeBL62HbnXqI4evrYY7QYzFYj+UD"
    "MB3fml1Jj+yZ7XG9svWZIhlep0hi+PoiWSNEWg4WyfJsGzjTaxWxY5Ea6vVFcs0ujcTo9SWyh0vkVbEbIBFwwezRmo2oa6sCY1RF"
    "90SMv//86/8va7uAVs0yTVguAp+HMb8UPe/JKsXSdMZjetQRJ0xsvjXkEa1gkjJQ2+kIbat9PqRGWovZcr6oLiFjtT0pkTLGn6X4"
    "uBr5FRR3RiheHT9DCu4cLJbL+ZWy+aTefqbiowruV9DbHaG3N1jvuTFzzIeyPF79gIUf6n2dA7ZR8CO9xWiH3gMl8kZI5A+WyPRd"
    "B3CVPiMl4Wem5CiJ/qOUTKTOycGLoXj3D6JC/jIgT17g8Ddp357KYB/tCPXQMv10iWkx5cvNKdM0fGDwd0G3VPBDqNMBlZfxFtQG"
    "lut7Zoej0w6mvD2eMm3b9UzD4FeIs0x5VJxhyltLi+k7puP6ZuuXk4bpdjDludh6dsBjaU3dDj/lmXiGKWvvKdOZWo5nAsM5z5Q5"
    "eYYpi8Up07Nsx3N9o8NPmcFnmP6HfgIDOLbvd8RTZZZklm1Zuu7/AQAA//8DAFBLAwQUAAYACAAAACEAFBgtPsc2AADLNQUADwAA"
    "AHdvcmQvc3R5bGVzLnhtbOx9S6/jSJbe3oD/g5Are5GdkqhnYbIHepCdhamqrqnM6l7zSrx51akrXUu6lVW1GsOezRjeTW+MNmAD"
    "4yd60RgDRm38hzqr/4PJICmRjAgyTvDwffICmXkp8UTEeX+HjBN/9dffP+573zmn8+54eP1i8Iv+i55z2By3u8P71y++fWe9nL3o"
    "nS/2YWvvjwfn9YsfnPOLv/7lv/wXf/Xxs/Plh71z7rkEDufPHjevXzxcLk+fvXp13jw4j/b5F8cn5+B+eH88PdoX99fT+1eP9unD"
    "89PLzfHxyb7s7nb73eWHV8N+f/IiIHNSoXK8v99tnPVx8/zoHC7s/lcnZ+9SPB7OD7unc0jtowq1j8fT9ul03Djns7vox71P79He"
    "Ha5kBiOO0ONuczqej/eXX7iLCWbESLm3D/rsf4/7G4ExjMCQIzDZON/DaMwCGq/cO6N0dlsYncmVzm4boaM3mQiB7TOIxNAI5+H9"
    "490eoXXeXrYPMHKhjF5599oX+8E+P8Qp3u9hFEcRir6C7Y+bD1GaDoxp4yvBHx49GT5uPvv8/eF4su/2LiVXK3uuYvUYYe9vVz7e"
    "P+y/zvfsuseW4D/3e+8/Ltd+6Zru9rhZO/f28/5y9n49fX0Kfg1+Y/9Yx8Pl3Pv4mX3e7Hbv3Pm6gz7u3PHfLA7n3Qv3E8c+Xxbn"
    "nR390AyueZ8/eF8U3rk5XyKXl7vt7sUrb9Dzj+6H39ku54fD8MrKm0Ts2t4+vA+vOYeX376NTiZy6c6l+/qFfXr5duHd+CpYm/9v"
    "ZMVPyd/YwE/2ZsfGse8vjuuYXL/gEd3vPD84nE7CX7559sRhP1+OwSCMgP/vlewrjumuv3K911vfibqfOvdfuOribN9e3A9ev2Bj"
    "uRe//fzr0+54ch3l6xfzeXDxrfO4e7Pbbp1D5IuHh93W+e2Dc/j27Gxv1//WYroYXNgcnw/u/w13+mwS5635/cZ58lyn++nB9mTy"
    "lXfD3vv28+42OLv934TEBoEkRPc/OLYXP3qDJAk2fRCJoXfHObJaMc3nxNrZt0ADGWUNNCproHFZAzFDKGOgaVkDzcoaiJEpcqDd"
    "YeuGAvZ9fhiOahYdiTWC6UiMDUxHYktgOhJTAdORWAKYjkTRwXQkegymI1FTAJ3LcSPTwoiyGxJtT6ebHSP06GaHBD262RFAj262"
    "w9ejm+3f9ehmu3M9utneW49utrOG0/VTrd7nrpkdLrmt7P54vByOF6d3cb7PT80+uLQYqMah5wU954SySAQyvmcLAnFuahub/Z6t"
    "IcxI9eP5xcN+veN97373/vnknHNP3Dl85+yPT07P3m5deogET87l+SThiI5On5x75+QcNg6mYuMR9ZBg7/D8eIegm0/2ezRazmGL"
    "zL6QIopTuCq0i58fPCPZISj1o705HfNP7Wij+Ycvduf8vPKI9JbP+72DROsrHBVjtPJjA0YmPzRgZPIjA0YmPzCIyAyLRQE1JE4F"
    "1JAYFlBD4puvn1h8C6gh8S2ghsS3gFp+vr3bXfbMxUezjoF67W61P3qPQXLP4+3u/cF2E4D84Saomfa+tk/2+5P99NDzCtNistE1"
    "Q8dZHrc/9N5hxLQrJay8nqnIyl317vCcn6ExaljGdaWHZF5XekgGdqWX38S+dNNkL0F7g4Nn3j7fXYRGyygpGe1be//sJ7T5rc2+"
    "5NewmwFYu9MZzQzEZBE0+CsvnfXEieH5brPMP7EbrfxmlfRKqNMLSCLM0ntmiuOG3/zw5JxcWPYhNyXruN8fPzpbPIpvL6ejr2tR"
    "kx8ykSiZvPn49GCfdwwrxUioh/rwBYrel/ZT7gV9vbd3Bxy5mS8f7d2+h5dBvHn35Re9d8cnD2Z6jMEhuDxeLsdHNJpBJfBf/da5"
    "+9c4E1y4IPjwA9JqF0jlIUZstUMIMj6l4xaJkptm7g47lBjK6P2N88Pd0T5tcah9fXL8V0ouDhLFt/bjk590INiW6xc/uv4HIRti"
    "9H5jn3ZeXQjLqN6hEIuUDc/Pd79zNvld3VfHHkpl6NfPF1Z/ZKkuuxuPXP40IUYuf4rApOmGB09/ERYbI5d/sTFyWItd7e3zeSd9"
    "hKpND2u5IT3s9eYHfwG94/54un/e4zEwJIjGwZAgGguP++fHwxlzxYwe4oIZPez1IqoMo4dQkmP0fnXabdGEwYhhSYIRwxIDI4Yl"
    "A0YMVQD539CJEMv/mk6EWP53dXxiSClAhBiWnqGGf6SnPBFiWHrGiGHpGSOGpWeMGJaeGeuec3/vJsF4ISZCEkvnIiTxAs3h4jw+"
    "HU/26QckkubeeW8jFEh9al+fjvfeZpbjwX+JG4GkV6PeIybbPjksIf/WuUObmkcLc14IFVF7vz8ekWprt4DD7owUDsfzzNvYTo7c"
    "U/h6b2+ch+N+65wka5Lf6+Llt/62jOT02TSUyp5f7N4/XHpvH67V/iiZST/zzhCwx27LHlDE80m4n0V025fOdvf8GE6U30wxMdRv"
    "Zhodu3mUffMtk4jdOVa8kx9zkn3nLUuO3TlVvJMfc6Z4J7PT2J1p9rC2Tx+EijBN058rxpMo3zRNi643C4dNU6TrnSIVnKZpUcxU"
    "eovNxntawEtHzWbk96sZj/x+iBXJqUDMSU5F2a7kJNIM7Bvnu50X2SFOk413fXsiOZzBkmglz/m3z0e/bh974KS+qetzN3E6nJ2e"
    "kI6h/uAq5mXkfFR2N3ISyn5HTkLZAclJKHki6e0glySnouyb5CSUnZScBNhb8REB5q34+2Heir9fx1vxVHS8VY4sQE5COR2QkwAb"
    "Kk8CbKg5MgU5CZChcrdrGSpPBWyoPAmwofIkwIbKJ2AwQ+Xvhxkqf7+OofJUdAyVpwI2VJ4E2FB5EmBD5UmADZUnATZUzdxeeruW"
    "ofJUwIbKkwAbKk8CbKgsX8xhqPz9MEPl79cxVJ6KjqHyVMCGypMAGypPAmyoPAmwofIkwIbKkwAZKne7lqHyVMCGypMAGypPAmyo"
    "/lZDfUPl74cZKn+/jqHyVHQMlacCNlSeBNhQeRJgQ+VJgA2VJwE2VJ4EyFC527UMlacCNlSeBNhQeRJgQ2UPC3MYKn8/zFD5+3UM"
    "laeiY6g8FbCh8iTAhsqTABsqTwJsqDwJsKHyJECGyt2uZag8FbCh8iTAhsqTSNPP4BGl7DX7AbzqKX1jX/3RVTCpb6JbuaOkDHVS"
    "4azktNT3IiyPxw894cZDg+ENNSK7u/3uyErUksfqUbrslQjQg89fr9J3+ESp52y6xLTOf2TKpXGjNPuL3DjuecaYvHucZn6RuyfX"
    "9wk5CooTn8oppFlehMKg50eawPw4Lo/SXGqEjry2PkrzqBEK8tL6KM2hRsUopwCSp5QMXLByUmAJy0npiZqrBoJFzVOAipqnABU1"
    "T0FL1BwZfVHzpLRFzZPSEzVXTwKLmqcAFTVPASpqnoKWqDky+qLmSWmLmielJ2o+lEFFzVOAipqnABV1zoAsJaMvap6Utqh5Unqi"
    "5jAtWNQ8BaioeQpQUfMUtETNkdEXNU9KW9Q8KT1Rc6gILGqeAlTUPAWoqHkKWqLmyOiLmielLWqeVJqoGWqOiRok4cjtsCQsciMs"
    "IEduhDnnyI0aaClytyZailDQREu8rEKZw9BSVGhyCqrSk1NQFaOcAkieUjJwwcpJgSUsJ6UnahhaEola31DlFKCihqElqahhaClV"
    "1DC0lCpqGFqSixqGlkSihqElkaj1nbOcgpaoYWgpVdQwtJQqahhakosahpZEooahJZGoYWhJJOqcAVlKRl/UMLSUKmoYWpKLGoaW"
    "RKKGoSWRqGFoSSRqGFqSihqGllJFDUNLqaKGoSW5qGFoSSRqGFoSiRqGlkSihqElqahhaClV1DC0lCpqGFr60r1lh9Dy5+2jfbr0"
    "8PqDvbHPDxc7fzO6bw8n53zcf+dse7hL/QK0ylcfY8cdebTZ6XHu9y8uz7yO15HtKVu/42dAkH3xc5eSPWBHFnmz6AWHPwUnFbHJ"
    "Bo/m2P9PZxdPB9/p963VZG5Y/reEh1u9fvHnf/5/f/7pTz//8x8+/f2///Snf/zLv/t97xvn/fPePnlTuB0tlf49dt5V6pfYHCKH"
    "XDGLiZxLxRabwZ4rQ3h+3E5eYgPd2a4Ifu2Jjn1qB5cPXse+xDVPa8Nrg9VDONWbIYUfBq5CjdvBUVofHOfpK3dQds37xdUe5+zz"
    "4nrK1p3X4crx3BnbCRQcuhW4lKPfQuiL7/bXgQLOBUOkHlpm/y7l0DLvQ+mhZbE7b4eWeZeX10PL7vy/V/6KNp5DCmdpTMbWnKkx"
    "u5U5K5fpzFXdLnvvS7hklwHXIvoRuPbYIWjsmq7ODOVKE/hJTaUZKihN3DcUrEfBcW1ZehT6p3rr0ciaDZZrmR4ltSZISmJaw67p"
    "ao0h15rgJRpNrTEaqjXMVpqvNbr6MJKqQ/C2nqY6jBTU4ZZx1FY7mE3UUTt2/t+l6spYqivBa5maujJuh64wg6mHrsT0YTgyrAlb"
    "gULmMmV80tUQ/yhSkYYEgUxTQybt0BBmJvXQkBT/UaK+TKX6EtRSNPVl2g59CTO9euvLqO/9JPXFO1Pqpi3vdt5ZxT4s0lWWmVRZ"
    "AmylqSyzdigLs5gahh8lwBOBN7kKKXOpjgSC1NSReTt0JCw4tMuh6CnQxpWovQk6ncsKlUNOnySHGElUI1TwDNWQz/LiFcfTZmhw"
    "M4y1SpeprLLOXu72vpa4//n84Knsx+AEeX+C2+9tn5T7+crZ77+0/W8fn+Rf3Tv3nrW5nw76rItl4vM7/0AG6f0n9iRHSuBVfDL+"
    "r+nK4B/RGGwpkbGax8ph1/ecXIZ6OVuMxJzAQ0l8XNSfpbozebUm7rTMyWAWWmDgUVx1Zc8E3H/D73l+3jeyp+PZe842C9xT5DtM"
    "oNevzA1/26QnuIDezeNdXVyPrcjzfb6v8H755nnvXrCfL8eoo1P0AZvns6t1b70vJEXP2JVk+af/+dOn//7HP//0d70bN5OsD59Y"
    "RVhvB2BZzngZl8FqwsOx4MhdiJoE0iI9UdSTPq8of/rHT//hv8EVJcDMJSgKj8Mi/TklkwjUBZTTVCwe/3FYcqk//5c/fPrf/6s3"
    "gEknWL40GwQ+uCwp67rz/67ZAy5V+flPpiTyG8LkN2ydAEt7sqQqLv+RkERcBkxcRvfEpYuBbb5Q4m/pFvJaEQB7wpKKYBBEalAs"
    "eFpu2b/XZJ991TuyNjjh6UcPr3r/cYOG45XXPT77XFLhXUifizmGj5Uzoo431sHDnM/2PoiEAd2a1VcGU9dzc7zwZs5edhZ67StL"
    "GIUgIf3gnK5qcfPm1y8FRhl1GP41PIfhK5nEYcByp2D+cl2tp79ommQlHogvw13Pxs3rhILJCgSrk5C6uN//z27Pv0gSfFj/eplC"
    "4I/rxSBQjEg+EDzuj+UD4RNNRPPmC2qf/u/fuRYONO8Qekq1oJ7mXWMhSizZ5uTF3tm9NSFPs2dePEYwQ5CR7ljZ0Ssaet2CAv+d"
    "Fp3Bq7zjVnk7Qz5tgUnnFLxClaNWslgM1rN1bPlcAjO4lrgRyxZ8HfnT//np0x9/AhrmXRYHZMuFymzIe5Loed2qUmOQMn+JS1Vs"
    "nh6zvDN4N7bIakewNIlYoXBZGnexBWukCTb9HcUYYCxXsBHJXUPMzTEPBGjbv4aMtqXihsJtbXEjr17isDfCsJSqHLhKEI1JXhna"
    "/eXBPrx39cL/nX0XM0YNeQfNInG6q8uxaGM4sebBQiSBuPBFG5JFp7uB4hY96Ptuu9hV27wVs1Uvn/d7pzAlFyz9io7c/3x+/WoC"
    "HiGreeraS9V28epDYFzI6o301Zep9uLVh+9DF7J6Cdb4ij2br3zp4cubxah92tKL0/rhfGLM44FNvPgwWhej9amLL0zplRcfvllX"
    "jNJvxav3DifdHZ7TC2a4q5cltny2wy6i6j//bk2MC6WagBofCqlEGFl8KNMa1PjAciJ0RtgOx4hHe3MKHs0k6y5BVQl3+YJ3VLwV"
    "e9Nkr5+MpwGukX1hMAhL69JvTMNyvOwbQ6MfLE76jZngRZnYN4zROGOmo74P+F6Fqw7FJimXeo25n087H5OyMujtChsouTHau5Lz"
    "Fc2M+hVvOJ/+xz/95d/+8dN//UPvL3//H3/+h3/6yz9Ai1mZr0/KdahKtkks6p7jkH98Y5pDUXkiI32fbRhwCbSfP6VE3md/vEUl"
    "n4cpPi+H6BP/iuXP//mnn/8T9DnkfRZzGsEQfzcnxxD/wewIxhK5rw40Pv6ybBPe6yh2M6aqkPxtlBIhjWFCkr+j2iAhFb+rTVU0"
    "/v5FiWgmMNHI3wptkGjK3n6oKih/46BEUFOYoIItjK0VFP6+P1Up+Tv2JFKawaQkfVGosZ4O9h5oCem5v3lOIq45TFwhniOjEryC"
    "XQBm4AvSG5t1ERNLLAU2CF4GCV9gwd04eStbKG4suCrDnb+m3C9XDwTvwvvXgEYVEQT/XuTby+kYbs+QWI4AdgSfgDBZlC/6K+Br"
    "3NeDtIBrAL1uLIBR+ubAl6rD47dwsLQ0uTOKesWavVaN9oq1eD/yMKxQReuGc4M1eGRbuvzfsk2yHkDLF1RSEf78p9//5fd/6OlV"
    "CaQvjoVyB9loqYyKWAdfg0sevgcz9AFyAWnW935UgqsWrohwgq8ecYcHwlgRblfS9ttV6QRvKNzZhzBOhDuBQJw4P9r7/cp+Eqx9"
    "1R/3R5K1h5MIvxu4zlwM4fda8ic4AjmSN6bn5Y941wEq1/g9TZGzKoHsCpwtGrtEy9dNMd79esWtNHripWipsKRbDJTyJt18rxE2"
    "ThjP5Wy4NriICJvfPuL3iPYO1xELWxo/x5lldtnDTsjW5I+xNhnL42nrnPxnh6wNRtQKvJUm0i5GPEi7btvPgiYZWveqJnziu8P2"
    "Glo371zsu3Xe5Lv9N3q3v4qz3/8Vrn6i95Aih30DNTB8XaJoDQwcVcx3pz0gCi8GldR4jS6u0CwL/ub4cWkftm93P17ZEjiZ8Bsu"
    "efk3VIwiut0zKt+M51zp+q5JVKpLt/V+7fHZV6P73el8cTnENIeTZ4iKophIS7jRKJQQlc/nTXBlUyi7o67psAu6thcgg4Qrug6V"
    "8DGJ61fnEVz3RHfjB/st8Ak3QUrEurdJqu2TKjNWd1ZMrOklLiUNQSJ15/rOwW+c04URyxC/UFj4fHXl+3CNWpu9Y/u1kUjYdX+9"
    "3+33Hk7wfq5St9jFeKjxrvn42g99IKkx9rw5nn7sPHsg2cxLvm4cS2bCw1bCQNqErEajU0098xqFhwo6/lelmtPV3EaB5WhxUGEs"
    "ym9IsoVIlnIcZb4CgvjaWFumeZW7H8STgafLWQ4qg0B5jmgHlSDPCarCjchz5iNjYrCal6jyHl5uQJ6j8AhBxwtnku1wnqPAcrRo"
    "qDAW5Tkk2UIkS3mOMl8BYdy03EB+CzHRMB6/2tU8B5VBoDxHtFlakOdI9onWMs+ZTuZDYyX2MuwU96bkOfPlcjmeyxai7YUzyXY4"
    "z1FgOVo0VBiL8hySbCGSpTxHma+QMD4xzfX4KvdoGL8Fnk7nOZgMAuU5okYYgjxHckptLfOcsTWaTxdiL3Mr8zQgz5n1J6PFLc9N"
    "LETbC2eS7XCeo8BytGioMBblOSTZQiRLeY4yXwFhfG2tZyZTCi6M3wJPp59bYTIIlOfwO4WEeY7khPV6vp8zmI3mS7GXueWSTXg/"
    "Z7lYrdguRdFCtL1wJtkO5zkKLMd7iyN7LMpzSLKFSJbyHGW+QsL40FxY8RdN+MDT6TwHk0GgPIffByzMcyaB0TchzzGNyaovqefc"
    "nE8D8hxrOp+MJO4yFIiGF84k2+E8R4HlaNFQYSzKc0iyhUiW8hxlvgLCuLU2R2umCFwYv2lDl/McVAZB8hz7nj9Lzk905AcUyZOb"
    "cC9BwclNnXIR0DZVJJLFb3JFoVn7dCmRH1vsT3I5d/bmw/vT8dn1LwFBaSBWNvcE43xzj/GtZanI9vgcep4fvZbAmjrabVvqcEpD"
    "brkWqgTKskhmZcsMkvi9HIjObgsTP+1t9h3MAIEbedGIlrERGIkqZYI3ZsUzwQTrosCv27kgWZWSVVFOSI66HipVbG5IskOSHShH"
    "HKrkiOAWBcIccbXsTybBO53tyxGBm6DRiJaxiRqJKuWIN2bFc8QE6+Kbd0PD72KOSFalZFWUI5KjrodKFZsjkuyQZAfKEQ2VHBHc"
    "3qGDOSJwAzka0TI2oCNRpRzxxqx4jphgXXzjc2j4XcwRyaqUrIpyRHLU9VCpYnNEkh2S7EA54kglRwS3xuhgjgjcfI9GtIzN+0hU"
    "KUe8MSueIyZYF980Hhp+F3NEsiolq6IckRx1PVSq2ByRZIckO1COOFbJEcFtRTqYIwIbF6ARLaPxARJVyhFvzEq8jxhnXXzDfWj4"
    "XcwRyaqUrIpyRHLU9VCpgt9HJNnhyA6UI05UckRwS5YO5ojApg9oRMtoGoFElXLEG7MSzQrirIs3KwgNv4s5IlmVklVRjkiOuh4q"
    "VWyOSLJDkh0kR7TvN0yYfI74q9NuGzBIOTUMX8XsUGpIu/BTaCbaQaFSvTaTQqFay/T1ZB0Pl7NH5LzZ7d5583794tH+3fH0ZuHy"
    "wKPiuJF5cd7Z0Q/N4Jr3+YP3ReGdmzNjRHB5udvuAuUSBHZ8Z183sxnU2W5kbdTQNF8zHygHjDTSCqiDU1eMQISzaqOymuCOAk8X"
    "TS4vMqTuZFXL7CNkl7T3c6Ub7akYvVZOh1nSiCbgo9rq2ZD0rJ16hlHhejkYMNUQV7i0O/d1sNRFTYUUqGqVvJTpAo2tmZ0GW4w+"
    "6mhC0tJXXWwoH/IvsokbFcByFsAopLTJHKgURsGIGnyW8zI0yQ4mu4+AUxSNtWWaV8rxJtPRqzUtjpFuVIqnaqxzxRXKSOdqonMo"
    "JbOhSskM3Mi4gyUz6rGoQFWrZKZMF2huzWy83GKUUkcTkpbM6mJD+WoERfa0pZJZzpIZhZQ2mQOVzCgYUb/zUkpmJDug7D6qly9M"
    "a22sxb3y4ldrWjIj3agUT9VY54ormZHO1UTnUEpmhkrJDHyuQ8hbOteBMZBaTgdUtUpmynSB5tbMcyhajFLqaELSklldbChfjaDI"
    "Fv9UMstZMqOQ0iZzoJIZBSM6/qWUkhnJDii7j4DyxcQ01+Mr5fgRZNGrNS2ZkW5UiqdqrHPFlcxI52qicygls5FKyQx8zFUHS2Z0"
    "AocCVa2SmTJdoLk181iuFqOUOpqQtGRWFxvKVyMo8sQjKpnlLJlRSGmTOVDJjIIRnYZXSsmMZAeU3UfAJjlrPTPZsIxy/ETW6NWa"
    "lsxINyrFUzXWueJKZqRzNdE5lJLZWKVkBj71s4MlMzqQTIGqXi8zVbrQfdCNPKW0ze1jamhC8l5mNbGhnM2bCjwAkkpmOUtmFFLa"
    "ZA5UMqNgRIcDl9PLjGQHk91HQPliaC6seAepG+Ho1br2MiPdqBJP1VjnCuxlRjpXD51DKZlNVEpm4EPQO1gyo/NZFahqlcyU6QLN"
    "rZmHtrcYpdTRhKQls7rYUL4aQZHnYVPJLGfJjEJKm8yBSmYUjNplfLUtmZHsgLL7qF6+sNbmaM2GYpSj5YvbeDUumZFuVIqnaqxz"
    "xZXMSOdqonMYJbNBn+lItGL2pbPdPT/23j7YWw/dhIcvKJfLwg2vXXrDrO/9JAUXP/nWt+Yls2atojJ4DK0aM3gUnZIzeBC9t9OA"
    "w9SyQnbni5Lhg2BlIRcs9ie5vjt78+H96fjsph0BwaIew5PSl6r0smJAcF07H0mc0uznI/FDmiGgulyTKK5eRepdY/XOXasS6ZBm"
    "2UiXVFbVQNmAV33vR2jA0WvlQNiSnFYZa1aGUEVoMiiJf8k/905m8eHTb0rns33ZdLnoryQ9pfGcv84oOu5fZxyNAKAzjFZaDx+I"
    "EntSf9A41ah/UbEycapktODY7RSfFL0Bik6pfn2Pqy7JgdXswOTq0/0hm6ZKuh+eA0Ppvtyjrazpcirph3nb8583DOiMohMGdMbR"
    "CAM6w2il+/CBKN0n9QeNU436F1cai52IFY2WtwOxIElQW9J9UvQGKDql+/U9arMkB1azwx6rT/cNNk2VdD/sYU/pvtyjLY3Vaibp"
    "5XU7GSFvGNAZRScM6IyjEQZ0htFK9+EDUbpP6g8apxr1LypaJk7ziEbL22EekCSoLek+KXoDFJ3S/foeE1ZWul+vg6qqT/dHbJoq"
    "6X7Yf5fSfblHm1uzxVJSwLh1dc4bBnRG0QkDOuNohAGdYbTSffhAlO6T+oPGqUb9i4qWiU7k0Wh5a0QOSYLaku6TojdA0Sndr+8R"
    "J2W9zFOvQzaqT/fHbJoq6X7YO5DSfblHm85W/fVNkWIe7QYy84YBnVG03unUGEfnnU6NYbTSffhAlO6T+oPGqUb9C3t3P95FNRot"
    "b01UIUlQa97dJ0Wvv6JTul/f9uxlpfv1ahBefbo/YdNUSfcngT+hdF/u0az5sj+VeLRb2MwbBnRG0WpXpTGOTv8qjWG00n34QJTu"
    "k/qDxqlG/YuKlokOcNFoyaJBZ9N9UvQGKDql+/VtLVuWA6tXc9Oq0/1hZnEfvEE3PBy7hUm+9HiKcHrp3jbzdkouddkeDaVXI8ol"
    "i0RAS1hm4X6qcU0PY4oR+VZB+VaLRZ+wKl/0vFHhZTJlWfWVYWSt2qliobISyqUtjM1KnGvOifXM+0nxCtFPvOTN8e7R41JK9lq3"
    "+R6clbP3raXz6QWMcR+vjBNFcErtKgBj2W1Sh9ptUgmWESwj280bhRvYqpKAGZbwCZiRvRI0K4S1BM6UuUTgjMAZJXgVgLPMprZX"
    "cEbPzAiceX8TOMtvu4Co1sDGogTOsIRP4IzslcBZIawlcKbMJQJnBM4owasAnGW2IL6CM3ALYgJnBM7IdvNG4Qa2gSVwhiV8Amdk"
    "rwTOCmEtgTNlLhE4I3BGCV4F4CyzYfQVnIEbRhM4I3BGtps3CjewaS+BMyzhEzgjeyVwVghrCZwpc4nAGYEzSvAqAGfZHUC023sT"
    "OCNwRrabNwo3sMUygTMs4RM4I3slcFYIawmcKXOJwBmBM0rwKgBnmc3Yr+AM3IydwBmBM7Ld3Pl58xpiEzhDEz6Bs87bK4GzQlhL"
    "4EyZSwTOCJxRglc6OBtImzV+sTtrNGgMnq4VjcdE6pPoshuoT7TNblRxoipa4claGZNO1zpNor66aAG/gIMn63i4nD3hnTe73Ttv"
    "kNcvHu3fHU9vFq7ieiM7blaxOO/s6IdmcM37/MH7ovDOzZnNOri83G13unavZLo5mKjpBqMAKTvhGlij+fS2CT4yj3APfE5/WDfV"
    "1ORqBSfddJ7nWfktoNOD93Odop+5+TOMXivsEJ8qJgoK0y8HfTZ1aaDWbqlMERtmTYmOkcGkIy0j9Yw0kyxFbTQ2tjtuV6mgXY7d"
    "DeQ7WvxeG2vLNK+TjFbN6hbBUacKi+FDxRgO7rxJMRxmU4nGYjGbCpmvYaqZZCmGo7Gx3TG8SgXtcgxvIN/RYrhpuaFRvGcsfrX6"
    "GI46VVgMNxRjOLhBG8VwmE0l+s/EbCpkvoapZpKlGI7GxnbH8CoVtMsxvIF8x4vhE9Ncj6+TjAZGI3a1BjEcc6qwGD5SjOHgPj4U"
    "w2E2lWhTELOpkPkapppJlmI4GhvbHcOrVNAux/AG8h2vlm6tZyZbARcYR7GrNailY04VFsPHijEc3O6BYjjMphK7WWM2FTJf57FX"
    "FlmK4WhsbPnz8AoVtNPPw5vHd7wYPjQXVvxx8m2O0as1iOGYU4XF8IliDAfvCqYYDrOpxKanmE2FzNcw1UyyFMPR2NjuGF6lgnY5"
    "hjeQ72gx3Fqbo/XNk0cDY3x/aPUxHHWqkBg+lLb2YBEc/CJboFMlBe4KgkcbsoWMSce3TyKRLGDfy8f4lkwUmvkTmvOP4YSGgcqf"
    "f1x5Ohq5lisDAe9tHbJSXm7+lr6RH7utRTRVwo6zIJOSSK0m/RMK7SZSjL6LOYckuSarOxLbC4kBLVX3hrcMqXojJxCAt409kQ4b"
    "eeAEtlKf06aVIbI4r1FAS9Zm2ms7wtDxEH5pE35R2KgHj15V7f7TSlpK2GzcIBwD5TEhGTzjIiyDrPN5sIyC7AjNFBUPCM/UsgVi"
    "5Rvba45oimcQYRpvnZmYZqiIaeiZTBsxjcLGZXgMq2o3tFbyUkLzhQZhGiiPCdPgGRdhGmSdz4NpFGRHmKaoeECYpo6YpvpGHzXH"
    "NMUziDCNt85MTJPRbOWKacDNVgjTNADTKDRygMewqrpDaCUvJTSjaRCmgfKYMA2ecRGmQdb5PJhGQXaEaYqKB4RpaolpKm98VHdM"
    "UziDCNN468zENBnNp66YBtx8ijBNAzCNQmMbeAyrqluOVvJSQnOuBmEaKI8J0+AZF2EaZJ3Pg2kUZEeYpqh4QJimlu+eVd4IruaY"
    "pngGEabx1pmJaTKa8V0xDbgZH2GaBmAahUZfGu9PV9Q9TCt5KaFZYZP20wB5TJgGz7gI0yDrfB5MoyA7wjRFxQPCNLXENJU3xqw7"
    "pimcQYRpvHVmYpqM5qRXTANuTkqYpgGYRqHxITyGVdVNUSt5KaF5a4MwDZTHhGnwjIswDbLO58E0CrIjTFNUPCBMU0dMU32j4Jpj"
    "muIZRJjGW2cGphkM2UwEiOZXJxeLBImzOpCZBrZXMJCpEywY9b2fpMeON/DzlXjJlFir0gUeQ+tBCHgUnQIoeJCERyp2mKuDK2aY"
    "NIAEdU15GkdGPYb3ZSkeu/O1Set4AgVSwfqVnGBcIgNMkWQ7VTnDwvQQgWFIpNCOO5j1vR9FDZvi5zBVTBQUPV8OBmzq0vip3TW0"
    "g4F0ulz0V9L2aCLT1QmlOqPoBFOdcTTCqc4wWgFVfyBgSIUPhBlU8/UuE3q5doRVPbFQYBVq2WK6dAGwspZVGVpRpwoLrkPF4Apu"
    "X9fB4LqypsupZE/r7bXwvMFVZxStdnEa4+g0itIYRiu46g8EDK7wgTCDa74mOkI/147gqicWCq7iFM5aTBfxInKallUZXFGnCguu"
    "hmJwBfdR6mBwXRqr1ewmsJj53nbS5w2uOqPoBFedcTSCq84wWsFVfyBgcIUPhBpcc3VzEPq5dgRXPbFQcBVq2Wq9Xi9WylpWZXBF"
    "nSosuI4Ugyu4oUcHg+vcmi2Wktz4tqUzb3DVGUWrKaDGODrtwDSG0Qqu+gMBgyt8IMzgmm9bsdDPtSO46omFgqtQy5bWciB5KUqk"
    "ZVUGV9SpwoLrWDG4gneWd/GZ62zVX98EFjPfG5bIG1x1RtF65qoxjs4zV41h9J65ag8EfeYKHgg1uOba3yb0cy155qolFgquQi1b"
    "jNdjU/zwQaRllT5zxZwqLLhOFIPrhIJrple15sv+VGK+t6J/3uCqM4rWlkKNcXQ2E2kMoxVc9QcCBlf4QJjBNd9GC6Gfa0dw1RML"
    "BVexli1Xi0h7wSwtqzK4ok4VElyHUzZ1WWgFv8U0KyeiBipGTQNS3YzMy2RMWifgZ5LUie6ZRDVCeSZNrbitSBUYpLOVCy0il75v"
    "J5cd6TlZc+L9KK5xMAdHg4y0AW/FikkGiyd5zb/DTQ0qkWH+/giNF4Iol7xyPODDrf1KcIXkwrl09kfVpV+PaYMlv23Y0Rdjb2r4"
    "Z8EDPfxrUtVWjKL5XX2TAxACytovee2gBt4vSWCoAWBoZM0GS+keMs7loxHVaoaQTVan+0E2Vb12B6p0of0NsuhiAqMK9l5WAY3W"
    "1tASv0pE4IjAEYGj2giBwBGKXNZLcyx5vaNu8Kj6fdn5AVLeVECbrraCFM/1psGkjJ3vV5hEz4zaCJNW/XF/JLG0QOIaW/OziWrt"
    "xM8mq7PxPpuq3j57VbrQbfVZdDFhUgW76CuASdbMXJs3lmatkmBSAzJBgkltFALBJBy5DNfL9VLdrVcIk6rvsJEfJuVNBbTpaitI"
    "8VxvGkzK6GFyhUngHiYEkxoAk+bL5XIs2XoaiFZjH3g2Ua1t39lkdXZ5Z1PVgknKdKF7uLPoosKk8vuhVAGTxi5QEtfORKskmNSE"
    "TJBgUguFQDAJRS7eVuG1uC4kdOsVwqTqeyXlh0l5UwFtutoKUjzXmwaTRoowCdyNimBSA2DSrD8ZRfa7xSwtEBMcJikQ1YFJCmQ1"
    "YJICVS2YpEwXCJMy6WLCpAo6W1UBk4amZYlrZ6JVEkxqQCZIMKmNQiCYhCIXc7y2THEWLHTrFcKk6rve5YdJeVMBbbraClI815sG"
    "kzL6Cl5hErivIMGkBsCk0XKxWt0yxJilBRLX2JuUTVRrb1I2WZ29SdlUtWCSMl3o3qQsuqgwqfwehRXAJHNtTaxbVS5rlQSTGpAJ"
    "EkxqoxAIJqHIZb0wTSv+WCDVrVe5N6ny/qUIe5NypgLadLUVpHiuNw0mZXSIvcIkcIdYgkkNgEnWdD4ZSSwtkLhGC9tsoloda7PJ"
    "6jSozaaqBZOU6ULbz2bRxYRJFXSbrQAmWaY1khRXRaskmNSATJBgUhuFQDAJRy5rc74W14WEbr1CmFR9J+r8MClvKqBNV19BCud6"
    "s2CSMWYzlYEk8IYklkEUj41qhTTiogqUNSErLbShRlgLcaiR1kEdapT13A2ENtTlqNDGRCAVdNfehTML8oddekahypKcOXsT7WrI"
    "Xj1ogGHJEi3EMbVjcSJjjVoANOhmwtB6qn5eNaKYQqqfDd8r1f3CVDRFj/KaVUlgFdn/1VcHskoKuNKrY10CUdRZBYj2Jl1NDzx5"
    "lahWZyKQZlUOk1Fo1yotQtRwUBHsZfqrAob2aQ9UDiM7z23nkwLtPBdtzHJYJWcqUEEMy7KoIIYQ/hKdy+M2kBLrhNpPJTGKK51V"
    "fiqKoRtWSSAX3QfWVwuoLIYqbCqMdbUwVrPTcEi3KofM7SuNoeo4rDg2ZEqfWRwDn/FDxTGy9NyWPinQ0nPRxiyOVXKSDhXHsCyL"
    "imMIATBxXkXcBlJinVD7qThGcaWzyk/FMXTDKgnoovvA+moBFcdQhU3Fsa4Wx2p2BhrpVuWQuX3FMVQdhxXHDKb0mcUx2khJll4I"
    "ZSqOVXJ+GhXHsCyLimMIATBxSlHcBlJinVD7qThGcaWzyk/FMXTDKgnoovvA+moBFcdQhU3Fsa4Wx2p28iXpVuWQuX3FMVQdhxXH"
    "RkzpM4tj4PM8qThGlp7b0icFWnou2pjFsUpOzaTiGJZlUXEMIQAmzqaL20BKrBNqPxXHKK50VvmpOIZuWCUBXXQfWF8toOIYqrCp"
    "ONbV4ljNzjsm3aocMrevOIaq47DiWEbnfe1TnKk4Rpae29InBVp6LtqoxbEqzkqm4hiWZVFxDCEAJk4kjdtASqwTaj8VxyiudFb5"
    "qTiGblglAV10H1hfLaDiGKqwqTjW1eJYzU65J92qHDK3rziGquOw4tiEKX1mcSw8dJ2KY2TpVBxDLY5Za3MkqYzHsycqjtXRsqg4"
    "hhAAE+dQx20gJdYJtZ+KYxRXOqv8VBxDN6ySgC66D6yvFlBxDFXYVBzranHMWq4WkhcQRdiBimOth8ztK46h6jikOGbfb5nSR4tj"
    "a/v0offF7nwJ9Ee1HjYNgmrR9bAg/uaO3VFvoF9Sy1UXUTzKWei1sgogaJnK9WKaJxlgupKyMxXkw9NDoehGkry8VBBePTLEFL5H"
    "r719sLeOVmSPIcVizEDMSVyB1lIcS7g4oogNFT7lZXDzjUNDGrqQKr91tJOVGgii26yEJKovB/zJUddEVfdIdcpYwRmr8imLQjuh"
    "nJVy1obnrMORYU2SfA8NIHqVstZyBGJMxtZcvPGG8tYKDKQEeXQncy2LmZ3IXXGZCcpehwrZK/TMU8pewdmr8jFIQkuh7JWy14Zn"
    "r5PhcDSMd+q8GUD0KmWv5QhkPjImRrw7RKpAKHttvDy6k72WxcxOZK+4zARlr4ZC9go9lIyyV3D2qnxOgdBSKHul7LXh2evInAyG"
    "bLECA4hepey1HIFMJ/OhoXIcAGWvbZFHd7LXspjZiewVl5mg7HWkkL1CTw2h7BWcvSo3EhZaCmWvlL02PHs1LGMwFj9+isN6yl7L"
    "EcjYGs2nC3WBUPbaeHl0J3sti5mdyF5xmQnKXscK2Su0rTdlr+DsVbnTn9BSKHul7LXh2euwP55NphIDiF6l7LUkODGYjeZLdYFQ"
    "9tp4eXTovdeSmNmN915RmQnKXicK2esk4Bxlr4Vlr8qteISWQtkrZa8Nz17n09G0LzOA6FXKXssRiGlMVn1xRUUoEMpeGy+P7mSv"
    "ZTGzE9krLjMh2at97zDuRrNXZtP3z/ueN5rrFQK+KeeuYRJSQe6ayD0C7xRNPvCzVvYNlUxE2gAzsWUrmHX43nOGp5YRzWBFegKq"
    "TTQlamvTTBihGlVQKnA1ZETavlFe9SIXtjEn3k/ME4g6kvid8wbszIYmwZ08hlGp304ogi8XgR5oAR6BWBCQupIo4+JhipfHDLBh"
    "ZPRaEBHncaUHgxmsQgi6pcASlPweN7wY42uZNQIc4QoaZzddtkU5MZgsCgBF6hUO9idzpn4YnMSrfhiAA/AeqPejOFGNOszBWTl7"
    "uXIrZsFKI30sYCQIQHk5MNjoaQBFuzcaIZXw8yqRSqLfWoxoDqyiQFYDrShQJbxirq2hJd7ERoiFEEu2JjYSsQxXo9VUvHudMEte"
    "zJJgbsLvhpeLRS0lCJhwS52kgYZclrPVylSZa/XYZTFdWmtTeaqEXtKb4qmjF3BvPEIv4edVohcFojroBZpIolEl9GLNzLV5Y1LU"
    "Aw5jVwm9EHppDXqZToerofgdhLjWE3qB+8cEcxOGFV4uFr2UIGBCL3WSBhp6McfL2VK8L0oUEKtEL2trMV0k35SQT5XQS3pTRHX0"
    "Au6NSOgl/FwxOUs0l4nFkIBlcPSS6LcYIxqKVAO9KJDVQC8KVAm9WGMXv4jrN/EuVK1DLwoGQuglWxMbiV7G5nRs3Fy6XOsJvcD9"
    "Y4K5Cb8bXi4WvZQgYEIvdZIGGnpZT0xjqdKWsXr0slqv1wv1qeKdkf5ypJ7pg/tIUqYffq6YyCjkufBMXyE70sn0oUkXGlXK9K2h"
    "aVniWkf8UXPrMn0ovKRMvz2Z/mhlLCfxWp5Y6ynTh/vHBHMTfje8XGymX4KAKdOvkzTQMv3VatVfiw+PEQXEKjP9pbUcrMUASjRV"
    "ek6R3kBUHb2A+4gSegk/V0zOEo2YYjEk6IOlsUck3ps0RjQUqc4ekWyyOntEsqkSejHX1sQSO+tx7Grr0IuCgRB6ydbERqKX4XSy"
    "mIorfHGtJ/QC948J5ib8bni54D0ixQuY0EudpIG3R2SyXpviTZOigFjpHpHxemyKQaFoqoRe0hvIqqMXcB9ZQi/h54rJmULuDkcv"
    "ChmfDnqBJpJoVAm9WKY1MsWBJf5EpnXoBQqZCb20B70sJ+NJX5yixLWe0AvcPyaYm/C74eVi0UsJAib0UidpoKEXa7keLeOPLtIC"
    "YpXoxVquFpJTWkVTJfSS3kD4no0uRC/eERiBeSkjlnCfaUcQS64stPS+sIX66rR0dMDstfSEZ24aC0PsgeMOJPDAq5VWPFTneGJC"
    "sSThNqPkbHImm1LmZxtMTXJ4UVYZ5TLE5QoSGF1S3TkuYNX3fhQ9lYF/VgDkHRL3j+pEWRiTTRQSRl8ORmzm0jCq3eOS4inomXTZ"
    "fQspolJEjTCLIipFVDWJro21JdnnW7eYul6aYys+qdSp4kXVoWJUBfdeo6gK8D3l99OiqEpRNcIsiqoUVdUkalpuXBXXhkW+qsqo"
    "ag3Xy/VSfap4UdVQjKrgnkAUVQHiL7/PS9eiqqt25kzsDIR7nwqPqokJxaIqtz2PompMrymqVhhVJ6a5jr8Zmuarqoyq5sIarMVg"
    "RThVvKg6Uoyq4P4bFFUhSVXpPRW6FlWn5nzla1OGhZUVVRMTikVVrr0VRdWYXlNUrbACbK1nkg0jIl9VaVQdry3J69LCqeJF1bFi"
    "VAXvC6eoCnquWvZe365FVWs4NfrifuLCd0ULj6qJCcWiKvc6M0XVmF5TVK0wqg49CKjsqyp9rrowTSueM6dOFS+qThSjKni/IkVV"
    "gO8pfw9a16KqMZqtF+LCVfxqWVE1MaFYVL1dpqhKUZW7XO2uzLU5kvQ7E/mqSp+rrs25pIuccKpoUdW+v++zyQvD6q9Ouy00mgaf"
    "dHX3f3Hb1H0pX4fKFcbzvHEutIQ8G8nV3zhifxRnrbHbDxYWYdpWzxUXv/NZeakJdsqWGqQZS7izLnwTcH2WirbDdtb3fhT1T2PT"
    "KloARZwoJHy+HIzZzKXRU3sLDYXRBoTRfC+ZCy2ilEBa/NHw9Qulxa+5RsHUmIytuYpitiGclrBYtIC6mC5drKeshVWGVNSpgoLq"
    "UDGognfQUFBtQFDN94650CZKCarFn1hcv6Ba/JprFFTnI2NiiDOI+GLbEFRLWCxaUC3+AG68x6eYUwUFVUMxqII30FBQbUBQzfeK"
    "udAmSgmqxR+kWb+gWvyaaxRUp5P50FBZbBuCagmLRQuqqOfCyqfa8CNsFYMqeP8MBdUmBNVcb5gLbaKUoFr8mXX1C6rFr7lGQXVs"
    "jeZT8YvZwk3EjQ6qJSwWLagWfwQjWlBFnSooqI4Vgyp4+wwF1SY8U831grnQJsp5plr4UUo1fKZa+Jrr9Ex1MBvNxc8lhO9lN/uZ"
    "avGLxXumWvjJYHjPVDGnCgqqE8WgykyUgmrbnqnmer9caBOlBNXiT/ioX1Atfs01CqqmMVlBthA3OqiWsFi0oFr8gTV4m2cwp3oL"
    "quH/zr/8/wAAAP//AwBQSwMEFAAGAAgAAAAhAC73ULlqAQAAJAQAABQAAAB3b3JkL3dlYlNldHRpbmdzLnhtbJzTzW/CIBQA8PuS"
    "/Q9N70p1akxj9bJsWbKvZB93BGqJwGsAV/vf71Grq/FidymPtu8XeDwWq71W0Y+wToLJ4tEwiSNhGHBpNln89fkwmMeR89RwqsCI"
    "LK6Fi1fL25tFlVZi/SG8xz9dhIpxqWZZXHhfpoQ4VghN3RBKYfBjDlZTj1O7IZra7a4cMNAl9XItlfQ1GSfJLG4Ze40CeS6ZuAe2"
    "08L4Jp9YoVAE4wpZuqNWXaNVYHlpgQnncD9aHTxNpTkxo8kFpCWz4CD3Q9xMu6KGwvRR0kRa/QHTfsD4Apgxse9nzFuDYGbXkbyf"
    "Mzs5knec/y2mA/BdL2J8d1xHGEJ6x3Lc86IfdzwjEnKppwV1xbmYq37ipCMeGkwB23ZN0a9o0xNY63CGmqVPGwOWrhVK2JURNlbU"
    "wOGJ5xOGJhT75n0oSxvkKgRYtSXeX7wt9Zv5fnkmYUaVgur99REn5OxuL38BAAD//wMAUEsDBBQABgAIAAAAIQAUzx2rOAMAAKkM"
    "AAASAAAAd29yZC9mb250VGFibGUueG1s3JXPbtowGMDvk/YOUe5t/hAIoNKqpWXaoT203QMYxyFWYxvFoZRru73BdtsrTKu6Sduh"
    "expu7aa+wj47AYKAlqxSDyOCmM/2D/uX7zNbOxcsNs5JIqngLdPZtE2DcCwCynst891pZ6NuGjJFPECx4KRljog0d7Zfv9oaNkPB"
    "U2nAfC6bDLfMKE37TcuSOCIMyU3RJxw6Q5EwlMLXpGcxlJwN+htYsD5KaZfGNB1Zrm3XzByTrEMRYUgx2Rd4wAhP9XwrITEQBZcR"
    "7csJbbgObSiSoJ8ITKSEPbM44zFE+RTjeAsgRnEipAjTTdhMviKNgumOrVssngGq5QDuAqCGyUU5Rj1nWDCzyKFBOU5tyqFBgfNv"
    "iykAgkEphFuZrEPd1PQCSwZpEJXDTZ6RpeaiFEVIRvPEMC5H9ArELMFigc+KTFJOWnUKHDH1DBluvu1xkaBuDCTISgMSy9Bg9QnP"
    "R910k1zouNKSN8JYNcDadl65xrDJEQPQyYh1RazjfcSFJA50nSPYvV2Fy7FVRvt2De5V2zctNRBHKJFEMbKBbhYOEaPxaBJNBEM8"
    "6+jTFEeT+DlKqNpD1iVpDzoGsmsDJ3+ZWcSBA2k+4i6MqcxHsObU5yNOYQz8ppUJWBBxShmRxhEZGsd65cuMqMyp2RUw4cHbhZa3"
    "3Ij+pecbOYA1uwedzsxIGyJ+vbq3YKTxmBH91ck46xtpI9aFla0woQxkJpQR9wVM2LWiCQ9OatebRpQJd7bvx000Spo4PDEOKceR"
    "eCQrGmBC1Uh9ZVbUl7pgIiDJMhkhvSDBChNu0URtt+139jvFnNAl4LhPmABfZU3c3fy6+3n9++bz/Yf399cf/1x9Mo5JbxCjRKtB"
    "cXoE4yZbyEbke1tiTWWQOl+yU8ZZbi0PPzeDtBJ31y94a+z7nfactyw9nsoghXLK1lJMoZhW5E9Hnybq8krXkhxSKcvVkrdwquha"
    "8l/kVIFaeiPSiOJlCfNw++Xh9psxvvw+vvwxvroaX35dlT17uub8vOZWOfsfaq4tBgklWYEtevAhZxp5Nan8KZM7aJCKUqmTr7Yy"
    "E1Hc5Pp/0XngqSLKG3L7LwAAAP//AwBQSwMEFAAGAAgAAAAhAPpvmAmSAQAAHQMAABEACAFkb2NQcm9wcy9jb3JlLnhtbCCiBAEo"
    "oAABAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAHySwU7jMBBA70j8Q+R74jjRVqsoNRIgDgiklbYrEDdjT1NDYlv2QJq/"
    "XydtU1ohbjOe5+fx2PXVtmuTT/BBW7MkLMtJAkZapU2zJP9Wd+lvkgQURonWGliSAQK54pcXtXSVtB7+eOvAo4aQRJMJlXRLskF0"
    "FaVBbqATIYuEicW19Z3AmPqGOiHfRQO0yPMF7QCFEijoKEzdbCR7pZKz0n34dhIoSaGFDgwGyjJGjyyC78K3G6bKF7LTODj4Fj0U"
    "Z3ob9Az2fZ/15YTG/hl9fnz4O1011WaclQTCayUr1NgCr+kxjFH4eH0DibvlOYmx9CDQeu4G3FiTKiu3E3NYHyf+DkNvvQpx90kW"
    "MQVBeu0wviNvwIAXCCp5HZJz31dwdLYi4GN88LUGdT3wjbXr5N5m0wlntRH38KnHz8IXEzGn9X7yu37j0XFi1W6+h8pTeXO7uiO8"
    "yFmZsiItylVRVuxXlecvY2sn+4/Cbt/Aj8ZikbI8zYsVi8by1HgQTP3LKG+sH3YjPMtOPjT/DwAA//8DAFBLAwQUAAYACAAAACEA"
    "5zdrLOcBAAAJBAAAEAAIAWRvY1Byb3BzL2FwcC54bWwgogQBKKAAAQAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACcU01v"
    "GyEQvVfqf1hxj7Fdx4osTNQ6qnJoGkveJOcpO2ujsoCAWHF/fYfdmq7bnrqn92bYx5sPxO1bZ6ojhqidXbPZZMoqtMo12u7X7Kn+"
    "fHXDqpjANmCcxTU7YWS38v07sQ3OY0gaY0USNq7ZISW/4jyqA3YQJ5S2lGld6CARDXvu2lYrvHPqtUOb+Hw6XXJ8S2gbbK58EWSD"
    "4uqY/le0cSr7i8/1yZOeFDV23kBC+TX/aQQvAVG7BKbWHcr5guKFiS3sMcq54AMQLy40Ud7MBB+Q2BwggErUPLlYXgs+4uKj90Yr"
    "SNRW+aBVcNG1qXrsvVb5f8HHRwT536F6DTqd5FTwMRVftKX7Pwg+ADIWYB/AH6KcZXeFiZ0CgxuqXLZgIgr+OyDuEfJUt6Czv2Na"
    "HVElF6qof9Bc56z6BhFzv9bsCEGDTWw4NpAeGx9TkLVOhrQL7+H42BjrRTY5gMuDPek9EL50198QH1uqLf3D7GxstvcwWB3ZGTs7"
    "3/GH6gNYGmxOFLRxnQdLLecFUc+/xydfu7u8Lb/aehkcLcKLToedB0Vjul5crMQoI3YUxYZmXMZUAuKeigqG9D9Rhbkxl7zQSNJ2"
    "j81Z4u9E3sHn4WXL2XIypa9funOMVqc8OfkTAAD//wMAUEsDBBQABgAIAAAAIQB0Pzl6wgAAACgBAAAeAAgBY3VzdG9tWG1sL19y"
    "ZWxzL2l0ZW0xLnhtbC5yZWxzIKIEASigAAEAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAjM+xisMwDAbg/eDewWhvnNxQ"
    "yhGnSyl0O0oOuhpHSUxjy1hqad++5qYrdOgoif/7Ubu9hUVdMbOnaKCpalAYHQ0+TgZ++/1qA4rFxsEuFNHAHRm23edHe8TFSgnx"
    "7BOrokQ2MIukb63ZzRgsV5QwlstIOVgpY550su5sJ9Rfdb3W+b8B3ZOpDoOBfBgaUP094Ts2jaN3uCN3CRjlRYV2FxYKp7D8ZCqN"
    "qrd5QjHgBcPfqqmKCbpr9dN/3QMAAP//AwBQSwECLQAUAAYACAAAACEAara5hYoBAABcBgAAEwAAAAAAAAAAAAAAAAAAAAAAW0Nv"
    "bnRlbnRfVHlwZXNdLnhtbFBLAQItABQABgAIAAAAIQAekRq37wAAAE4CAAALAAAAAAAAAAAAAAAAAMMDAABfcmVscy8ucmVsc1BL"
    "AQItABQABgAIAAAAIQDMl4OVzBUAACboAAARAAAAAAAAAAAAAAAAAOMGAAB3b3JkL2RvY3VtZW50LnhtbFBLAQItABQABgAIAAAA"
    "IQBWNWonLwEAAMgEAAAcAAAAAAAAAAAAAAAAAN4cAAB3b3JkL19yZWxzL2RvY3VtZW50LnhtbC5yZWxzUEsBAi0ACgAAAAAAAAAh"
    "AIWpIi+BQAAAgUAAABUAAAAAAAAAAAAAAAAATx8AAHdvcmQvbWVkaWEvaW1hZ2UxLmpwZ1BLAQItABQABgAIAAAAIQAhWqKEYgYA"
    "ANsdAAAVAAAAAAAAAAAAAAAAAANgAAB3b3JkL3RoZW1lL3RoZW1lMS54bWxQSwECLQAUAAYACAAAACEAJXEbwrQEAAD7DQAAEQAA"
    "AAAAAAAAAAAAAACYZgAAd29yZC9zZXR0aW5ncy54bWxQSwECLQAUAAYACAAAACEA/uqhBqoAAAAEAQAAEwAAAAAAAAAAAAAAAAB7"
    "awAAY3VzdG9tWG1sL2l0ZW0xLnhtbFBLAQItABQABgAIAAAAIQDAWgds4QAAAFUBAAAYAAAAAAAAAAAAAAAAAH5sAABjdXN0b21Y"
    "bWwvaXRlbVByb3BzMS54bWxQSwECLQAUAAYACAAAACEAXHGVdc4EAAD7HAAAEgAAAAAAAAAAAAAAAAC9bQAAd29yZC9udW1iZXJp"
    "bmcueG1sUEsBAi0AFAAGAAgAAAAhABQYLT7HNgAAyzUFAA8AAAAAAAAAAAAAAAAAu3IAAHdvcmQvc3R5bGVzLnhtbFBLAQItABQA"
    "BgAIAAAAIQAu91C5agEAACQEAAAUAAAAAAAAAAAAAAAAAK+pAAB3b3JkL3dlYlNldHRpbmdzLnhtbFBLAQItABQABgAIAAAAIQAU"
    "zx2rOAMAAKkMAAASAAAAAAAAAAAAAAAAAEurAAB3b3JkL2ZvbnRUYWJsZS54bWxQSwECLQAUAAYACAAAACEA+m+YCZIBAAAdAwAA"
    "EQAAAAAAAAAAAAAAAACzrgAAZG9jUHJvcHMvY29yZS54bWxQSwECLQAUAAYACAAAACEA5zdrLOcBAAAJBAAAEAAAAAAAAAAAAAAA"
    "AAB8sQAAZG9jUHJvcHMvYXBwLnhtbFBLAQItABQABgAIAAAAIQB0Pzl6wgAAACgBAAAeAAAAAAAAAAAAAAAAAJm0AABjdXN0b21Y"
    "bWwvX3JlbHMvaXRlbTEueG1sLnJlbHNQSwUGAAAAABAAEAAXBAAAn7YAAAAA"
)


# ------------------------------------------------------------------ 명령행
def template_bytes_from(path=None):
    if path:
        return open(path, 'rb').read()
    return base64.b64decode(TEMPLATE_B64)


def cli(args):
    records, comps, _, led, comp = load_excel(args.xlsx)
    print('장비대장 %d건 (시트: %s) / 장비구성 시트: %s' % (len(records), led, comp or '없음'))
    index = scan_images(args.images) if args.images else []
    if args.only:
        want = {norm_key(x) for x in args.only.split(',')}
        records = [r for r in records if r['key'] in want]
    out = args.out or ('시험장비관리대장_%s.docx' % datetime.now().strftime('%Y%m%d'))
    generate(records, comps, index, out, template_bytes_from(args.template), args.zip, args.date,
             args.limit, args.target, catalog=not args.no_catalog, source=os.path.basename(args.xlsx))


# ------------------------------------------------------------------ 창(GUI)
def gui():
    import queue
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title('시험장비 관리대장 생성기')
    root.geometry('980x760')
    st = dict(records=[], comps={}, index=[], sheets=[])
    q_ = queue.Queue()
    v = {k: tk.StringVar() for k in ('xlsx', 'img', 'led', 'comp', 'mode', 'date', 'limit', 'target')}
    v_cat = tk.BooleanVar(value=True)
    v['mode'].set('하나의 Word 파일'); v['date'].set('dot'); v['limit'].set('1024'); v['target'].set('300')

    pad = dict(padx=6, pady=3)
    f = ttk.Frame(root, padding=8); f.pack(fill='both', expand=True)
    f.columnconfigure(1, weight=1)

    def row(r, label, var, cmd=None):
        ttk.Label(f, text=label).grid(row=r, column=0, sticky='w', **pad)
        ttk.Entry(f, textvariable=var).grid(row=r, column=1, sticky='ew', **pad)
        if cmd:
            ttk.Button(f, text='찾아보기', command=cmd).grid(row=r, column=2, **pad)

    def refresh_table():
        for i in tree.get_children():
            tree.delete(i)
        for i, r in enumerate(st['records']):
            e = None
            if st['index']:
                e, _ = find_image(r, st['index'])
            tree.insert('', 'end', iid=str(i), values=(r['no'], r['id'], r['name'],
                        (e['name'] if e else ('없음' if st['index'] else '-')), len(st['comps'].get(r['key'], [])) or '-'))
        tree.selection_set(tree.get_children())
        lab_cnt.config(text='%d건 (Ctrl/Shift 클릭으로 선택 변경, 기본 전체)' % len(st['records']))

    def reload_excel(*_):
        p = v['xlsx'].get()
        if not p or not os.path.isfile(p):
            return
        try:
            recs, comps, sheets, led, comp = load_excel(p, v['led'].get() or None,
                                                       v['comp'].get() if v['comp'].get() != '' else None)
            if not v['led'].get():
                v['led'].set(led); v['comp'].set(comp)
                cb_led['values'] = sheets; cb_comp['values'] = [''] + sheets
            st.update(records=recs, comps=comps)
            log('장비대장 %d건 읽음 · 구성내역 있는 장비 %d건' % (len(recs), sum(1 for r in recs if r['key'] in comps)))
            if not comp and not v['comp'].get():
                log('※ 장비구성 시트가 없어 구성내역은 빈 칸입니다.')
            refresh_table()
        except Exception as ex:
            messagebox.showerror('엑셀 오류', str(ex))

    def pick_xlsx():
        p = filedialog.askopenfilename(filetypes=[('Excel', '*.xlsx *.xlsm'), ('All', '*.*')])
        if p:
            v['xlsx'].set(p); v['led'].set(''); v['comp'].set(''); reload_excel()

    def pick_img():
        p = filedialog.askdirectory()
        if p:
            v['img'].set(p); st['index'] = scan_images(p)
            log('사진 %d장 발견' % len(st['index'])); refresh_table()

    row(0, '엑셀 파일', v['xlsx'], pick_xlsx)
    row(1, '사진 폴더', v['img'], pick_img)
    ttk.Label(f, text='장비대장 시트').grid(row=2, column=0, sticky='w', **pad)
    cb_led = ttk.Combobox(f, textvariable=v['led'], state='readonly'); cb_led.grid(row=2, column=1, sticky='ew', **pad)
    cb_led.bind('<<ComboboxSelected>>', reload_excel)
    ttk.Label(f, text='장비구성 시트').grid(row=3, column=0, sticky='w', **pad)
    cb_comp = ttk.Combobox(f, textvariable=v['comp'], state='readonly'); cb_comp.grid(row=3, column=1, sticky='ew', **pad)
    cb_comp.bind('<<ComboboxSelected>>', reload_excel)

    opt = ttk.Frame(f); opt.grid(row=4, column=0, columnspan=3, sticky='ew', **pad)
    ttk.Label(opt, text='출력').pack(side='left')
    ttk.Combobox(opt, textvariable=v['mode'], state='readonly', width=18,
                 values=['하나의 Word 파일', '장비별 파일(ZIP)']).pack(side='left', padx=6)
    ttk.Label(opt, text='날짜').pack(side='left')
    ttk.Combobox(opt, textvariable=v['date'], state='readonly', width=6, values=['dot', 'dash', 'kr', 'han']).pack(side='left', padx=6)
    ttk.Label(opt, text='압축 기준(KB)').pack(side='left')
    ttk.Entry(opt, textvariable=v['limit'], width=6).pack(side='left', padx=6)
    ttk.Label(opt, text='목표(KB)').pack(side='left')
    ttk.Entry(opt, textvariable=v['target'], width=6).pack(side='left', padx=6)

    ttk.Checkbutton(opt, text='카탈로그(HTML)도 생성', variable=v_cat).pack(side='left', padx=12)
    lab_cnt = ttk.Label(f, text='')
    lab_cnt.grid(row=5, column=0, columnspan=3, sticky='w', **pad)
    tree = ttk.Treeview(f, columns=('no', 'id', 'name', 'img', 'cmp'), show='headings', height=12, selectmode='extended')
    for c, t, w in (('no', '순번', 50), ('id', '관리번호', 110), ('name', '물품명', 330), ('img', '사진', 200), ('cmp', '구성', 50)):
        tree.heading(c, text=t); tree.column(c, width=w, anchor='w')
    tree.grid(row=6, column=0, columnspan=3, sticky='nsew', **pad)
    f.rowconfigure(6, weight=1)
    sb = ttk.Scrollbar(f, command=tree.yview); sb.grid(row=6, column=3, sticky='ns'); tree.config(yscrollcommand=sb.set)

    btn = ttk.Button(f, text='Word 문서 만들기')
    btn.grid(row=7, column=0, columnspan=3, **pad)
    bar = ttk.Progressbar(f, maximum=100); bar.grid(row=8, column=0, columnspan=3, sticky='ew', **pad)
    txt = tk.Text(f, height=10, state='disabled'); txt.grid(row=9, column=0, columnspan=3, sticky='nsew', **pad)

    def log(m):
        q_.put(('log', m))

    def pump():
        try:
            while True:
                k, a = q_.get_nowait()
                if k == 'log':
                    txt.config(state='normal'); txt.insert('end', a + '\n'); txt.see('end'); txt.config(state='disabled')
                elif k == 'bar':
                    bar['value'] = a
                elif k == 'done':
                    btn.config(state='normal')
        except queue.Empty:
            pass
        root.after(100, pump)

    def run():
        if not st['records']:
            messagebox.showinfo('안내', '먼저 엑셀 파일을 선택하세요.'); return
        sel = [int(i) for i in tree.selection()] or list(range(len(st['records'])))
        recs = [dict(st['records'][i]) for i in sorted(sel)]
        zip_mode = v['mode'].get().startswith('장비별')
        out = filedialog.asksaveasfilename(defaultextension='.zip' if zip_mode else '.docx',
                                           initialfile='시험장비관리대장_%s' % datetime.now().strftime('%Y%m%d'),
                                           filetypes=[('ZIP' if zip_mode else 'Word', '*.zip' if zip_mode else '*.docx')])
        if not out:
            return
        btn.config(state='disabled')
        try:
            lim, tgt = int(v['limit'].get()), int(v['target'].get())
        except ValueError:
            lim, tgt = 1024, 300

        def work():
            try:
                generate(recs, st['comps'], st['index'], out, template_bytes_from(None), zip_mode, v['date'].get(),
                         lim, tgt, log, lambda a, b: q_.put(('bar', a * 100 / b)),
                         catalog=v_cat.get(), source=os.path.basename(v['xlsx'].get()))
            except Exception as ex:
                log('오류: %s' % ex)
            q_.put(('done', None))
        threading.Thread(target=work, daemon=True).start()

    btn.config(command=run)
    pump()
    if not HAVE_HEIF:
        log('※ pillow-heif 미설치 — 아이폰 HEIC 사진을 쓰려면: pip install pillow-heif')
    root.mainloop()


def main():
    ap = argparse.ArgumentParser(description='시험장비 관리대장 생성기')
    ap.add_argument('--xlsx'); ap.add_argument('--images'); ap.add_argument('--out')
    ap.add_argument('--zip', action='store_true'); ap.add_argument('--date', default='dot', choices=['dot', 'dash', 'kr', 'han'])
    ap.add_argument('--limit', type=int, default=1024); ap.add_argument('--target', type=int, default=300)
    ap.add_argument('--only'); ap.add_argument('--template')
    ap.add_argument('--no-catalog', action='store_true', help='카탈로그(HTML) 생성 안 함')
    a = ap.parse_args()
    if a.xlsx:
        cli(a)
    else:
        gui()


if __name__ == '__main__':
    main()

