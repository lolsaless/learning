"""실내공기질 실적 자료의 검증, 병합, 집계, Excel 출력을 담당한다.

UI 프레임워크에 의존하지 않으므로 Streamlit, stlite, 데스크톱 앱에서
같은 처리 규칙을 사용할 수 있다.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import BinaryIO, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


RECEIPT_COLUMN = "접수번호"
RECEIPT_GROUP_COLUMN = "접수번호_10자리"
APARTMENT_GROUPS = {"기존공동주택", "신축공동주택"}

MEASUREMENT_COLUMNS = [
    "미세먼지(PM10)",
    "초미세먼지(PM2.5)",
    "이산화탄소",
    "폼알데하이드",
    "총부유세균",
    "일산화탄소",
    "라돈",
]

OTHER_ANALYTE_COLUMNS = ["라돈(밀폐)", "벤젠", "톨루엔", "에틸벤젠", "자일렌", "스틸렌"]

OUTPUT_COLUMNS = [
    "시군명",
    "시설군",
    "세부시설군",
    "시설명",
    "주소",
    "채취지점",
    "시료명",
    *MEASUREMENT_COLUMNS,
    *OTHER_ANALYTE_COLUMNS,
    "부적합항목",
    "확인일",
    "검사결과",
    RECEIPT_COLUMN,
    RECEIPT_GROUP_COLUMN,
    "접수일자",
]

NOTE_REQUIRED_COLUMNS = {
    RECEIPT_COLUMN,
    "검체유형",
    "채취장소",
    "시료명",
    "PM10",
    "PM2.5",
    "CO2",
    "폼알데하이드",
    "부유세균",
    "일산화탄소",
    "라돈",
    *OTHER_ANALYTE_COLUMNS,
}

STAT_REQUIRED_COLUMNS = {
    RECEIPT_COLUMN,
    "접수일자",
    "시설명",
    "의뢰기관",
    "배출시설",
    "부적합항목",
    "확인일",
    "검사결과",
}


class ProcessingError(ValueError):
    """사용자가 입력 파일을 고쳐야 하는 처리 오류."""


@dataclass
class ProcessingBundle:
    merged: pd.DataFrame
    filtered: pd.DataFrame
    extracted: pd.DataFrame
    apartments: pd.DataFrame
    multi_facilities: pd.DataFrame
    summary: pd.DataFrame
    checks: pd.DataFrame


def _read_excel(payload: bytes | BinaryIO, label: str) -> pd.DataFrame:
    source = BytesIO(payload) if isinstance(payload, bytes) else payload
    try:
        frame = pd.read_excel(
            source,
            converters={RECEIPT_COLUMN: lambda value: "" if pd.isna(value) else str(value).strip()},
        )
    except Exception as exc:  # pandas/openpyxl의 여러 예외를 사용자 메시지로 통합
        raise ProcessingError(f"{label} 파일을 읽을 수 없습니다. 정상적인 .xlsx 파일인지 확인해주세요.") from exc

    normalized = [str(column).strip() for column in frame.columns]
    if len(normalized) != len(set(normalized)):
        duplicates = sorted({column for column in normalized if normalized.count(column) > 1})
        raise ProcessingError(f"{label} 파일에 중복된 열 이름이 있습니다: {', '.join(duplicates)}")
    frame.columns = normalized
    return frame


def _require_columns(frame: pd.DataFrame, required: set[str], label: str) -> None:
    missing = sorted(required.difference(frame.columns))
    if missing:
        raise ProcessingError(f"{label} 파일에 필수 열이 없습니다: {', '.join(missing)}")


def _validate_receipts(note: pd.DataFrame, stat: pd.DataFrame) -> None:
    for frame, label in ((note, "실험노트"), (stat, "조회통계")):
        receipt_values = frame[RECEIPT_COLUMN].astype("string").str.strip()
        blank_count = int((receipt_values.isna() | receipt_values.eq("")).sum())
        if blank_count:
            raise ProcessingError(f"{label} 파일에 접수번호가 빈 행이 {blank_count}개 있습니다.")
        duplicate_count = int(frame[RECEIPT_COLUMN].duplicated(keep=False).sum())
        if duplicate_count:
            raise ProcessingError(
                f"{label} 파일에 중복 접수번호 행이 {duplicate_count}개 있습니다. "
                "다대다 병합으로 행이 늘어나는 것을 방지하기 위해 처리를 중단했습니다."
            )

    note_ids = set(note[RECEIPT_COLUMN])
    stat_ids = set(stat[RECEIPT_COLUMN])
    note_only = note_ids - stat_ids
    stat_only = stat_ids - note_ids
    if note_only or stat_only:
        raise ProcessingError(
            "두 파일의 접수번호가 일치하지 않습니다. "
            f"실험노트에만 {len(note_only)}건, 조회통계에만 {len(stat_only)}건입니다."
        )


def _unique_values(series: pd.Series) -> list[object]:
    values: list[object] = []
    seen: set[str] = set()
    for value in series:
        if pd.isna(value):
            continue
        if isinstance(value, str):
            value = value.strip()
            if not value:
                continue
        key = repr(value)
        if key not in seen:
            seen.add(key)
            values.append(value)
    return values


def _join_unique(series: pd.Series) -> object:
    values = _unique_values(series)
    if not values:
        return pd.NA
    if len(values) == 1:
        return values[0]
    return " / ".join(str(value) for value in values)


def _aggregate_date(series: pd.Series) -> object:
    dates = pd.to_datetime(series, errors="coerce").dropna().drop_duplicates().sort_values()
    if dates.empty:
        return pd.NaT
    if len(dates) == 1:
        return dates.iloc[0]
    return " / ".join(value.strftime("%Y-%m-%d") for value in dates)


def _datetime_sort_key(series: pd.Series) -> pd.Series:
    try:
        return pd.to_datetime(series, errors="coerce", format="mixed")
    except TypeError:  # pandas 1.x compatibility
        return pd.to_datetime(series, errors="coerce")


def _add_check(checks: list[dict[str, object]], severity: str, item: str, count: int, details: str) -> None:
    checks.append({"수준": severity, "검증항목": item, "건수": int(count), "내용": details})


def process_excel_files(note_payload: bytes | BinaryIO, stat_payload: bytes | BinaryIO) -> ProcessingBundle:
    note = _read_excel(note_payload, "실험노트")
    stat = _read_excel(stat_payload, "조회통계")
    _require_columns(note, NOTE_REQUIRED_COLUMNS, "실험노트")
    _require_columns(stat, STAT_REQUIRED_COLUMNS, "조회통계")
    _validate_receipts(note, stat)

    try:
        merged = note.merge(
            stat,
            on=RECEIPT_COLUMN,
            how="inner",
            suffixes=("_실험노트", "_조회통계"),
            validate="one_to_one",
        )
    except Exception as exc:
        raise ProcessingError("접수번호 기준 병합에 실패했습니다. 접수번호 형식과 중복 여부를 확인해주세요.") from exc

    specimen_type = merged["검체유형_실험노트"].astype("string")
    filtered = merged[specimen_type.str.contains("실내공기질", na=False)].copy()
    if filtered.empty:
        raise ProcessingError("실험노트의 검체유형에 '실내공기질'이 포함된 행이 없습니다.")

    type_parts = filtered["검체유형_실험노트"].astype("string").str.split("/")
    malformed_type_count = int(type_parts.str.len().lt(3).fillna(True).sum())
    if malformed_type_count:
        raise ProcessingError(
            f"검체유형이 '실내공기질/구분/시설군' 형식이 아닌 행이 {malformed_type_count}개 있습니다."
        )

    extracted = pd.DataFrame(
        {
            "시군명": filtered["의뢰기관"],
            "시설군": type_parts.str[2].str.strip(),
            "세부시설군": filtered["배출시설"],
            "시설명": filtered["시설명"],
            "주소": filtered["채취장소_실험노트"],
            "시료명": filtered["시료명_실험노트"],
            "미세먼지(PM10)": filtered["PM10"],
            "초미세먼지(PM2.5)": filtered["PM2.5"],
            "이산화탄소": filtered["CO2"],
            "폼알데하이드": filtered["폼알데하이드"],
            "총부유세균": filtered["부유세균"],
            "일산화탄소": filtered["일산화탄소"],
            "라돈": filtered["라돈"],
            **{column: filtered[column] for column in OTHER_ANALYTE_COLUMNS},
            "부적합항목": filtered["부적합항목"],
            "확인일": filtered["확인일"],
            "검사결과": filtered["검사결과"],
            RECEIPT_COLUMN: filtered[RECEIPT_COLUMN],
            "접수일자": filtered["접수일자"],
        }
    ).reset_index(drop=True)

    receipt_parts = extracted[RECEIPT_COLUMN].astype("string").str.extract(r"^(.*)-(\d+)$")
    malformed_receipt_count = int(receipt_parts[0].isna().sum())
    if malformed_receipt_count:
        raise ProcessingError(
            f"접수번호가 '기본번호-순번' 형식이 아닌 실내공기질 행이 {malformed_receipt_count}개 있습니다."
        )
    extracted[RECEIPT_GROUP_COLUMN] = receipt_parts[0]

    group_keys = ["시설명", RECEIPT_GROUP_COLUMN]
    multi_mask = ~extracted["시설군"].isin(APARTMENT_GROUPS)
    multi_source = extracted[multi_mask].copy()
    apartments = extracted[~multi_mask].copy()

    if multi_source["시설명"].isna().any():
        count = int(multi_source["시설명"].isna().sum())
        raise ProcessingError(f"다중이용시설 자료 중 시설명이 빈 행이 {count}개 있습니다.")

    extracted["채취지점"] = extracted.groupby(group_keys, dropna=False)["시료명"].transform(_join_unique)
    apartments = extracted[~multi_mask].copy()
    multi_source = extracted[multi_mask].copy()

    checks: list[dict[str, object]] = []
    non_detect_count = 0
    for column in MEASUREMENT_COLUMNS:
        original = multi_source[column]
        non_detect_count += int(original.astype("string").str.strip().eq("불검출").sum())
        cleaned = original.replace(r"^\s*불검출\s*$", pd.NA, regex=True)
        numeric = pd.to_numeric(cleaned, errors="coerce")
        invalid = cleaned.notna() & numeric.isna()
        if invalid.any():
            raise ProcessingError(
                f"{column} 열에 숫자 또는 '불검출'이 아닌 값이 {int(invalid.sum())}개 있습니다."
            )
        multi_source[column] = numeric
        if numeric.notna().sum() == 0:
            _add_check(checks, "경고", f"{column} 측정값", len(multi_source), "다중이용시설의 유효한 숫자값이 없습니다.")

    if non_detect_count:
        _add_check(
            checks,
            "경고",
            "불검출 평균 처리",
            non_detect_count,
            "현재 규칙에 따라 불검출은 결측값으로 처리하여 평균에서 제외했습니다.",
        )

    conflict_columns = [
        "시군명",
        "시설군",
        "세부시설군",
        "주소",
        "부적합항목",
        "확인일",
        "검사결과",
        "접수일자",
        *OTHER_ANALYTE_COLUMNS,
    ]
    grouped = multi_source.groupby(group_keys, sort=False, dropna=False)
    for column in conflict_columns:
        conflict_count = int((grouped[column].nunique(dropna=False) > 1).sum())
        if conflict_count:
            _add_check(
                checks,
                "경고",
                f"집계 그룹 내 {column} 불일치",
                conflict_count,
                "임의의 첫 행을 선택하지 않고 고유값을 ' / '로 모두 보존했습니다.",
            )

    records: list[dict[str, object]] = []
    for (facility_name, receipt_group), group in grouped:
        record: dict[str, object] = {
            "시설명": facility_name,
            RECEIPT_GROUP_COLUMN: receipt_group,
            "채취지점": _join_unique(group["시료명"]),
            "시료명": _join_unique(group["시료명"]),
            RECEIPT_COLUMN: _join_unique(group[RECEIPT_COLUMN]),
        }
        for column in MEASUREMENT_COLUMNS:
            record[column] = group[column].mean(skipna=True)
        for column in ["시군명", "시설군", "세부시설군", "주소", *OTHER_ANALYTE_COLUMNS]:
            record[column] = _join_unique(group[column])
        record["부적합항목"] = _join_unique(group["부적합항목"])
        record["검사결과"] = _join_unique(group["검사결과"])
        record["확인일"] = _aggregate_date(group["확인일"])
        record["접수일자"] = _aggregate_date(group["접수일자"])
        if record["검사결과"] == "적합":
            record["부적합항목"] = ""
        records.append(record)

    multi_facilities = pd.DataFrame.from_records(records, columns=OUTPUT_COLUMNS)
    apartments = apartments.reindex(columns=OUTPUT_COLUMNS)

    for frame in (apartments, multi_facilities):
        frame["__확인일_정렬"] = _datetime_sort_key(frame["확인일"])
        frame.sort_values(["__확인일_정렬", RECEIPT_GROUP_COLUMN], na_position="last", inplace=True)
        frame.drop(columns="__확인일_정렬", inplace=True)
        frame.reset_index(drop=True, inplace=True)

    missing_confirmation = int(pd.to_datetime(extracted["확인일"], errors="coerce").isna().sum())
    if missing_confirmation:
        _add_check(checks, "경고", "확인일 누락", missing_confirmation, "확인일이 없거나 날짜로 변환할 수 없습니다.")
    if not checks:
        _add_check(checks, "정상", "입력 및 집계 검증", 0, "확인된 경고가 없습니다.")

    summary = pd.DataFrame(
        [
            {"처리단계": "실험노트 입력", "건수": len(note)},
            {"처리단계": "조회통계 입력", "건수": len(stat)},
            {"처리단계": "접수번호 병합", "건수": len(merged)},
            {"처리단계": "실내공기질 필터", "건수": len(extracted)},
            {"처리단계": "다중이용시설 원본", "건수": len(multi_source)},
            {"처리단계": "다중이용시설 집계", "건수": len(multi_facilities)},
            {"처리단계": "기존·신축공동주택", "건수": len(apartments)},
            {"처리단계": "최종 출력", "건수": len(multi_facilities) + len(apartments)},
        ]
    )

    return ProcessingBundle(
        merged=merged,
        filtered=filtered,
        extracted=extracted,
        apartments=apartments,
        multi_facilities=multi_facilities,
        summary=summary,
        checks=pd.DataFrame(checks, columns=["수준", "검증항목", "건수", "내용"]),
    )


def _safe_for_excel(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _prepare_for_excel(frame: pd.DataFrame) -> pd.DataFrame:
    prepared = frame.copy()
    for column in prepared.select_dtypes(include=["object", "string"]).columns:
        prepared[column] = prepared[column].map(_safe_for_excel)
    return prepared


def _format_workbook(file_bytes: BytesIO, data_sheets: Iterable[str]) -> bytes:
    workbook = load_workbook(file_bytes)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        if worksheet.max_row >= 1:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            worksheet.row_dimensions[1].height = 24

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="center")
                if hasattr(cell.value, "strftime"):
                    cell.number_format = "yyyy-mm-dd"

        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:200]]
            max_length = max((len(value) for value in values), default=10)
            header = str(column_cells[0].value or "")
            upper_bound = 60 if header in {"주소", "채취지점", "내용", "부적합항목"} else 28
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 11), upper_bound)

        if worksheet.title in data_sheets and worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions

        if worksheet.title == "검증결과":
            for row in range(2, worksheet.max_row + 1):
                if worksheet.cell(row=row, column=1).value == "경고":
                    for cell in worksheet[row]:
                        cell.fill = warning_fill

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_final_workbook(bundle: ProcessingBundle) -> bytes:
    output = BytesIO()
    data_sheets = ["다중이용시설", "기존신축공동주택"]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _prepare_for_excel(bundle.summary).to_excel(writer, index=False, sheet_name="처리요약")
        _prepare_for_excel(bundle.multi_facilities).to_excel(writer, index=False, sheet_name=data_sheets[0])
        _prepare_for_excel(bundle.apartments).to_excel(writer, index=False, sheet_name=data_sheets[1])
        _prepare_for_excel(bundle.checks).to_excel(writer, index=False, sheet_name="검증결과")
    output.seek(0)
    return _format_workbook(output, data_sheets)


def build_dataframe_workbook(frame: pd.DataFrame, sheet_name: str = "데이터") -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _prepare_for_excel(frame).to_excel(writer, index=False, sheet_name=sheet_name[:31])
    output.seek(0)
    return _format_workbook(output, [sheet_name[:31]])
