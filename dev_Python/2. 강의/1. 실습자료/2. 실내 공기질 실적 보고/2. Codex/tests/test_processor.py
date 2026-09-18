from __future__ import annotations

import sys
import unittest
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from src.processor import ProcessingError, build_final_workbook, process_excel_files  # noqa: E402


class ProcessorRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.note_bytes = (PROJECT_ROOT / "samples" / "실험노트.xlsx").read_bytes()
        cls.stat_bytes = (PROJECT_ROOT / "samples" / "조회통계.xlsx").read_bytes()
        cls.bundle = process_excel_files(cls.note_bytes, cls.stat_bytes)

    def test_reference_counts(self) -> None:
        counts = dict(zip(self.bundle.summary["처리단계"], self.bundle.summary["건수"]))
        self.assertEqual(counts["접수번호 병합"], 713)
        self.assertEqual(counts["실내공기질 필터"], 459)
        self.assertEqual(counts["다중이용시설 집계"], 178)
        self.assertEqual(counts["기존·신축공동주택"], 84)
        self.assertEqual(counts["최종 출력"], 262)

    def test_collection_points_are_scoped_to_receipt_group(self) -> None:
        grouped = self.bundle.extracted.groupby(["시설명", "접수번호_10자리"], dropna=False)
        for _, group in grouped:
            expected = " / ".join(dict.fromkeys(group["시료명"].dropna().astype(str).str.strip()))
            actual = group["채취지점"].dropna().astype(str).unique().tolist()
            self.assertEqual(actual, [expected])

    def test_output_contains_summary_and_validation_sheets(self) -> None:
        workbook = load_workbook(BytesIO(build_final_workbook(self.bundle)), read_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["처리요약", "다중이용시설", "기존신축공동주택", "검증결과"],
        )
        self.assertEqual(workbook["다중이용시설"].freeze_panes, "A2")
        self.assertIsNotNone(workbook["다중이용시설"].auto_filter.ref)

    def test_duplicate_receipt_is_rejected(self) -> None:
        note = pd.read_excel(BytesIO(self.note_bytes))
        note = pd.concat([note, note.iloc[[0]]], ignore_index=True)
        output = BytesIO()
        note.to_excel(output, index=False)
        with self.assertRaisesRegex(ProcessingError, "중복 접수번호"):
            process_excel_files(output.getvalue(), self.stat_bytes)

    def test_unmatched_receipt_is_rejected(self) -> None:
        stat = pd.read_excel(BytesIO(self.stat_bytes))
        stat.loc[0, "접수번호"] = "NOT-MATCHED-01"
        output = BytesIO()
        stat.to_excel(output, index=False)
        with self.assertRaisesRegex(ProcessingError, "일치하지 않습니다"):
            process_excel_files(self.note_bytes, output.getvalue())


if __name__ == "__main__":
    unittest.main()
