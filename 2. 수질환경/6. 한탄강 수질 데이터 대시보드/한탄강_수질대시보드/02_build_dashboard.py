"""
02_build_dashboard.py
한탄강_수질DB.xlsx 를 읽어 인터랙티브 대시보드 HTML(단일 파일)을 만든다.

사용법
    python 02_build_dashboard.py                       # 같은 폴더의 한탄강_수질DB.xlsx → 한탄강_수질대시보드.html
    python 02_build_dashboard.py DB.xlsx 출력.html
    python 02_build_dashboard.py --cdn                 # 라이브러리를 CDN에서 불러오는 가벼운 버전(인터넷 필요)

기본값은 lib/ 폴더의 ECharts·Leaflet을 HTML 안에 넣어(인라인) 인터넷 없이도 그래프가 열린다.
(지도 배경 타일만 인터넷이 필요하며, 오프라인이면 시군 경계 배경만 표시된다.)
"""
import argparse
import json
import math
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
LIB = HERE / "lib"
ITEMS_ORDER = ["TOC", "COD", "BOD", "SS", "T-N", "T-P", "색도"]
CDN = {
    "echarts": "https://cdnjs.cloudflare.com/ajax/libs/echarts/5.5.1/echarts.min.js",
    "leaflet": "https://cdnjs.cloudflare.com/ajax/libs/leaflet/1.9.4/leaflet.js",
}


def sheet_dicts(wb, name):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    return [dict(zip(hdr, r)) for r in it if any(v not in (None, "") for v in r)]


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    try:
        return float(str(v).strip().replace(",", ""))
    except ValueError:
        return None


def norm_month(v):
    if isinstance(v, datetime):
        return f"{v.year}-{v.month:02d}"
    s = str(v).strip().replace(".", "-").replace("/", "-")
    if len(s) == 6 and s.isdigit():
        return f"{s[:4]}-{s[4:]}"
    y, m = s.split("-")[:2]
    return f"{int(y)}-{int(m):02d}"


def load(db_path):
    wb = openpyxl.load_workbook(db_path, data_only=True)
    warns = []

    stations = []
    for r in sheet_dicts(wb, "지점정보"):
        if not r.get("지점명"):
            continue
        stations.append({
            "id": r.get("지점ID") or "",
            "name": str(r["지점명"]).strip(),
            "river": r.get("하천명") or "",
            "basin": r.get("수계") or "기타",
            "sigun": r.get("시군") or "",
            "addr": r.get("주소") or "",
            "lat": to_num(r.get("위도")),
            "lon": to_num(r.get("경도")),
            "coord": r.get("좌표상태") or "",
            "note": r.get("비고") or "",
            "order": to_num(r.get("정렬순서")) or 999,
        })
    stations.sort(key=lambda s: s["order"])
    names = {s["name"] for s in stations}

    items = []
    for r in sheet_dicts(wb, "항목정보"):
        items.append({"code": str(r["항목"]).strip(), "name": r.get("표시명") or r["항목"],
                      "unit": r.get("단위") or "", "dec": int(to_num(r.get("소수자리"))) if to_num(r.get("소수자리")) is not None else 2,
                      "desc": r.get("설명") or ""})
    if not items:
        items = [{"code": c, "name": c, "unit": "", "dec": 2, "desc": ""} for c in ITEMS_ORDER]
    codes = [i["code"] for i in items]

    limits, target = defaultdict(list), {}
    for r in sheet_dicts(wb, "기준값"):
        it = str(r["항목"]).strip()
        if r.get("구분") == "목표":
            target[it] = to_num(r.get("상한값"))
        else:
            limits[it].append({"code": r.get("등급코드"), "name": r.get("등급명"), "max": to_num(r.get("상한값"))})

    # 측정결과 → 지점×월 (같은 달 여러 번 채취하면 평균)
    acc = defaultdict(lambda: defaultdict(list))
    flags, dates = defaultdict(dict), defaultdict(dict)
    months = set()
    for i, r in enumerate(sheet_dicts(wb, "측정결과"), start=2):
        st = str(r.get("지점명") or "").strip()
        if not st or not r.get("조사연월"):
            continue
        if st not in names:
            warns.append(f"측정결과 {i}행: 지점정보에 없는 지점명 '{st}' — 건너뜀")
            continue
        try:
            ym = norm_month(r["조사연월"])
        except Exception:  # noqa: BLE001
            warns.append(f"측정결과 {i}행: 조사연월 형식 오류 '{r['조사연월']}' — 건너뜀")
            continue
        months.add(ym)
        for c in codes:
            v = to_num(r.get(c))
            if v is not None:
                acc[(st, ym)][c].append(v)
        d = r.get("채취일")
        if isinstance(d, datetime):
            dates[st][ym] = d.strftime("%Y-%m-%d")
        if r.get("검토표시"):
            flags[st][ym] = str(r.get("비고") or r.get("검토표시"))
    months = sorted(months)

    values = {}
    for s in stations:
        values[s["name"]] = {}
        for c in codes:
            row = []
            for m in months:
                vs = acc.get((s["name"], m), {}).get(c)
                row.append(round(sum(vs) / len(vs), 6) if vs else None)
            values[s["name"]][c] = row
        multi = [m for m in months if any(len(v) > 1 for v in acc.get((s["name"], m), {}).values())]
        for m in multi:
            warns.append(f"{s['name']} {m}: 같은 달 채취가 2회 이상 → 평균값 사용")

    yearly = defaultdict(dict)
    yinfo = {}
    for r in sheet_dicts(wb, "연평균색도"):
        st, y, v = str(r["지점명"]).strip(), int(to_num(r["연도"])), to_num(r["평균색도"])
        if v is not None:
            yearly[st][y] = v
            yinfo.setdefault(st, {"river": r.get("하천명") or "", "addr": r.get("주소") or ""})

    geo_path = LIB / "sigungu.geojson"
    geo = json.loads(geo_path.read_text(encoding="utf-8")) if geo_path.exists() else None

    data = {
        "title": "한탄강 수계 수질 대시보드",
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "source": Path(db_path).name,
        "items": items, "limits": limits, "target": target,
        "stations": stations, "months": months, "values": values,
        "flags": flags, "dates": dates,
        "yearly": yearly, "yearlyInfo": yinfo, "geo": geo,
    }
    return data, warns


def build_html(data, mode):
    tpl = (HERE / "template.html").read_text(encoding="utf-8")
    css = (LIB / "leaflet.css").read_text(encoding="utf-8")
    if mode == "inline":
        libs = "".join(f"<script>{(LIB / f).read_text(encoding='utf-8')}</script>\n"
                       for f in ("echarts.min.js", "leaflet.js"))
    else:
        libs = "".join(f'<script src="{u}"></script>\n' for u in CDN.values())
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    body = (tpl.replace("/*__LEAFLET_CSS__*/", css)
               .replace("<!--__LIBS__-->", libs)
               .replace("/*__DATA__*/null", payload))
    if mode == "artifact":
        return body
    return ('<!doctype html>\n<html lang="ko">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
            "</head>\n<body>\n" + body + "\n</body>\n</html>\n")


def main():
    ap = argparse.ArgumentParser(description="한탄강 수질 대시보드 생성")
    ap.add_argument("db", nargs="?", default=str(HERE / "한탄강_수질DB.xlsx"))
    ap.add_argument("out", nargs="?", default=None)
    ap.add_argument("--cdn", action="store_true", help="라이브러리를 CDN에서 불러옴(파일 크기 작음, 인터넷 필요)")
    ap.add_argument("--artifact", action="store_true", help=argparse.SUPPRESS)
    a = ap.parse_args()
    if not Path(a.db).exists():
        sys.exit(f"[오류] DB 파일이 없습니다: {a.db}")
    out = Path(a.out) if a.out else Path(a.db).with_name("한탄강_수질대시보드.html")
    data, warns = load(a.db)
    mode = "artifact" if a.artifact else ("cdn" if a.cdn else "inline")
    out.write_text(build_html(data, mode), encoding="utf-8")
    for w in warns:
        print("[확인]", w)
    print(f"[완료] {out}  ({out.stat().st_size / 1024:,.0f} KB)")
    print(f"  지점 {len(data['stations'])}개 · 항목 {len(data['items'])}개 · 기간 "
          f"{data['months'][0] if data['months'] else '-'} ~ {data['months'][-1] if data['months'] else '-'}")


if __name__ == "__main__":
    main()
