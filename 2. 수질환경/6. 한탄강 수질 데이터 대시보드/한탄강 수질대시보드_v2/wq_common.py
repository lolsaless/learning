"""
wq_common.py — 한탄강 수질 DB 공통 모듈 (읽기·쓰기·값 해석)
모든 스크립트(01~04)가 이 모듈로 DB를 읽고 써서 양식이 항상 같게 유지된다.
"""
import math
import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
DEFAULT_DB = HERE / "한탄강_수질DB.xlsx"
ITEMS = ["TOC", "COD", "BOD", "SS", "T-N", "T-P", "색도"]

MEAS_COLS = ["조사연월", "채취일", "지점명"] + ITEMS + ["검토표시", "비고"]
ST_COLS = ["정렬순서", "지점ID", "지점명", "하천명", "수계", "시군", "주소", "위도", "경도", "좌표상태", "좌표출처", "비고"]
ITEM_COLS = ["항목", "표시명", "단위", "소수자리", "설명"]
LIM_COLS = ["항목", "구분", "등급코드", "등급명", "상한값", "비고"]
YEAR_COLS = ["지점명", "하천명", "주소", "연도", "평균색도", "구간"]
SET_COLS = ["설정", "값", "설명"]
NET_COLS = ["수계", "중권역", "지점명", "주소", "채수지점", "위도", "경도", "분류번호"]

ITEM_INFO = [
    ("TOC", "총유기탄소(TOC)", "mg/L", 2, "하천 생활환경기준 등급 판정 항목"),
    ("COD", "화학적산소요구량(COD)", "mg/L", 1, "일부 지점만 분석"),
    ("BOD", "생물화학적산소요구량(BOD)", "mg/L", 1, "하천 생활환경기준 등급 판정 항목"),
    ("SS", "부유물질(SS)", "mg/L", 1, "하천 생활환경기준 등급 판정 항목"),
    ("T-N", "총질소(T-N)", "mg/L", 3, "하천 기준 없음(참고)"),
    ("T-P", "총인(T-P)", "mg/L", 3, "하천 생활환경기준 등급 판정 항목"),
    ("색도", "색도", "도", 0, "목표 기준(기준값 시트) 대비 달성 여부"),
]
GRADES = ["매우좋음", "좋음", "약간좋음", "보통", "약간나쁨", "나쁨", "매우나쁨"]
GRADE_CODE = ["Ia", "Ib", "II", "III", "IV", "V", "VI"]
GRADE_LIMITS = {
    "BOD": [1, 2, 3, 5, 8, 10],
    "TOC": [2, 3, 4, 5, 6, 8],
    "SS": [25, 25, 25, 25, 100, None],
    "T-P": [0.02, 0.04, 0.1, 0.2, 0.3, 0.5],
}
COLOR_TARGET = 15

# 설정: (키, 기본값, 설명, 선택지)
SETTINGS = [
    ("VWORLD_KEY", "", "V월드 인증키(지도 배경·좌표 검색). 비우면 시군 경계만 표시", None),
    ("VWORLD_DOMAIN", "", "인증키 발급 때 등록한 서비스 URL(예: http://localhost:8765). 좌표 검색에 사용", None),
    ("기본배경지도", "자동", "자동=라이트 일반지도·다크 야간지도", ["자동", "일반", "백지도", "위성", "하이브리드", "야간", "없음"]),
    ("기본항목", "색도", "대시보드를 열었을 때 처음 보이는 항목", ITEMS),
    ("평균대비_판정", "표준편차", "▲▼ 표시 기준. 표준편차=평균±k·σ 밖, 퍼센트=평균±p% 밖, 전체=평균과 다르면 모두", ["표준편차", "퍼센트", "전체"]),
    ("평균대비_k", 1, "표준편차 배수 k (보통 1)", None),
    ("평균대비_p", 30, "퍼센트 기준 p (%)", None),
    ("이상값_기준", 6, "지점별 중앙값에서 벗어난 정도(로버스트 z). 이 값 이상이고 아래 배율도 넘으면 '이상값 의심'", None),
    ("이상값_배율", 5, "지점 중앙값의 몇 배 이상(또는 몇 분의 1 이하)일 때 이상값으로 볼지", None),
    ("전월동일_항목수", 4, "전월과 같은 값이 이 개수 이상인 항목이면 '전월 동일' 점검", None),
    ("정량한계_처리", "한계값", "'<0.5'처럼 입력된 값의 계산 방법: 한계값=0.5, 절반=0.25, 제외=계산 제외", ["한계값", "절반", "제외"]),
]

BASIN = {
    "한탄강": "한탄강 본류",
    "포천천": "영평천 수계", "금현천": "영평천 수계", "송우천": "영평천 수계",
    "영평천": "영평천 수계", "외북천": "영평천 수계", "수입천": "영평천 수계",
    "신천": "신천 수계", "능안천": "신천 수계", "효촌천": "신천 수계", "덕계천": "신천 수계",
    "마개미천": "신천 수계", "회암천": "신천 수계", "청담천": "신천 수계", "상패천": "신천 수계",
}

HEAD_FILL = PatternFill("solid", fgColor="1F4E5A")
HEAD_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="맑은 고딕", size=10)
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="C9D3D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ── 값 해석 ──────────────────────────────────────────────
def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    try:
        return float(str(v).strip().replace(",", ""))
    except ValueError:
        return None


LOQ_RE = re.compile(r"^\s*(?:<|＜|ND\s*<?)\s*([\d.]+)\s*$", re.I)


def parse_value(v):
    """측정값 해석 → (숫자, 정량한계 문자열 또는 None). '<0.5' → (0.5, '<0.5'), 'ND' → (None,'ND')."""
    if isinstance(v, str):
        s = v.strip()
        m = LOQ_RE.match(s)
        if m:
            return float(m.group(1)), "<" + m.group(1)
        if s.upper() == "ND":
            return None, "ND"
    return to_num(v), None


def norm_month(v):
    if isinstance(v, datetime):
        return f"{v.year}-{v.month:02d}"
    s = str(v).strip().replace(".", "-").replace("/", "-")
    if len(s) == 6 and s.isdigit():
        return f"{s[:4]}-{s[4:]}"
    y, m = s.split("-")[:2]
    return f"{int(y)}-{int(m):02d}"


def clean_river(name):
    return re.sub(r"^\d+\)\s*", "", str(name)).strip() if name else None


# ── 읽기 ─────────────────────────────────────────────────
def sheet_dicts(wb, name):
    if name not in wb.sheetnames:
        return []
    it = wb[name].iter_rows(values_only=True)
    try:
        hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    except StopIteration:
        return []
    return [dict(zip(hdr, r)) for r in it if any(v not in (None, "") for v in r)]


def read_db(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    db = {
        "meas": sheet_dicts(wb, "측정결과"),
        "stations": sheet_dicts(wb, "지점정보"),
        "items": sheet_dicts(wb, "항목정보"),
        "limits": sheet_dicts(wb, "기준값"),
        "yearly": sheet_dicts(wb, "연평균색도"),
        "settings": sheet_dicts(wb, "설정"),
        "network": sheet_dicts(wb, "측정망지점"),
    }
    # v1 DB 호환: 없는 시트는 기본값으로 채움
    if not db["items"]:
        db["items"] = [dict(zip(ITEM_COLS, r)) for r in ITEM_INFO]
    if not db["limits"]:
        db["limits"] = default_limits()
    have = {r.get("설정") for r in db["settings"]}
    for k, v, d, _ in SETTINGS:
        if k not in have:
            db["settings"].append({"설정": k, "값": v, "설명": d})
    for s in db["stations"]:
        s.setdefault("좌표출처", "주소 기반 추정" if s.get("좌표상태") == "추정" else "")
    return db


def settings_dict(db):
    out = {k: v for k, v, _, _ in SETTINGS}
    for r in db["settings"]:
        if r.get("설정"):
            out[str(r["설정"]).strip()] = r.get("값") if r.get("값") is not None else ""
    return out


def default_limits():
    rows = []
    for it, lims in GRADE_LIMITS.items():
        for code, gname, lim in zip(GRADE_CODE, GRADES, lims + [None]):
            rows.append({"항목": it, "구분": "생활환경기준", "등급코드": code, "등급명": gname,
                         "상한값": lim, "비고": "상한 없음" if lim is None else ""})
    rows.append({"항목": "색도", "구분": "목표", "등급코드": "-", "등급명": "목표", "상한값": COLOR_TARGET,
                 "비고": "연평균 이하이면 달성(원본 판정과 일치하도록 설정, 확인 필요)"})
    return rows


# ── 쓰기 ─────────────────────────────────────────────────
def _style(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 24
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BODY_FONT, BORDER
    ws.freeze_panes = "A2"


def _table(ws, name, style="TableStyleLight11"):
    ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 2)}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)


def _sheet(wb, title, cols, rows, widths, tname, tstyle="TableStyleLight11"):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
    _style(ws, widths)
    _table(ws, tname, tstyle)
    return ws


GUIDE = [
    ("한탄강 수계 수질 DB (v2)", ""),
    ("", ""),
    ("시트", "용도 / 입력 규칙"),
    ("측정결과", "1행 = 1지점 1회 채취. 매월 지점 수만큼 행을 아래에 추가(또는 '월별결과_가져오기'로 자동 추가)."),
    ("", "조사연월은 YYYY-MM 형식(예: 2026-10). 값이 없으면 빈칸. 정량한계 미만은 '<0.5'처럼 입력."),
    ("", "검토표시 = '검토필요'로 적으면 대시보드 '자료 점검'에 표시. 비고에 항목명(BOD 등)을 적으면 그 항목만 표시."),
    ("지점정보", "지점 추가·변경 시 한 줄 추가. 정렬순서 = 상류→하류 순서(종단 그래프 순서). 위도·경도 = 지도 위치."),
    ("", "좌표상태: 확인(현장 확인) · 측정망(국가측정망 좌표) · 브이월드(자동 검색) · 추정(주소 기반 추정)."),
    ("항목정보", "단위·소수자리."),
    ("기준값", "하천 생활환경기준 등급 상한값과 색도 목표값."),
    ("연평균색도", "과거 연도별 평균 색도(2019~). 연말에 해당 연도 행 추가."),
    ("설정", "V월드 인증키, 기본 배경지도, ▲▼ 판정 기준, 이상값 기준, 정량한계 처리 방법."),
    ("측정망지점", "국가수질측정망 지점 좌표(지도 참고 레이어)."),
    ("", ""),
    ("갱신 방법", "도구모음.bat 실행 → 메뉴에서 선택. 대시보드 만들기 → 생성된 HTML 열기"),
]


def write_db(path, db, backup=True):
    """DB 전체를 표준 양식으로 저장. 기존 파일은 '백업' 폴더에 복사 후 덮어씀."""
    path = Path(path)
    if backup and path.exists():
        bdir = path.parent / "백업"
        bdir.mkdir(exist_ok=True)
        shutil.copy2(path, bdir / f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "안내"
    for g in GUIDE:
        ws.append(g)
    ws["A1"].font = Font(name="맑은 고딕", bold=True, size=14, color="1F4E5A")
    for r in ws.iter_rows(min_row=3, max_row=ws.max_row):
        r[0].font = Font(name="맑은 고딕", bold=True, size=10)
        r[1].font = Font(name="맑은 고딕", size=10)
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 110

    # 측정결과 (조사연월·정렬순서 순)
    order = {s["지점명"]: (s.get("정렬순서") or 999) for s in db["stations"]}
    meas = sorted(db["meas"], key=lambda r: (str(r.get("조사연월")), order.get(r.get("지점명"), 999)))
    ws = _sheet(wb, "측정결과", MEAS_COLS, meas, [10, 12, 12, 8, 8, 8, 8, 8, 8, 8, 10, 44], "측정결과", "TableStyleLight9")
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
        for c in row[3:10]:
            c.alignment = Alignment(horizontal="right")
        if row[10].value:
            for c in row:
                c.fill = NOTE_FILL
    dv = DataValidation(type="list", formula1="=지점정보!$C$2:$C$300", allow_blank=False,
                        error="지점정보 시트에 등록된 지점명만 입력할 수 있습니다.", showErrorMessage=True)
    dv.add("C2:C50000")
    ws.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1='"검토필요"', allow_blank=True)
    dv2.add("K2:K50000")
    ws.add_data_validation(dv2)

    ws = _sheet(wb, "지점정보", ST_COLS, db["stations"], [8, 8, 12, 10, 12, 9, 50, 10, 10, 9, 26, 18], "지점정보")
    dv3 = DataValidation(type="list", formula1='"확인,측정망,브이월드,추정,미입력"', allow_blank=True)
    dv3.add("J2:J300")
    ws.add_data_validation(dv3)
    for row in ws.iter_rows(min_row=2):
        if row[9].value == "추정":
            row[9].fill = NOTE_FILL
    _sheet(wb, "항목정보", ITEM_COLS, db["items"], [8, 26, 8, 9, 46], "항목정보")
    _sheet(wb, "기준값", LIM_COLS, db["limits"], [8, 12, 9, 10, 9, 56], "기준값")
    _sheet(wb, "연평균색도", YEAR_COLS, db["yearly"], [14, 10, 56, 8, 10, 22], "연평균색도")
    ws = _sheet(wb, "설정", SET_COLS, db["settings"], [18, 42, 90], "설정")
    for i, r in enumerate(db["settings"], start=2):
        choice = next((c for k, _, _, c in SETTINGS if k == r.get("설정")), None)
        if choice:
            dv = DataValidation(type="list", formula1='"' + ",".join(choice) + '"', allow_blank=False)
            dv.add(f"B{i}")
            ws.add_data_validation(dv)
    _sheet(wb, "측정망지점", NET_COLS, db["network"], [8, 12, 10, 34, 30, 10, 10, 10], "측정망지점")
    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"[오류] '{path.name}' 파일이 엑셀에서 열려 있습니다. 엑셀을 닫고 다시 실행하세요.")


# ── 기존(월별 시트) 양식 읽기 ─────────────────────────────
def _cell_value(v):
    """원본 셀 값: 숫자는 숫자, '7.3 ' 같은 문자 숫자는 숫자로, '<0.5'는 문자 그대로 보존."""
    if isinstance(v, str):
        s = v.strip()
        if s in ("", "-", "—"):
            return None
        if LOQ_RE.match(s) or s.upper() == "ND":
            return s
        n = to_num(s)
        return n if n is not None else s
    return v


def read_legacy(path, only_months=None):
    """YYYYMM 이름 시트들을 읽어 (측정행 목록, 지점 메타 dict, 지점 순서, 연평균색도 목록) 반환."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows, stations, order = [], {}, []
    sheets = sorted([ws for ws in wb.worksheets if re.fullmatch(r"\d{6}", ws.title.strip())], key=lambda w: w.title)
    for ws in sheets:
        ym = f"{ws.title[:4]}-{ws.title[4:6]}"
        if only_months and ym not in only_months:
            continue
        hdr_row = None
        for r in range(1, 10):
            vals = [str(c.value).strip() if c.value else "" for c in ws[r]]
            if "시료번호" in vals:
                hdr_row, hdr = r, vals
                break
        if not hdr_row:
            print(f"[경고] {ws.title}: 헤더(시료번호)를 찾지 못해 건너뜀")
            continue
        col = {h: i for i, h in enumerate(hdr) if h}
        addr_i = col.get("주 소", col.get("주소", 3))
        river = None
        for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            st = r[col["시료번호"]] if len(r) > col["시료번호"] else None
            if not st:
                continue
            st = str(st).strip()
            if "하천명" in col and r[col["하천명"]]:
                river = clean_river(r[col["하천명"]])
            addr = r[addr_i]
            date = r[col["시료채취일"]] if "시료채취일" in col else None
            note = r[len(hdr) - 1] if len(r) >= len(hdr) else None
            if st not in stations:
                order.append(st)
                stations[st] = {"하천명": river, "주소": str(addr).strip() if addr else "", "비고": ""}
            if note and str(note).strip() in ("변경지점", "신규지점"):
                stations[st]["비고"] = str(note).strip()
            rec = {"조사연월": ym, "채취일": date if isinstance(date, datetime) else None, "지점명": st,
                   "검토표시": None, "비고": None}
            for it in ITEMS:
                rec[it] = _cell_value(r[col[it]]) if it in col else None
            rows.append(rec)
    yearly = []
    if "연도별 평균색도" in wb.sheetnames:
        ws = wb["연도별 평균색도"]
        hdr = [c.value for c in ws[2]]
        years = [(i, int(h)) for i, h in enumerate(hdr) if isinstance(h, (int, float))]
        river = None
        for r in ws.iter_rows(min_row=3, values_only=True):
            if not r[2]:
                continue
            if r[1]:
                river = clean_river(r[1])
            for i, y in years:
                v = to_num(r[i])
                if v is not None:
                    yearly.append({"지점명": str(r[2]).strip(), "하천명": river, "주소": str(r[3] or "").strip(),
                                   "연도": y, "평균색도": round(v, 2), "구간": (r[12] or "") if len(r) > 12 else ""})
    return rows, stations, order, yearly
