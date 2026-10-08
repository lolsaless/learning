# -*- coding: utf-8 -*-
"""
시험장비관리대장(HWP 5.0 바이너리, 확장자 .hwpx 포함) -> 장비 DB 변환 스크립트

설치:  pip install olefile pandas openpyxl Pillow
실행:  python equipment_extract.py                      (아래 설정값 사용)
       python equipment_extract.py 파일.hwpx --out 결과폴더
       python equipment_extract.py --limit 5             (앞 5건만 시험 실행)
       python equipment_extract.py --debug-record 04-장비-43   (해당 장비 셀 구조 출력)
       python equipment_extract.py --thumbs               (엑셀 photos 시트에 썸네일 삽입)

결과물(OUT_DIR):
  equipment.db / equipment.xlsx / csv/*.csv / images/*.png / report.txt
"""
import argparse
import io
import json
import os
import re
import sqlite3
import struct
import sys
import traceback
import zlib

import olefile
import pandas as pd

# ─────────────────────────── 설정 ───────────────────────────
HWP_PATH = r"장비관리대장.hwpx"      # 입력 파일 (명령행 인자로 덮어쓸 수 있음)
OUT_DIR = r"output"                  # 출력 폴더
EXCEL_THUMBS = False                 # True: 엑셀 photos 시트에 썸네일 삽입 (--thumbs 와 동일)
# ────────────────────────────────────────────────────────────

# 라벨 정규화 사전: (공백 제거한 라벨의 시작 문자열, 필드명). 새 양식이 나오면 여기에만 추가하면 된다.
LABELS = [
    ("형식및규격", "spec"),
    ("모델명", "spec"),
    ("제작회사", "maker"),
    ("도입년월일", "date"),
    ("도입가격", "price"),
    ("회사명", "vendor"),
    ("주소", "addr"),
    ("전화번호", "tel"),
    ("용도및특기사항", "use"),
]
# 장비구성내역 표 머리 라벨 -> 표준 컬럼
COMP_HEAD = [("연번", "seq"), ("장비명", "name"), ("구성품", "name"), ("품명", "name"),
             ("수량", "qty"), ("단위", "unit")]
CHECKED = "■☑∎✓✔"          # 체크된 상자로 보는 기호
UNCHECKED = "□☐"

nospace = lambda s: re.sub(r"\s", "", s or "")


# ═════════════ 1. HWP 저수준 파싱 (OLE → 레코드 트리) ═════════════
def read_records(data):
    """레코드 스트림을 (tag, level, payload)로 순회"""
    p = 0
    while p + 4 <= len(data):
        h = struct.unpack_from("<I", data, p)[0]
        p += 4
        tag, lvl, size = h & 0x3FF, (h >> 10) & 0x3FF, h >> 20
        if size == 0xFFF:
            size = struct.unpack_from("<I", data, p)[0]
            p += 4
        yield tag, lvl, data[p:p + size]
        p += size


# 문단 텍스트 안의 컨트롤 문자 중 뒤에 14바이트 부가 데이터가 붙는 것
_EXTRA14 = set(range(1, 10)) | {11, 12} | set(range(14, 24))


def decode_text(b):
    out, i = [], 0
    while i + 2 <= len(b):
        c = struct.unpack_from("<H", b, i)[0]
        i += 2
        if c < 32:
            if c in _EXTRA14:
                i += 14
                if c == 9:
                    out.append("\t")
            elif c in (10, 13):
                out.append("\n")
            elif c == 24:
                out.append("-")
            elif c in (30, 31):
                out.append(" ")
        else:
            out.append(chr(c))
    return "".join(out)


class Node:
    __slots__ = ("tag", "lvl", "data", "kids")

    def __init__(self, tag, lvl, data):
        self.tag, self.lvl, self.data, self.kids = tag, lvl, data, []


def build_tree(data):
    root = Node(0, -1, b"")
    stack = [root]
    for tag, lvl, payload in read_records(data):
        n = Node(tag, lvl, payload)
        while stack[-1].lvl >= lvl:
            stack.pop()
        stack[-1].kids.append(n)
        stack.append(n)
    return root


def walk(n):
    yield n
    for k in n.kids:
        yield from walk(k)


def para_text(p):
    return "".join(decode_text(k.data) for k in p.kids if k.tag == 67)


def pic_ids(n):
    """노드 하위의 그림(tag 85) BinItem ID들 (문서 순서)"""
    return [struct.unpack_from("<H", x.data, 71)[0] for x in walk(n) if x.tag == 85]


def open_hwp(path):
    ole = olefile.OleFileIO(path)
    if not ole.exists("FileHeader"):
        raise ValueError("HWP(OLE) 형식이 아닙니다. 한글에서 .hwp로 다시 저장해 보세요.")
    header = ole.openstream("FileHeader").read()
    compressed = struct.unpack_from("<I", header, 36)[0] & 1

    def read_stream(name):
        d = ole.openstream(name).read()
        return zlib.decompress(d, -15) if compressed else d

    return ole, read_stream


def load_bindata_map(ole, read_stream):
    """BinItem ID -> OLE 스트림 이름 (대소문자 무시)"""
    streams = {"/".join(s): 1 for s in ole.listdir() if s[0] == "BinData"}
    lower = {k.lower(): k for k in streams}
    m = {}
    for tag, _, b in read_records(read_stream("DocInfo")):
        if tag == 18:
            _, bid, el = struct.unpack_from("<HHH", b, 0)
            ext = b[6:6 + el * 2].decode("utf-16le")
            name = f"BinData/BIN{bid:04X}.{ext}"
            m[bid] = lower.get(name.lower())
    return m


# ═════════════ 2. 표/셀 파싱 ═════════════
def parse_cells(tbl_ctrl):
    """표 컨트롤 -> 셀 리스트. 셀 좌표는 LIST_HEADER 오프셋 8의 (col,row,colSpan,rowSpan)"""
    cells, cur = [], None
    for k in tbl_ctrl.kids:
        if k.tag == 72:
            col, row, cs, rs = struct.unpack_from("<HHHH", k.data, 8)
            cur = dict(col=col, row=row, cs=cs, rs=rs, paras=[], pics=[])
            cells.append(cur)
        elif k.tag == 66 and cur is not None:
            cur["paras"].append(para_text(k).rstrip("\n"))
            cur["pics"] += pic_ids(k)
    for c in cells:
        c["text"] = "\n".join(c["paras"])
    return cells


def below(cells, lab):
    """라벨 셀 바로 아래(같은 열, row = 라벨.row + 라벨.rowSpan)의 값 셀"""
    return [c for c in cells if c["col"] == lab["col"] and c["row"] == lab["row"] + lab["rs"]]


def top_blocks(root):
    """최상위 문단 -> (index, 문단텍스트, [표 셀 리스트...])"""
    out = []
    for i, p in enumerate(root.kids):
        tbls = [parse_cells(c) for c in p.kids if c.tag == 71 and c.data[:4] == b" lbt"]
        out.append((i, para_text(p).strip(), tbls))
    return out


def group_records(blocks):
    """'관리대장' 제목 표가 나오면 새 장비 시작. 각 구간에서 앞면표/뒷면표/헤더줄을 찾는다."""
    starts = [i for i, _, tb in blocks
              if tb and "관리대장" in nospace(" ".join(c["text"] for c in tb[0][:3]))]
    recs = []
    for n, s in enumerate(starts):
        e = starts[n + 1] if n + 1 < len(starts) else len(blocks)
        recs.append(blocks[s:e])
    return recs


# ═════════════ 3. 값 정규화 ═════════════
def nz(s):
    """빈 값은 None"""
    s = (s or "").strip()
    return s or None


def norm_date(raw):
    if not raw:
        return None, None
    dates = re.findall(r"(\d{4})\s*\.\s*(\d{1,2})\s*\.\s*(\d{1,2})", raw)
    iso = f"{int(dates[0][0]):04d}-{int(dates[0][1]):02d}-{int(dates[0][2]):02d}" if dates else None
    note = None
    if len(dates) > 1 or "~" in raw or "(" in raw:
        note = "여러 날짜/기간 표기 - 첫 날짜만 정규화"
    elif not dates:
        note = "날짜 형식 인식 실패"
    return iso, note


def norm_price(raw):
    """(price_krw, note). 해석이 애매하면 (None, 사유)"""
    if not raw:
        return None, None
    t = raw.replace("원", "").strip()
    num = lambda s: int(re.sub(r"[^\d]", "", s))
    if re.fullmatch(r"\d{1,3}(?:[ ,]\d{3})*|\d+", t):
        return num(t), None
    if "+" in t:
        parts = re.findall(r"\d{1,3}(?:[ ,]\d{3})+|\d+", t)
        tot = sum(num(x) for x in parts)
        return None, f"합산 표기 (합계 {tot:,})"
    if re.search(r"[x×X]\s*\d+\s*대?", t):
        return None, "단가×수량 표기"
    if "(" in t:
        first = re.findall(r"\d{1,3}(?:[ ,]\d{3})+|\d+", t)
        return None, f"괄호 병기 값 있음 (첫 값 {first[0]})" if first else "괄호 표기"
    return None, "가격 형식 인식 실패"


def split_maker(raw):
    """'SANYO (일본)' -> ('SANYO','일본'). '(주)' 같은 법인 표기는 국가로 보지 않음"""
    if not raw:
        return None, None
    m = re.match(r"^(.*?)\s*\(([^()]*)\)\s*$", raw)
    if m and m.group(2).strip() not in ("주", "유", "사", "재", "합", "") and len(m.group(2).strip()) >= 2:
        return nz(m.group(1)), nz(m.group(2))
    return raw, None


def split_tel(raw):
    if not raw:
        return None, None
    m = (re.search(r"\b0\d{1,2}[-\s]\d{3,4}\s?-?\s?\d{4}\b", raw)      # 02-592-2997, 02 6340 6300, 02-3281 -3376
         or re.search(r"\b\d{4}-\d{4}\b", raw))                        # 1544-7777
    if not m:
        return None, None
    rest = (raw[:m.start()] + " " + raw[m.end():]).strip()
    return m.group(0), nz(re.sub(r"\s+", " ", rest))


def parse_condition(items):
    """['■ 신품','□중고품'] -> ('신품'|'중고품'|None, 경고)"""
    hit = []
    for t in items:
        if "신품" in t and any(ch in t for ch in CHECKED):
            hit.append("신품")
        elif "중고" in t and any(ch in t for ch in CHECKED):
            hit.append("중고품")
    if len(hit) == 1:
        return hit[0], None
    return None, "신품/중고품 판정 불가"


def perf_extras(perf):
    if not perf:
        return None, None, None
    cls = re.findall(r"\b\d{8}-\d{8}\b", perf)
    sn = re.findall(r"(?:S\s*/\s*N|S\.N\.?|Serial\s*No\.?)\s*[:.]?\s*([A-Za-z0-9][A-Za-z0-9\-_]*)", perf, re.I)
    tr = [ln.strip() for ln in perf.splitlines() if "이관" in ln or "관리전환" in ln]
    return (nz("; ".join(dict.fromkeys(cls))), nz("; ".join(dict.fromkeys(sn))), nz("\n".join(tr)))


# ═════════════ 4. 장비 1건 파싱 ═════════════
def parse_record(seg, rid):
    rec = {"record_id": rid}
    warns = []
    # 제목 표 / 헤더 줄
    title_cells = seg[0][2][0]
    title = next((nospace(c["text"]) for c in title_cells if "관리대장" in nospace(c["text"])), "")
    rec["header_title"] = title
    rec["is_marked_x"] = int("(X)" in title.upper())
    hdr = next((t for _, t, _ in seg if "물품명" in t), "")
    m = re.match(r"물품명\s*:\s*(.*?)\s{2,}관리번호\s*\((.*?)\)(.*)$", hdr, re.S)
    if m:
        rec["name"], rec["mgmt_no"], rec["header_tail"] = nz(m.group(1)), nz(m.group(2)), nz(m.group(3))
    else:
        rec["name"] = rec["mgmt_no"] = rec["header_tail"] = None
        warns.append("헤더 줄(물품명/관리번호) 파싱 실패: " + hdr[:40])
    rec["mgmt_no_is_range"] = int(bool(rec["mgmt_no"] and "~" in rec["mgmt_no"]))
    if rec["mgmt_no_is_range"]:
        warns.append("관리번호 범위 표기(2대 이상 묶음)")
    if rec["is_marked_x"]:
        warns.append("헤더에 (X) 표기")
    if rec["header_tail"]:
        warns.append("헤더 뒤 추가 문구: " + rec["header_tail"])

    main = next((tb[0] for _, _, tb in seg if tb and any(
        nospace(c["text"]).startswith(("형식", "모델명")) for c in tb[0])), None)
    hist = next((tb[0] for _, _, tb in seg if tb and any(
        "수리" in nospace(c["text"]) for c in tb[0][:2])), None)
    if main is None:
        raise ValueError("앞면 표를 찾지 못함")

    # 기본정보: 라벨(상단 2행) 바로 아래 셀이 값
    raw = {}
    for c in main:
        if c["row"] > 1:
            continue
        key = nospace(c["text"])
        for pre, field in LABELS:
            if key.startswith(pre) and field not in raw:
                v = below(main, c)
                raw[field] = nz(" ".join(x["text"].replace("\n", " ") for x in v))
    for f in ("spec", "maker", "date", "price", "vendor", "addr", "tel", "use"):
        raw.setdefault(f, None)

    rec["spec"] = raw["spec"]
    rec["maker_raw"] = raw["maker"]
    rec["maker"], rec["country"] = split_maker(raw["maker"])
    rec["date_raw"] = raw["date"]
    rec["date_iso"], rec["date_note"] = norm_date(raw["date"])
    rec["price_raw"] = raw["price"]
    rec["price_krw"], rec["price_note"] = norm_price(raw["price"])
    rec["vendor"], rec["addr"] = raw["vendor"], raw["addr"]
    rec["tel_raw"] = raw["tel"]
    rec["tel_first"], rec["contact_note"] = split_tel(raw["tel"])
    rec["use"] = raw["use"]

    # 입고시 상태
    lab = next((c for c in main if nospace(c["text"]) == "입고시상태"), None)
    items = []
    if lab:
        items = [x["text"] for x in main
                 if x["row"] >= lab["row"] + lab["rs"]
                 and lab["col"] <= x["col"] < lab["col"] + lab["cs"] + 3
                 and ("신품" in x["text"] or "중고" in x["text"])]
    rec["condition_raw"] = " / ".join(i.strip() for i in items) or None
    rec["condition"], w = parse_condition(items)
    if w:
        warns.append(w)

    # 주요 성능 및 특징 (라벨 문구는 제거)
    pf = next((c for c in main if nospace(c["text"]).startswith("주요성능")), None)
    perf = None
    if pf:
        perf = re.sub(r"^\s*주요\s*성능\s*및\s*특징[ \t]*", "", "\n".join(pf["paras"])).strip("\n")
        perf = "\n".join(l.rstrip() for l in perf.splitlines()).strip() or None
    rec["perf"] = perf
    rec["item_class_no"], rec["serial_no"], rec["transfer_note"] = perf_extras(perf)

    # 장비구성내역
    comps = []
    hl = next((c for c in main if nospace(c["text"]) == "연번"), None)
    if hl:
        colmap = {}
        for c in main:
            if c["row"] == hl["row"]:
                k = nospace(c["text"])
                for pre, std in COMP_HEAD:
                    if k.startswith(pre):
                        colmap[c["col"]] = std
        rows = {}
        for c in main:
            if c["row"] > hl["row"] and c["col"] in colmap:
                rows.setdefault(c["row"], {})[colmap[c["col"]]] = c["text"].strip()
        for r in sorted(rows):
            d = rows[r]
            if any(v for k, v in d.items() if k != "seq"):      # 연번만 있는 빈 행 제외
                comps.append({"seq": nz(d.get("seq")), "name": nz(d.get("name")),
                              "qty": nz(d.get("qty")), "unit": nz(d.get("unit"))})
    rec["components"] = comps

    # 사진 (앞면 표 안의 그림, 중복 ID 제거, 문서 순서 유지)
    ids = []
    for c in main:
        for p in c["pics"]:
            if p not in ids:
                ids.append(p)
    rec["pic_ids"] = ids

    # 뒷면 이력
    hist_rows = []
    if hist is None:
        warns.append("뒷면 이력 표를 찾지 못함")
    else:
        dh = [c for c in hist if nospace(c["text"]) == "일자"]
        ch = [c for c in hist if nospace(c["text"]) == "내용"]
        for d in dh:
            nxt = [c for c in ch if c["col"] > d["col"]]
            if not nxt:
                continue
            cc = min(nxt, key=lambda c: c["col"])["col"]
            start = d["row"] + d["rs"]
            for r in sorted({c["row"] for c in hist if c["row"] >= start}):
                dt = next((c["text"].strip() for c in hist if c["row"] == r and c["col"] == d["col"]), "")
                ct = next((c["text"].strip() for c in hist if c["row"] == r and c["col"] == cc), "")
                if not dt and not ct:
                    continue
                cert = re.search(r"성적서\s*번호\s*([^\n]+)", ct)
                unc = re.search(r"불확도\s*([^\n]+)", ct)
                hist_rows.append({"hist_date": nz(dt), "content": nz(ct),
                                  "is_calibration": int("교정" in ct),
                                  "cert_no": nz(cert.group(1)) if cert else None,
                                  "uncertainty": nz(unc.group(1)) if unc else None})
    rec["history"] = hist_rows

    # 필수필드 경고
    for f, lbl in (("name", "물품명"), ("mgmt_no", "관리번호"), ("date_raw", "도입년월일"), ("price_raw", "도입가격")):
        if not rec.get(f):
            warns.append(f"필수필드 누락: {lbl}")
    for f, lbl in (("spec", "형식및규격"), ("maker_raw", "제작회사"), ("vendor", "도입처 회사명"),
                   ("addr", "도입처 주소"), ("tel_raw", "전화번호"), ("use", "용도")):
        if not rec.get(f):
            warns.append(f"값 없음: {lbl}")
    if rec["price_note"]:
        warns.append("가격: " + rec["price_note"])
    if rec["date_note"]:
        warns.append("날짜: " + rec["date_note"])
    if not ids:
        warns.append("사진 없음")
    rec["warnings"] = warns
    return rec


# ═════════════ 5. 사진 저장 ═════════════
def safe_name(s):
    return re.sub(r'[\\/:*?"<>|~\s]+', "_", s or "unknown")


def save_photos(recs, ole, binmap, out_dir):
    from PIL import Image
    img_dir = os.path.join(out_dir, "images")
    os.makedirs(img_dir, exist_ok=True)
    rows = []
    for rec in recs:
        for n, bid in enumerate(rec["pic_ids"], 1):
            stream = binmap.get(bid)
            if not stream:
                rec["warnings"].append(f"사진 BinItem {bid} 스트림 없음")
                continue
            data = ole.openstream(stream).read()
            ext = os.path.splitext(stream)[1].lower() or ".png"
            fn = f"{safe_name(rec['mgmt_no'] or rec['record_id'])}_{n}{ext}"
            with open(os.path.join(img_dir, fn), "wb") as f:
                f.write(data)               # 원본 바이트 그대로(재인코딩 없음)
            try:
                w, h = Image.open(io.BytesIO(data)).size
            except Exception:
                w = h = None
            rows.append({"record_id": rec["record_id"], "mgmt_no": rec["mgmt_no"], "seq": n,
                         "file_path": os.path.join("images", fn).replace("\\", "/"),
                         "bin_id": bid, "bin_stream": stream, "width": w, "height": h})
    return rows


# ═════════════ 6. 저장 (SQLite / CSV / Excel) ═════════════
EQ_COLS = ["record_id", "mgmt_no", "mgmt_no_is_range", "name", "header_title", "is_marked_x", "header_tail",
           "spec", "maker_raw", "maker", "country", "date_raw", "date_iso", "date_note",
           "price_raw", "price_krw", "price_note", "vendor", "addr", "tel_raw", "tel_first", "contact_note",
           "use", "condition_raw", "condition", "perf", "item_class_no", "serial_no", "transfer_note",
           "n_components", "n_history", "n_photos", "warnings"]
CO_COLS = ["record_id", "mgmt_no", "seq", "name", "qty", "unit"]
HI_COLS = ["record_id", "mgmt_no", "hist_date", "content", "is_calibration", "cert_no", "uncertainty"]
PH_COLS = ["record_id", "mgmt_no", "seq", "file_path", "bin_id", "bin_stream", "width", "height"]

DDL = """
CREATE TABLE equipment (record_id INTEGER PRIMARY KEY, mgmt_no TEXT, mgmt_no_is_range INTEGER, name TEXT,
 header_title TEXT, is_marked_x INTEGER, header_tail TEXT, spec TEXT, maker_raw TEXT, maker TEXT, country TEXT,
 date_raw TEXT, date_iso TEXT, date_note TEXT, price_raw TEXT, price_krw INTEGER, price_note TEXT, vendor TEXT,
 addr TEXT, tel_raw TEXT, tel_first TEXT, contact_note TEXT, use TEXT, condition_raw TEXT, condition TEXT,
 perf TEXT, item_class_no TEXT, serial_no TEXT, transfer_note TEXT, n_components INTEGER, n_history INTEGER,
 n_photos INTEGER, warnings TEXT);
CREATE INDEX idx_eq_mgmt ON equipment(mgmt_no);
CREATE TABLE components (id INTEGER PRIMARY KEY AUTOINCREMENT, record_id INTEGER REFERENCES equipment(record_id),
 mgmt_no TEXT, seq TEXT, name TEXT, qty TEXT, unit TEXT);
CREATE TABLE history (id INTEGER PRIMARY KEY AUTOINCREMENT, record_id INTEGER REFERENCES equipment(record_id),
 mgmt_no TEXT, hist_date TEXT, content TEXT, is_calibration INTEGER, cert_no TEXT, uncertainty TEXT);
CREATE TABLE photos (id INTEGER PRIMARY KEY AUTOINCREMENT, record_id INTEGER REFERENCES equipment(record_id),
 mgmt_no TEXT, seq INTEGER, file_path TEXT, bin_id INTEGER, bin_stream TEXT, width INTEGER, height INTEGER);
"""


def build_tables(recs, photo_rows):
    eq, co, hi = [], [], []
    for r in recs:
        row = {c: r.get(c) for c in EQ_COLS}
        row["n_components"], row["n_history"] = len(r["components"]), len(r["history"])
        row["n_photos"] = sum(1 for p in photo_rows if p["record_id"] == r["record_id"])
        row["warnings"] = json.dumps(r["warnings"], ensure_ascii=False) if r["warnings"] else None
        eq.append(row)
        co += [{"record_id": r["record_id"], "mgmt_no": r["mgmt_no"], **c} for c in r["components"]]
        hi += [{"record_id": r["record_id"], "mgmt_no": r["mgmt_no"], **h} for h in r["history"]]
    return (pd.DataFrame(eq, columns=EQ_COLS), pd.DataFrame(co, columns=CO_COLS),
            pd.DataFrame(hi, columns=HI_COLS), pd.DataFrame(photo_rows, columns=PH_COLS))


def save_sqlite(path, dfs):
    if os.path.exists(path):
        os.remove(path)
    con = sqlite3.connect(path)
    con.executescript(DDL)
    for name, df in zip(("equipment", "components", "history", "photos"), dfs):
        cols = [c for c in df.columns]
        con.executemany(f"INSERT INTO {name} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                        [tuple(None if pd.isna(v) else (int(v) if hasattr(v, 'item') and isinstance(v.item(), int) else v)
                               for v in row) for row in df.itertuples(index=False)])
    con.commit()
    con.close()


def save_excel(path, dfs, out_dir, thumbs):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    names = ("equipment", "components", "history", "photos")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in zip(names, dfs):
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for j, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [max(len(l) for l in str(v).split("\n")) for v in df[col].dropna()][:2000] or [8])
                ws.column_dimensions[get_column_letter(j)].width = min(max(width * 1.3 + 2, 8), 60)
                ws.cell(1, j).font = Font(name="Malgun Gothic", bold=True)
                for i in range(2, ws.max_row + 1):
                    ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
                    ws.cell(i, j).font = Font(name="Malgun Gothic", size=10)
            if name == "photos" and thumbs and len(df):
                from openpyxl.drawing.image import Image as XImg
                from PIL import Image
                col = len(df.columns) + 1
                ws.column_dimensions[get_column_letter(col)].width = 22
                ws.cell(1, col, "thumbnail").font = Font(name="Malgun Gothic", bold=True)
                for i, fp in enumerate(df["file_path"], 2):
                    full = os.path.join(out_dir, fp)
                    try:
                        with Image.open(full) as im:
                            w, h = im.size
                        img = XImg(full)
                        s = 110 / max(w, h)
                        img.width, img.height = int(w * s), int(h * s)
                        ws.add_image(img, f"{get_column_letter(col)}{i}")
                        ws.row_dimensions[i].height = img.height * 0.78 + 4
                    except Exception:
                        pass


# ═════════════ 7. 리포트 ═════════════
def write_report(path, recs, errors, dfs, photo_rows):
    eq = dfs[0]
    L = ["=== 시험장비관리대장 추출 리포트 ===", f"인식된 장비 수: {len(recs)}",
         f"오류로 건너뛴 구간: {len(errors)}"]
    for e in errors:
        L.append(f"  - 구간 {e[0]}: {e[1]}")
    L.append(f"장비구성내역 행 수: {len(dfs[1])} / 이력 행 수: {len(dfs[2])} (교정 {int(dfs[2]['is_calibration'].sum()) if len(dfs[2]) else 0}) / 사진 파일 수: {len(photo_rows)}")

    def lst(title, cond):
        ids = [r["mgmt_no"] or f"#{r['record_id']}" for r in recs if cond(r)]
        L.append(f"\n[{title}] {len(ids)}건")
        L.append("  " + ", ".join(ids) if ids else "  (없음)")

    lst("필수필드 누락(물품명/관리번호/도입년월일/도입가격)", lambda r: any(w.startswith("필수필드 누락") for w in r["warnings"]))
    lst("신품/중고 판정 불가", lambda r: r["condition"] is None)
    lst("도입가격 정규화 불가(원문 확인 필요)", lambda r: r["price_raw"] and r["price_krw"] is None)
    lst("사진 없음", lambda r: not r["pic_ids"])
    lst("사진 2장 이상", lambda r: len(r["pic_ids"]) >= 2)
    lst("주요 성능 및 특징 비어 있음", lambda r: not r["perf"])
    lst("(X) 표기 헤더", lambda r: r["is_marked_x"])
    lst("관리번호 범위 표기", lambda r: r["mgmt_no_is_range"])
    L.append("\n[경고가 있는 장비 전체 목록]")
    for r in recs:
        if r["warnings"]:
            L.append(f"  {r['mgmt_no']}: " + " | ".join(r["warnings"]))
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


# ═════════════ 8. 디버그 ═════════════
def debug_record(seg, rec):
    for _, _, tbs in seg:
        for cells in tbs:
            print(f"--- 표 (셀 {len(cells)}개)")
            for c in cells:
                print(f"  col={c['col']:>2} row={c['row']:>2} cs={c['cs']} rs={c['rs']} pics={c['pics'] or ''} | {c['text']!r}"[:200])
    print("\n=== 파싱 결과")
    print(json.dumps({k: v for k, v in rec.items()}, ensure_ascii=False, indent=1, default=str))


# ═════════════ main ═════════════
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("hwp", nargs="?", default=HWP_PATH)
    ap.add_argument("--out", default=OUT_DIR)
    ap.add_argument("--limit", type=int, default=0, help="앞 N건만 처리")
    ap.add_argument("--debug-record", default=None, help="관리번호를 주면 해당 장비의 셀 구조와 결과 출력")
    ap.add_argument("--thumbs", action="store_true", default=EXCEL_THUMBS)
    a = ap.parse_args()

    if not os.path.exists(a.hwp):
        sys.exit(f"입력 파일을 찾을 수 없습니다: {a.hwp}")
    os.makedirs(a.out, exist_ok=True)

    ole, read_stream = open_hwp(a.hwp)
    root = build_tree(read_stream("BodyText/Section0"))
    binmap = load_bindata_map(ole, read_stream)
    segs = group_records(top_blocks(root))
    print(f"장비 구간 {len(segs)}건 인식 / 내장 이미지 {len(binmap)}개")
    if a.limit:
        segs = segs[:a.limit]

    recs, errors = [], []
    for n, seg in enumerate(segs, 1):
        try:
            rec = parse_record(seg, len(recs) + 1)
            recs.append(rec)
            if a.debug_record and rec["mgmt_no"] == a.debug_record:
                debug_record(seg, rec)
        except Exception as e:
            errors.append((n, f"{type(e).__name__}: {e}"))
            traceback.print_exc()
    if a.debug_record:
        return

    photo_rows = save_photos(recs, ole, binmap, a.out)
    dfs = build_tables(recs, photo_rows)
    save_sqlite(os.path.join(a.out, "equipment.db"), dfs)
    csv_dir = os.path.join(a.out, "csv")
    os.makedirs(csv_dir, exist_ok=True)
    for name, df in zip(("equipment", "components", "history", "photos"), dfs):
        df.to_csv(os.path.join(csv_dir, f"{name}.csv"), index=False, encoding="utf-8-sig")
    save_excel(os.path.join(a.out, "equipment.xlsx"), dfs, a.out, a.thumbs)
    write_report(os.path.join(a.out, "report.txt"), recs, errors, dfs, photo_rows)
    print(f"완료: 장비 {len(recs)}건, 구성 {len(dfs[1])}행, 이력 {len(dfs[2])}행, 사진 {len(photo_rows)}장 -> {os.path.abspath(a.out)}")


if __name__ == "__main__":
    main()
